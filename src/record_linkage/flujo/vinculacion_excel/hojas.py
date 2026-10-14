"""Contenido de cada hoja del Excel de entrega, en el orden en que se escriben.

El ORDEN de las hojas y el de las escrituras es parte del contrato del archivo: cambiar cualquiera
cambia el libro (índices de estilo incluidos). ``escribir_hojas`` es el único sitio que lo fija.
"""

from __future__ import annotations

import re
from collections.abc import Callable, Collection, Mapping, Sequence
from functools import partial
from typing import cast

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from record_linkage.flujo.vinculacion.config import ConfigVinculacion, Fuente
from record_linkage.flujo.vinculacion.resultado import Resultado
from record_linkage.flujo.vinculacion_excel.estilos import (
    ALINEACION_ARRIBA_ENVUELTA,
    ALINEACION_DERECHA,
    ALINEACION_ENCABEZADO,
    ALINEACION_ENVUELTA,
    AUTOR_COMENTARIOS,
    COLORES_PESTANA,
    Estilos,
    altura_de_fila,
)
from record_linkage.flujo.vinculacion_excel.leeme import (
    hoja_leeme,
    prefijo_de_bloque,
    textos_leeme,
)
from record_linkage.flujo.vinculacion_excel.tablas import (
    colorear_columna,
    escribir_celda,
    escribir_tabla,
    formatear_tabla,
    hoja_desde_tabla,
)

__all__ = ["escribir_hojas"]

#: RESUMEN: fila del encabezado de sus tablas (los datos empiezan en la siguiente).
_FILA_ENCABEZADO = 3
#: RESUMEN: columna (G) donde empieza «Cobertura por país», a la derecha de los indicadores.
_COLUMNA_COBERTURA = 7
#: Cuántas advertencias de CALIDAD_DATOS se repiten al pie del RESUMEN.
_ADVERTENCIAS_EN_RESUMEN = 8
#: Decisiones que admite la lista desplegable de la columna DECISION de la hoja REVISION.
_LISTA_DE_DECISIONES = '"EMPRESA,GRUPO,NO_VINCULAR,RELACION_PROPIEDAD,MISMA_ENTIDAD,DISTINTA"'
#: Filas de más (después de la última de REVISION) a las que llega la lista desplegable.
_FILAS_EXTRA_DE_VALIDACION = 200
#: Columnas de REVISION que una persona llena.
_COLUMNAS_DE_ENTRADA = ("DECISION", "AUTOR", "FECHA", "RAZON")
#: Columnas de VALIDACION que se muestran como porcentaje.
_COLUMNAS_PORCENTUALES = ("PRECISION", "RECALL", "F1", "RECALL_CONTANDO_CANDIDATOS_C")
#: Ancho de las 14 primeras columnas de VALIDACION.
_ANCHOS_DE_VALIDACION = (16, 60, 40, 16, 40, 16, 16, 16, 16, 16, 16, 16, 16, 16)


# ── RESUMEN ──────────────────────────────────────────────────────────────────
def _rango_de_exportadores(letras: Mapping[str, str], ultima_fila: int, columna: str) -> str:
    """Rango de una columna de EXPORTADORES (sin el encabezado), listo para una fórmula."""
    return f"EXPORTADORES!${letras[columna]}$2:${letras[columna]}${ultima_fila}"


def _formula_de_indicador(
    indicador: str,
    etiqueta_ancla: str,
    rango: Callable[[str], str],
    columnas: Collection[str],
) -> str | None:
    """Fórmula viva sobre EXPORTADORES que reproduce un indicador del RESUMEN (``None`` si no tiene).

    El orden de las preguntas importa: la primera que acierta decide, aunque no dé fórmula.
    """
    limpio = indicador.strip()
    if indicador == "Empresas vinculadas a al menos una base":
        return f'=COUNTIF({rango("RESULTADO")},"VINCULADO")'
    if indicador.startswith(f"Empresas {etiqueta_ancla}"):
        return f"=COUNTA({rango('ID_EXPORTADOR')})"
    if limpio == "· con la MISMA EMPRESA en alguna base":
        return f'=COUNTIFS({rango("RESULTADO")},"VINCULADO",{rango("NIVEL_VINCULO")},"EMPRESA")'
    if limpio == "· solo con su GRUPO empresarial":
        return f'=COUNTIFS({rango("RESULTADO")},"VINCULADO",{rango("NIVEL_VINCULO")},"GRUPO")'
    if limpio.startswith("· vinculadas a ") and "todas" not in indicador:
        etiqueta = limpio[len("· vinculadas a ") :]
        return f'=COUNTIF({rango("FUENTES_VINCULADAS")},"*{etiqueta}*")'
    if indicador.endswith("de las empresas vinculadas"):
        columna = indicador.split(" de las empresas vinculadas")[0]
        if columna in columnas:
            return (
                f'=IFERROR(SUMIF({rango("RESULTADO")},"VINCULADO",{rango(columna)})'
                f"/SUM({rango(columna)}),0)"
            )
        return None
    if indicador == "Empresas que requieren revisión":
        return f'=COUNTIF({rango("REQUIERE_REVISION")},"SÍ")'
    return None


def _formulas_del_resumen(
    indicadores: pd.DataFrame, exportadores: pd.DataFrame, etiqueta_ancla: str
) -> dict[int, str]:
    """Fórmula de cada indicador del RESUMEN que se recalcula sobre EXPORTADORES, por índice del indicador."""
    ultima_fila = len(exportadores) + 1
    letras = {c: get_column_letter(i + 1) for i, c in enumerate(exportadores.columns)}
    rango = partial(_rango_de_exportadores, letras, ultima_fila)
    formulas: dict[int, str] = {}
    for k, r in indicadores.iterrows():
        formula = _formula_de_indicador(r["INDICADOR"], etiqueta_ancla, rango, letras)
        if formula is not None:
            formulas[cast(int, k)] = formula
    return formulas


def _formato_de_indicadores(hoja: Worksheet, estilos: Estilos, indicadores: pd.DataFrame) -> None:
    """Formato del valor (porcentaje o entero) y negrita en los indicadores que no van sangrados."""
    for k, r in indicadores.iterrows():
        # Hallazgo F3: se supone un índice 0..n-1 en el resumen; con otro, las filas no coinciden con las fórmulas.
        fila = _FILA_ENCABEZADO + 1 + cast(int, k)
        celda = hoja.cell(row=fila, column=3)
        celda.number_format = "0.0%" if r["UNIDAD"] == "%" else "#,##0"
        celda.alignment = ALINEACION_DERECHA
        if not str(r["INDICADOR"]).startswith("  "):
            hoja.cell(row=fila, column=2).font = estilos.negrita


def _cobertura_por_pais(hoja: Worksheet, estilos: Estilos, por_pais: pd.DataFrame) -> None:
    """Tabla «Cobertura por país», a la derecha de los indicadores."""
    col0 = _COLUMNA_COBERTURA
    hoja.cell(
        row=_FILA_ENCABEZADO - 1, column=col0, value="Cobertura por país"
    ).font = estilos.titulo(12)
    escribir_tabla(hoja, por_pais, fila0=_FILA_ENCABEZADO, col0=col0)
    for j, nombre in enumerate(por_pais.columns):
        celda = hoja.cell(row=_FILA_ENCABEZADO, column=col0 + j)
        celda.font, celda.fill = estilos.encabezado, estilos.relleno_encabezado
        celda.alignment = ALINEACION_ENCABEZADO
        hoja.column_dimensions[get_column_letter(col0 + j)].width = max(
            12, min(18, len(str(nombre)) // 2 + 5)
        )
        for i in range(_FILA_ENCABEZADO + 1, _FILA_ENCABEZADO + 1 + len(por_pais)):
            celda = hoja.cell(row=i, column=col0 + j)
            celda.font = estilos.base
            if str(nombre).startswith("PCT"):
                celda.number_format = "0.0%"
    hoja.column_dimensions[get_column_letter(col0)].width = 16
    hoja.column_dimensions["F"].width = 3
    hoja.row_dimensions[_FILA_ENCABEZADO].height = 42


def _advertencias_de_calidad(
    hoja: Worksheet, estilos: Estilos, calidad: pd.DataFrame, fila: int
) -> None:
    """Las primeras advertencias de CALIDAD_DATOS, bajo los indicadores, desde ``fila``."""
    hoja.cell(
        row=fila, column=1, value="Advertencias de los datos (detalle en CALIDAD_DATOS)"
    ).font = estilos.titulo(12)
    for r in calidad.head(_ADVERTENCIAS_EN_RESUMEN).itertuples(index=False):
        fila += 1
        escribir_celda(hoja, fila, 1, r.FUENTE).font = estilos.negrita
        escribir_celda(hoja, fila, 2, r.TEMA).font = estilos.negrita
        texto = f"{r.HALLAZGO} {r.IMPACTO}"
        escribir_celda(hoja, fila, 3, texto).font = estilos.base
        hoja.cell(row=fila, column=3).alignment = ALINEACION_ARRIBA_ENVUELTA
        for col in (1, 2):
            hoja.cell(row=fila, column=col).alignment = ALINEACION_ARRIBA_ENVUELTA
        hoja.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=5)
        hoja.row_dimensions[fila].height = altura_de_fila(len(texto) // 110)


def _hoja_resumen(libro: Workbook, estilos: Estilos, res: Resultado) -> None:
    """RESUMEN: indicadores con fórmulas vivas sobre EXPORTADORES, cobertura por país y advertencias."""
    hoja = libro.create_sheet("RESUMEN")
    hoja.sheet_properties.tabColor = COLORES_PESTANA["negocio"]
    hoja["A1"] = "Resumen del cruce"
    hoja["A1"].font = estilos.titulo(13)
    indicadores = res.resumen.copy()
    formulas = _formulas_del_resumen(
        indicadores, res.tablas["EXPORTADORES"], res.cfg.ancla.etiqueta
    )
    escribir_tabla(hoja, indicadores, fila0=_FILA_ENCABEZADO)
    for k, formula in formulas.items():
        hoja.cell(row=_FILA_ENCABEZADO + 1 + k, column=3, value=formula)
    formatear_tabla(
        hoja,
        indicadores,
        estilos,
        fila0=_FILA_ENCABEZADO,
        congelar=None,
        filtro=False,
        anchos_max=90,
    )
    _formato_de_indicadores(hoja, estilos, indicadores)
    for letra, ancho in (("A", 18), ("B", 58), ("C", 14), ("D", 12), ("E", 90)):
        hoja.column_dimensions[letra].width = ancho
    _cobertura_por_pais(hoja, estilos, res.por_pais)
    _advertencias_de_calidad(
        hoja, estilos, res.calidad, _FILA_ENCABEZADO + 1 + len(indicadores) + 2
    )


# ── EXPORTADORES ─────────────────────────────────────────────────────────────
#: Un sitio web: dominio con una terminación de letras (``zelvora.es``), con o sin esquema, ruta o puerto. Un texto como
#: 'N.A.', '1.5' o 'x@y.com' no lo es (E6 de F3.1b: antes todo texto con un punto y sin espacios se volvía enlace).
_ES_SITIO_WEB = re.compile(
    r"(?:https?://)?[\w-]+(?:\.[\w-]+)*\.[a-z]{2,}(?:[/:?#]\S*)?", re.IGNORECASE
)


def _enlaces_web(hoja: Worksheet, estilos: Estilos, exportadores: pd.DataFrame) -> None:
    """Convierte en hipervínculo el primer sitio web de cada celda de las columnas ``*_WEB``."""
    columnas = list(exportadores.columns)
    for nombre in [c for c in columnas if c.endswith("_WEB") or c == "CONTACTO_WEB"]:
        col = columnas.index(nombre) + 1
        for fila in range(2, len(exportadores) + 2):
            celda = hoja.cell(row=fila, column=col)
            sitio = str(celda.value or "").split(" | ")[0].strip()
            if _ES_SITIO_WEB.fullmatch(sitio):
                celda.hyperlink = sitio if sitio.startswith("http") else f"https://{sitio}"
                celda.font = estilos.enlace


def _comentarios_de_bloques(
    hoja: Worksheet, cfg: ConfigVinculacion, columnas: Sequence[str]
) -> None:
    """Comentario en el encabezado ``<DESTINO>_NIVEL`` de cada destino: qué es ese bloque de columnas."""
    for destino in cfg.destinos:
        pre = prefijo_de_bloque(destino)
        if f"{pre}_NIVEL" in columnas:
            hoja.cell(row=1, column=columnas.index(f"{pre}_NIVEL") + 1).comment = Comment(
                f"Bloque {pre}_*: registro principal de {destino.etiqueta}. Si {pre}_NIVEL = GRUPO, "
                f"sus datos son de OTRA sociedad del grupo (contexto, no atribuibles al exportador).",
                AUTOR_COMENTARIOS,
            )


def _ocultar_columnas_constantes(
    hoja: Worksheet, ancla: Fuente, exportadores: pd.DataFrame
) -> None:
    """Oculta las columnas de métrica o contexto de la ancla que valen lo mismo en todas las filas."""
    metricas = {
        ancla.nombres_contexto.get(m, m).upper().replace(" ", "_") for m in ancla.cols_metricas
    } | {ancla.nombres_contexto.get(m, m).upper().replace(" ", "_") for m in ancla.cols_contexto}
    columnas = list(exportadores.columns)
    # Un ausente cuenta como un valor más (E2 de F3.1b): ``astype(str)`` los dejaba en NaN con pandas 3 y en 'nan' con
    # pandas 2, y una columna con un valor y ausentes se ocultaba o no según la versión.
    for nombre in columnas:
        serie = exportadores[nombre]
        distintos = serie.astype(object).where(serie.notna(), "ausente").astype(str).nunique()
        if nombre in metricas and distintos <= 1 and len(exportadores) > 1:
            hoja.column_dimensions[get_column_letter(columnas.index(nombre) + 1)].hidden = True


def _hoja_exportadores(libro: Workbook, estilos: Estilos, res: Resultado) -> None:
    """EXPORTADORES: una fila por empresa ancla, con semáforos, hipervínculos y comentarios."""
    exportadores = res.tablas["EXPORTADORES"]
    hoja = hoja_desde_tabla(
        libro,
        estilos,
        "EXPORTADORES",
        exportadores,
        "negocio",
        congelar="C2",
        anchos_max=45,
        envolver=False,
    )
    advertencia, ok = estilos.relleno_advertencia, estilos.relleno_ok
    colorear_columna(
        hoja, exportadores, "RESULTADO", [("VINCULADO", ok), ("CANDIDATO", advertencia)]
    )
    colorear_columna(hoja, exportadores, "CONFIANZA", [("MEDIA", advertencia)])
    colorear_columna(hoja, exportadores, "REQUIERE_REVISION", [("SÍ", advertencia)])
    _enlaces_web(hoja, estilos, exportadores)
    columnas = list(exportadores.columns)
    _comentarios_de_bloques(hoja, res.cfg, columnas)
    _ocultar_columnas_constantes(hoja, res.cfg.ancla, exportadores)
    hoja.cell(row=1, column=columnas.index("RESULTADO") + 1).comment = Comment(
        "VINCULADO: al menos un vínculo EMPRESA o GRUPO (automático A/B o decidido). "
        "CANDIDATO POR REVISAR: solo candidatos C. SIN COINCIDENCIA: nada con evidencia.",
        AUTOR_COMENTARIOS,
    )


# ── ENLACES, REVISION y DECISIONES ───────────────────────────────────────────
def _hoja_enlaces(libro: Workbook, estilos: Estilos, res: Resultado) -> None:
    """ENLACES: todos los vínculos finales, uno por par empresa–registro."""
    enlaces = res.tablas["ENLACES"]
    hoja = hoja_desde_tabla(libro, estilos, "ENLACES", enlaces, "negocio", congelar="C2")
    colorear_columna(
        hoja,
        enlaces,
        "CONFIANZA",
        [("MEDIA", estilos.relleno_advertencia), ("ALTA", estilos.relleno_ok)],
    )


def _validacion_de_decision(hoja: Worksheet, columnas: Sequence[str], filas: int) -> int:
    """Lista desplegable de la columna DECISION; devuelve el número de esa columna (desde 1)."""
    validacion = DataValidation(
        type="list",
        formula1=_LISTA_DE_DECISIONES,
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Decisión no válida",
        error=(
            "Use EMPRESA, GRUPO, NO_VINCULAR o RELACION_PROPIEDAD (vínculos) "
            "o MISMA_ENTIDAD / DISTINTA (duplicados)."
        ),
    )
    hoja.add_data_validation(validacion)
    col = columnas.index("DECISION") + 1
    letra = get_column_letter(col)
    validacion.add(f"{letra}2:{letra}{filas + _FILAS_EXTRA_DE_VALIDACION}")
    return col


def _columnas_de_entrada(
    hoja: Worksheet, estilos: Estilos, columnas: Sequence[str], filas: int
) -> None:
    """Rellena en amarillo las columnas de REVISION que una persona llena (encabezado más oscuro)."""
    for nombre in _COLUMNAS_DE_ENTRADA:
        col = columnas.index(nombre) + 1
        hoja.cell(row=1, column=col).fill = estilos.relleno_encabezado_entrada
        for fila in range(2, filas + 2):
            hoja.cell(row=fila, column=col).fill = estilos.relleno_entrada


def _hoja_revision(libro: Workbook, estilos: Estilos, res: Resultado) -> None:
    """REVISION: lo que el motor no decide solo, con la lista desplegable y las columnas para llenar."""
    revision = res.tablas["REVISION"]
    hoja = hoja_desde_tabla(
        libro, estilos, "REVISION", revision, "revision", congelar="D2", anchos_max=60
    )
    if len(revision):
        colorear_columna(
            hoja,
            revision,
            "ESTADO_REVISION",
            [("PENDIENTE", estilos.relleno_advertencia), ("DECIDIDO", estilos.relleno_ok)],
        )
        columnas = list(revision.columns)
        col_decision = _validacion_de_decision(hoja, columnas, len(revision))
        _columnas_de_entrada(hoja, estilos, columnas, len(revision))
        # Hallazgo F3: aquí dice ARCHIVOS_DECISIONES y el paso 3 del LEEME dice ARCHIVO_DECISIONES.
        hoja.cell(row=1, column=col_decision).comment = Comment(
            "Llene esta columna (lista desplegable). Luego vuelva a correr el notebook "
            "con este Excel en ARCHIVOS_DECISIONES.",
            AUTOR_COMENTARIOS,
        )


def _hoja_decisiones(libro: Workbook, estilos: Estilos, res: Resultado) -> None:
    """DECISIONES: historial de todas las decisiones humanas aplicadas en la corrida."""
    decisiones = res.tablas["DECISIONES"]
    hoja = hoja_desde_tabla(
        libro, estilos, "DECISIONES", decisiones, "revision", anchos_max=70, envolver=False
    )
    if len(decisiones):
        hoja.cell(row=1, column=1).comment = Comment(
            "Historial de TODAS las decisiones aplicadas en esta corrida (de todos los archivos). "
            "Al usar este Excel como archivo de decisiones se leen esta hoja y, después, "
            "las decisiones nuevas de REVISION.",
            AUTOR_COMENTARIOS,
        )


# ── Hojas de datos ───────────────────────────────────────────────────────────
def _hoja_contactos(libro: Workbook, estilos: Estilos, res: Resultado) -> None:
    """CONTACTOS: web, teléfono, correo y dirección de cada registro vinculado."""
    hoja_desde_tabla(
        libro, estilos, "CONTACTOS", res.tablas["CONTACTOS"], "negocio", envolver=False
    )


def _hojas_detalle(libro: Workbook, estilos: Estilos, res: Resultado) -> None:
    """Una hoja por cada tabla ``DETALLE_*`` (la ancla y cada destino), en el orden de las tablas."""
    tablas = res.tablas
    for nombre in [k for k in tablas if k.startswith("DETALLE_")]:
        hoja_desde_tabla(
            libro,
            estilos,
            nombre,
            tablas[nombre],
            "detalle",
            congelar="C2",
            anchos_max=45,
            envolver=False,
        )


def _hoja_deduplicacion(libro: Workbook, estilos: Estilos, res: Resultado) -> None:
    """DEDUPLICACION: registros que son la misma entidad dentro de cada base."""
    hoja_desde_tabla(
        libro, estilos, "DEDUPLICACION", res.tablas["DEDUPLICACION"], "detalle", envolver=False
    )


# ── VALIDACION y hojas de auditoría ──────────────────────────────────────────
def _bloque_de_validacion(
    hoja: Worksheet, estilos: Estilos, titulo: str, tabla: pd.DataFrame | None, fila: int
) -> int:
    """Escribe un bloque de VALIDACION (título y tabla, o ``(sin datos…)``); devuelve la fila del siguiente."""
    hoja.cell(row=fila, column=1, value=titulo).font = estilos.titulo(11)
    fila += 1
    if tabla is None or not len(tabla):
        hoja.cell(
            row=fila, column=1, value="(sin datos: no se declaró referencia o no hay pares)"
        ).font = estilos.base
        return fila + 2
    escribir_tabla(hoja, tabla, fila0=fila)
    for col, nombre in enumerate(tabla.columns, 1):
        encabezado = hoja.cell(row=fila, column=col)
        encabezado.font, encabezado.fill = estilos.encabezado, estilos.relleno_encabezado
        encabezado.alignment = ALINEACION_ENVUELTA
        for i in range(fila + 1, fila + 1 + len(tabla)):
            celda = hoja.cell(row=i, column=col)
            celda.font = estilos.base
            if str(nombre) in _COLUMNAS_PORCENTUALES:
                celda.number_format = "0.0%"
    return fila + len(tabla) + 2


def _hoja_validacion(libro: Workbook, estilos: Estilos, res: Resultado) -> None:
    """VALIDACION: métricas del motor automático, discrepancias y pares decididos en varias revisiones."""
    tablas = res.tablas
    hoja = libro.create_sheet("VALIDACION")
    hoja.sheet_properties.tabColor = COLORES_PESTANA["auditoria"]
    hoja["A1"] = "Motor automático (sin decisiones) medido contra la referencia revisada"
    hoja["A1"].font = estilos.titulo(12)
    hoja["A2"] = (
        "Precisión = vínculos automáticos A+B confirmados por la referencia / vínculos automáticos "
        "del mismo país. Recall = vínculos de la referencia que el motor propone solo. IC95 de "
        "Wilson. La referencia no es una verdad de campo: su propio error no está medido."
    )
    hoja["A2"].font = estilos.cursiva(9)
    fila = 4
    for titulo, tabla in (
        ("Métricas", tablas["VALIDACION"]),
        ("Discrepancias motor vs referencia", tablas["VALIDACION_DETALLE"]),
        ("Pares decididos en más de una revisión", tablas["COMPARACION_REVISIONES"]),
    ):
        fila = _bloque_de_validacion(hoja, estilos, titulo, tabla, fila)
    for col, ancho in enumerate(_ANCHOS_DE_VALIDACION, 1):
        hoja.column_dimensions[get_column_letter(col)].width = ancho


def _hoja_calidad_datos(libro: Workbook, estilos: Estilos, res: Resultado) -> None:
    """CALIDAD_DATOS: problemas de los insumos que cambian la lectura de las cifras."""
    hoja_desde_tabla(
        libro, estilos, "CALIDAD_DATOS", res.tablas["CALIDAD_DATOS"], "auditoria", anchos_max=70
    )


def _hoja_control_qa(libro: Workbook, estilos: Estilos, res: Resultado) -> None:
    """CONTROL_QA: invariantes verificadas en la corrida, con semáforo en ESTADO."""
    control = res.tablas["CONTROL_QA"]
    hoja = hoja_desde_tabla(libro, estilos, "CONTROL_QA", control, "auditoria", anchos_max=70)
    colorear_columna(
        hoja,
        control,
        "ESTADO",
        [
            ("OK", estilos.relleno_ok),
            ("ADVERTENCIA", estilos.relleno_advertencia),
            ("FALLA", estilos.relleno_falla),
        ],
    )


def _hoja_parametros(libro: Workbook, estilos: Estilos, res: Resultado) -> None:
    """PARAMETROS: versiones, huellas de los insumos, umbrales y listas de la corrida."""
    hoja_desde_tabla(
        libro, estilos, "PARAMETROS", res.tablas["PARAMETROS"], "auditoria", anchos_max=120
    )


def _hoja_metodologia(
    libro: Workbook, estilos: Estilos, metodo: Sequence[tuple[str, str, str]]
) -> None:
    """METODOLOGIA: el método paso a paso."""
    tabla = pd.DataFrame(metodo, columns=["PASO", "TEMA", "DESCRIPCION"])
    hoja_desde_tabla(libro, estilos, "METODOLOGIA", tabla, "auditoria", anchos_max=120)


def escribir_hojas(libro: Workbook, estilos: Estilos, res: Resultado) -> None:
    """Crea todas las hojas del libro, en el orden de entrega.

    LEEME · RESUMEN · EXPORTADORES · ENLACES · REVISION · DECISIONES · CONTACTOS · DETALLE_* (uno por
    base, en el orden de las tablas) · DEDUPLICACION · VALIDACION · CALIDAD_DATOS · CONTROL_QA ·
    PARAMETROS · METODOLOGIA.
    """
    hojas, glosario, metodo = textos_leeme(res.cfg)
    hoja_leeme(libro, estilos, res.cfg, hojas, glosario)
    _hoja_resumen(libro, estilos, res)
    _hoja_exportadores(libro, estilos, res)
    _hoja_enlaces(libro, estilos, res)
    _hoja_revision(libro, estilos, res)
    _hoja_decisiones(libro, estilos, res)
    _hoja_contactos(libro, estilos, res)
    _hojas_detalle(libro, estilos, res)
    _hoja_deduplicacion(libro, estilos, res)
    _hoja_validacion(libro, estilos, res)
    _hoja_calidad_datos(libro, estilos, res)
    _hoja_control_qa(libro, estilos, res)
    _hoja_parametros(libro, estilos, res)
    _hoja_metodologia(libro, estilos, metodo)

"""Hoja LEEME del Excel de entrega: cómo leerlo, qué trae cada hoja, glosario y guía para llenar la revisión.

Los textos no dependen de los datos sino de la configuración (la ancla y los destinos): ``textos_leeme``
los arma y ``hoja_leeme`` los escribe. La hoja METODOLOGIA reutiliza el tercer texto (``metodo``).
"""

from __future__ import annotations

from collections.abc import Sequence
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

import record_linkage
from record_linkage.flujo.vinculacion.config import ConfigVinculacion, Fuente
from record_linkage.flujo.vinculacion.constantes import NOMBRE_MOTOR
from record_linkage.flujo.vinculacion_excel.estilos import (
    ALINEACION_ARRIBA,
    ALINEACION_ARRIBA_ENVUELTA,
    COLORES_PESTANA,
    Estilos,
    altura_de_fila,
)
from record_linkage.flujo.vinculacion_excel.tablas import escribir_tabla

__all__ = ["hoja_leeme", "prefijo_de_bloque", "textos_leeme"]

# Hallazgo F3: el paso 3 dice ARCHIVO_DECISIONES y el comentario de REVISION, ARCHIVOS_DECISIONES (el campo es plural).
_PASOS_DE_LECTURA = (
    "1. RESUMEN: cuántas empresas quedaron vinculadas, con qué confianza y qué tan bueno es el motor.",
    "2. EXPORTADORES: una fila por empresa; filtre RESULTADO = VINCULADO. Los datos de cada base vienen en columnas con su prefijo.",
    "3. REVISION: llene DECISION (lista desplegable) en las filas PENDIENTE y vuelva a correr el notebook con este archivo "
    "como ARCHIVO_DECISIONES. Nada de lo que está en C se vincula sin esa decisión.",
)

# Hallazgo F3: la definición de RELACIÓN DE PROPIEDAD trae un ejemplo con dos empresas reales (y otros textos
# nombran productos de datos reales). No se cambian: la tabla de textos declarados del arnés solo cubre el
# «p. ej.» del RESUMEN, no esta frase; cambiarla rompe la paridad.
_GLOSARIO = (
    (
        "NIVEL EMPRESA",
        "El registro destino es la MISMA persona jurídica o empresa: coinciden todas las palabras del nombre "
        "(salvo forma legal, puntuación, tildes, una errata o el truncamiento) y el país. Una decisión de "
        "revisión con evidencia (web, aviso legal, descripción) puede afirmarlo aunque el país del registro "
        "sea el de la matriz (fDi Markets).",
    ),
    (
        "NIVEL GRUPO",
        "El registro es OTRA sociedad del mismo grupo o marca (filial, matriz, hermana): comparten la marca y lo que "
        "sobra es genérico (país, sector, 'holding'), o tienen el mismo nombre en otro país. Úselo para contexto de "
        "grupo; no atribuya a la empresa las cifras financieras de otra entidad.",
    ),
    (
        "RELACIÓN DE PROPIEDAD",
        "Vínculo por dueño o matriz SIN nombre compartido (Kaufland frente a Lidl). Solo lo aporta la "
        "revisión y no cuenta como vínculo.",
    ),
    ("BANDA A", "Automático, alta confianza: mismo nombre normalizado y mismo país."),
    (
        "BANDA B",
        "Automático, confianza media: errata, truncamiento, alias entre paréntesis, nombre de una palabra, marca "
        "compartida con sobrantes genéricos o mismo nombre en otro país. Se vincula, pero conviene confirmarlo.",
    ),
    (
        "BANDA C",
        "Candidato: comparte una palabra distintiva pero difiere en otras, es una posible sigla o la matriz reportada "
        "coincide. NO se vincula hasta que alguien decida.",
    ),
    (
        "CONFIANZA ALTA / MEDIA",
        "ALTA = automático banda A, o decidido por una PERSONA. MEDIA = automático banda B sin "
        "revisar, o decidido por una revisión ASISTIDA (AUTOR con la marca '(verificar)') que nadie ha "
        "validado todavía. Para validar: ponga su nombre en AUTOR (sin la marca) en la hoja REVISION.",
    ),
    (
        "ESTADO_REVISION",
        "PENDIENTE = nadie ha decidido. VALIDAR = lo decidió una revisión asistida y falta la firma de una "
        "persona. DECIDIDO = decisión validada, o descarte de un candidato que no cambia el resultado automático.",
    ),
    (
        "SCORE_RUES_LINKER",
        "Similitud 0–1 del comparador SimilitudNombre de rues-linker (Jaccard ponderado por IDF + "
        "Jaro-Winkler). Es evidencia, no la regla: la decisión la toma la alineación palabra a palabra.",
    ),
)

_COLUMNAS_FIJAS_DE_EXPORTADORES = (
    (
        "ID_EXPORTADOR · NOMBRE_EXPORTADOR",
        "La empresa tras unir sus grafías. El nombre es la grafía más representativa "
        "(la que coincide con un registro oficial, o la más completa y frecuente); todas están en GRAFIAS.",
    ),
    ("RESULTADO", "VINCULADO · CANDIDATO POR REVISAR (solo candidatos C) · SIN COINCIDENCIA."),
    (
        "NIVEL_VINCULO · CONFIANZA",
        "El mejor nivel alcanzado (EMPRESA > GRUPO) y la confianza de sus vínculos principales.",
    ),
    (
        "REQUIERE_REVISION · QUE_REVISAR",
        "SÍ cuando hay vínculos de confianza MEDIA, candidatos C, duplicados internos por "
        "decidir o una unión de grafías hecha por revisión asistida. QUE_REVISAR dice qué.",
    ),
    (
        "Métricas",
        "Sumadas por empresa. Si una métrica viene repetida en el origen, la columna *_SIN_REPETIDOS es la "
        "recomendada y la suma simple queda al lado (ver CALIDAD_DATOS). Columnas constantes (p. ej. un valor en 0 para "
        "todos) se ocultan.",
    ),
)

_COLUMNA_DE_CONTACTO = (
    "CONTACTO_*",
    "Web, teléfono y correo de los registros vinculados a nivel EMPRESA (nunca de otras filiales).",
)

_GUIA_DE_REVISION = (
    (
        "DECISION",
        "VINCULO: EMPRESA · GRUPO · NO_VINCULAR · RELACION_PROPIEDAD. DUPLICADO: MISMA_ENTIDAD · DISTINTA.",
    ),
    (
        "AUTOR · FECHA",
        "Quién decidió y cuándo (texto libre; p. ej. 'M. Pérez', '2026-10-02'). Un AUTOR con la marca "
        "'(verificar)' es una revisión asistida: para VALIDARLA basta con reemplazar el AUTOR por su nombre.",
    ),
    (
        "RAZON",
        "La evidencia: 'misma web vantiko.es y misma ciudad', 'es la matriz según su web', 'homónimo en Málaga'.",
    ),
    (
        "Ejemplo",
        "TIPO=VINCULO · FUENTE=ORBIS · CLAVE_A=España|VANTIKO CHOCOLATES, S.A., · CLAVE_B=ESB0000001 · "
        "DECISION=EMPRESA · AUTOR=M. Pérez · FECHA=2026-10-02 · RAZON=Mismo nombre y web vantiko.es",
    ),
    (
        "No edite",
        "TIPO, FUENTE, CLAVE_A, CLAVE_B ni NOMBRE_B: son la llave con la que la decisión se vuelve a aplicar.",
    ),
)


def prefijo_de_bloque(destino: Fuente) -> str:
    """Prefijo de las columnas de EXPORTADORES que vienen de ``destino``: su primera palabra en mayúsculas."""
    return destino.etiqueta.split()[0].upper()


def _textos_de_hojas(cfg: ConfigVinculacion) -> list[tuple[str, str, str]]:
    """Una fila por hoja del libro: nombre, qué contiene y cómo usarla."""
    ancla, destinos = cfg.ancla, list(cfg.destinos)
    dest = " y ".join(d.etiqueta for d in destinos) or "(sin destinos)"
    # Hallazgo F3: la fila «VALIDACION*» no coincide con el nombre de la hoja (VALIDACION) y su asterisco no se explica.
    return [
        (
            "RESUMEN",
            "Cifras principales, cobertura por país, calidad medida del motor y controles.",
            "Empiece aquí.",
        ),
        (
            "EXPORTADORES",
            f"UNA FILA POR EMPRESA de {ancla.etiqueta} (todas, vinculadas o no): su resultado, nivel y confianza, "
            f"métricas consolidadas y el mejor registro de {dest} con sus datos.",
            "Filtre RESULTADO = VINCULADO. NIVEL_VINCULO dice si es la misma empresa o su grupo.",
        ),
        (
            "ENLACES",
            "Todos los vínculos finales, uno por par empresa–registro, con el porqué (MOTIVO) y la evidencia.",
            "ES_PRINCIPAL = SÍ es el que aparece en EXPORTADORES.",
        ),
        (
            "REVISION",
            "Lo que el motor no decide solo: vínculos B por confirmar, candidatos C y posibles duplicados. "
            "Las columnas amarillas (DECISION, AUTOR, FECHA, RAZON) son para llenar.",
            "Llene DECISION y vuelva a correr el notebook apuntando a este Excel: la decisión se aplica y queda trazada.",
        ),
        (
            "DECISIONES",
            "Historial de todas las decisiones humanas aplicadas (archivo, autor, fecha y razón).",
            "No se edita: se conserva para que la próxima corrida reproduzca este resultado.",
        ),
        (
            "CONTACTOS",
            "Web, teléfono, correo y dirección de cada registro vinculado, sin mezclar fuentes.",
            "Uso comercial.",
        ),
        (
            f"DETALLE_{ancla.clave}",
            f"Cada fila original de {ancla.etiqueta} con su ID_EXPORTADOR y resultado.",
            "Para cruzar con otras tablas.",
        ),
        *[
            (
                f"DETALLE_{d.clave}",
                f"Cada registro de {d.etiqueta} con los exportadores a los que quedó vinculado.",
                "",
            )
            for d in destinos
        ],
        (
            "DEDUPLICACION",
            "Registros que son la misma entidad dentro de cada base y cómo se unieron.",
            "",
        ),
        (
            "VALIDACION*",
            "Precisión y recall del motor automático contra una referencia revisada, y cada discrepancia.",
            "Mide qué tan confiable es el automático sin revisión.",
        ),
        (
            "CALIDAD_DATOS",
            "Problemas de los insumos que cambian la lectura (métricas en cero o repetidas, truncamientos…).",
            "Léala antes de usar las cifras.",
        ),
        (
            "CONTROL_QA",
            "Invariantes verificadas en cada corrida (conservación de filas y métricas, coherencia de vínculos).",
            "",
        ),
        (
            "PARAMETROS · METODOLOGIA",
            "Versiones, huellas SHA-256 de los insumos, umbrales, listas y el método paso a paso.",
            "Trazabilidad.",
        ),
    ]


def _textos_de_metodo(cfg: ConfigVinculacion) -> list[tuple[str, str, str]]:
    """Los once pasos del método: número, tema y descripción (hoja METODOLOGIA)."""
    ancla = cfg.ancla
    return [
        (
            "1",
            "Ingesta",
            "Lee cada hoja declarada, valida que existan las columnas, conserva la fila de Excel y asigna un ID por registro "
            "(el nativo si existe, p. ej. BvD ID).",
        ),
        (
            "2",
            "Normalización (rues-linker)",
            "sanear_texto → ascii_mayusculas → unir_iniciales → normalizar_nombre: tildes, "
            "mayúsculas, entidades HTML, iniciales sueltas y formas legales (GMBH, S.L.U., S.P.A., STIFTUNG & CO. KG…). "
            "El nombre entre paréntesis se guarda como alias y la matriz (Parent/GUO) como superficie aparte.",
        ),
        (
            "3",
            "País (rues-linker)",
            "canonizar_pais contra el catálogo ISO 3166 de la librería + alias declarados (FRANCE, ITALY…). "
            "Una grafía desconocida DETIENE la corrida con la sugerencia. El país es evidencia, no bloqueo.",
        ),
        (
            "4",
            "Deduplicación interna",
            f"{ancla.etiqueta}: mismo nombre normalizado, mismas letras con otros espacios o errata/truncamiento "
            "con forma legal compatible; lo dudoso (siglas, contención) va a REVISION. Con ID nativo (Orbis) "
            "solo por raíz del ID. Dos grafías que son la MISMA EMPRESA que un mismo registro destino se unen.",
        ),
        (
            "5",
            "Candidatos",
            "Comparación exhaustiva ancla × destino (recall del bloqueo = 1): pares que comparten una palabra "
            "distintiva (con tolerancia de una letra), cadenas con Jaro-Winkler ≥ 0,85 o siglas.",
        ),
        (
            "6",
            "Clasificación",
            "Alineación palabra a palabra + SimilitudNombre de rues-linker. EMPRESA: todas las palabras se "
            "corresponden, mismas iniciales/números, score ≥ piso y mismo país. GRUPO: marca compartida y sobrantes "
            "genéricos. C: lo demás con evidencia. Nombres de pila solos no cuentan como evidencia.",
        ),
        (
            "7",
            "Variables adicionales",
            "Si ancla y destino comparten identificador, web, correo, teléfono o ciudad, se usan como "
            "evidencia. Además, un registro con el MISMO SITIO WEB que un vínculo EMPRESA se propone o se sube "
            "de nivel (web + ciudad + palabra compartida = EMPRESA).",
        ),
        (
            "8",
            "Decisiones de revisión",
            "Se leen de CSV/Excel (TIPO, FUENTE, CLAVE_A, CLAVE_B, DECISION…) en orden; la última manda. "
            "Confirman, cambian el nivel, agregan o descartan, y cada vínculo dice su ORIGEN. Una decisión "
            "cuyo registro destino cambió de nombre DETIENE la corrida.",
        ),
        (
            "9",
            "Consolidación",
            "Una fila por empresa ancla: métricas sumadas (y sin repetidos cuando la base repite totales), mejor "
            "vínculo por destino (EMPRESA > GRUPO, confianza, mismo país, menos sobrantes, score) y contactos.",
        ),
        (
            "10",
            "Validación",
            "Precisión/recall del motor AUTOMÁTICO (sin decisiones) contra una referencia revisada, con IC95 de Wilson.",
        ),
        (
            "11",
            "Controles",
            "Invariantes de conservación y coherencia; huellas SHA-256 de insumos y resultado para saber si algo cambió.",
        ),
    ]


def textos_leeme(
    cfg: ConfigVinculacion,
) -> tuple[list[tuple[str, str, str]], list[tuple[str, str]], list[tuple[str, str, str]]]:
    """Textos de la hoja LEEME y de la hoja METODOLOGIA, en este orden.

    Devuelve ``(hojas, glosario, metodo)``: qué trae cada hoja del libro, el glosario de términos y los
    pasos del método. Dependen solo de la configuración: la ancla y los destinos.
    """
    return _textos_de_hojas(cfg), list(_GLOSARIO), _textos_de_metodo(cfg)


def _definiciones_de_exportadores(cfg: ConfigVinculacion) -> list[tuple[str, str]]:
    """Descripción de las columnas de EXPORTADORES: las fijas, un bloque por destino y los contactos."""
    definiciones = list(_COLUMNAS_FIJAS_DE_EXPORTADORES)
    for destino in cfg.destinos:
        pre = prefijo_de_bloque(destino)
        texto = (
            f"Registro principal de {destino.etiqueta} y sus datos. Si {pre}_NIVEL = GRUPO, los datos (ingresos, empleados, "
            f"proyectos, contactos) son de OTRA sociedad del grupo: sirven como contexto, no para atribuirlos al exportador. "
            f"{pre}_N_VINCULOS cuenta todos sus vínculos en esa base; {pre}_OTROS_VINCULOS los lista."
        )
        if destino.cols_pais_extra:
            texto += (
                f" {pre}_PAIS es '{destino.col_pais}' tal como viene; también cuenta como mismo país "
                f"{', '.join(destino.cols_pais_extra)}."
            )
        definiciones.append((f"{pre}_…", texto))
    definiciones.append(_COLUMNA_DE_CONTACTO)
    return definiciones


def _titulo_y_subtitulo(hoja: Worksheet, estilos: Estilos, cfg: ConfigVinculacion) -> None:
    hoja["A1"] = cfg.titulo
    hoja["A1"].font = estilos.titulo(14)
    hoja["A2"] = (
        f"Vinculación de {cfg.ancla.etiqueta} con "
        + " y ".join(d.etiqueta for d in cfg.destinos)
        + f" · motor {NOMBRE_MOTOR} sobre rues-linker {record_linkage.__version__} · "
        + datetime.now().strftime("%Y-%m-%d %H:%M")
    )
    hoja["A2"].font = estilos.cursiva(10)


def _pasos_de_lectura(hoja: Worksheet, estilos: Estilos) -> int:
    """Escribe «Cómo leer este archivo» (fila 4 y los tres pasos); devuelve la fila del título siguiente."""
    fila = 4
    hoja.cell(
        row=fila, column=1, value="Cómo leer este archivo en tres pasos"
    ).font = estilos.titulo(12)
    for paso in _PASOS_DE_LECTURA:
        fila += 1
        hoja.cell(row=fila, column=1, value=paso).font = estilos.base
    return fila + 2


def _tabla_de_hojas(
    hoja: Worksheet, estilos: Estilos, hojas: Sequence[tuple[str, str, str]], fila: int
) -> int:
    """Escribe «Hojas» (qué trae cada hoja) desde ``fila``; devuelve la fila del título siguiente."""
    hoja.cell(row=fila, column=1, value="Hojas").font = estilos.titulo(12)
    fila += 1
    tabla = pd.DataFrame(hojas, columns=["HOJA", "QUE_CONTIENE", "COMO_USARLA"])
    escribir_tabla(hoja, tabla, fila0=fila)
    for j in range(1, 4):
        celda = hoja.cell(row=fila, column=j)
        celda.font, celda.fill = estilos.encabezado, estilos.relleno_encabezado
    for i in range(fila + 1, fila + 1 + len(tabla)):
        for j in range(1, 4):
            hoja.cell(row=i, column=j).font = estilos.negrita if j == 1 else estilos.base
            hoja.cell(row=i, column=j).alignment = ALINEACION_ARRIBA_ENVUELTA
        largo = max(
            len(str(hoja.cell(row=i, column=2).value or "")) / 95,
            len(str(hoja.cell(row=i, column=3).value or "")) / 60,
        )
        hoja.row_dimensions[i].height = altura_de_fila(int(largo))
    return fila + len(tabla) + 2


def _escribir_definicion(
    hoja: Worksheet,
    estilos: Estilos,
    fila: int,
    termino: str,
    definicion: str,
    *,
    envolver_termino: bool = False,
) -> None:
    """Escribe un término en negrita (columna A) y su definición combinada en B:C, con la altura que pide el texto."""
    hoja.cell(row=fila, column=1, value=termino).font = estilos.negrita
    hoja.cell(row=fila, column=1).alignment = (
        ALINEACION_ARRIBA_ENVUELTA if envolver_termino else ALINEACION_ARRIBA
    )
    hoja.cell(row=fila, column=2, value=definicion).font = estilos.base
    hoja.cell(row=fila, column=2).alignment = ALINEACION_ARRIBA_ENVUELTA
    hoja.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=3)
    hoja.row_dimensions[fila].height = altura_de_fila(len(definicion) // 150)


def _glosario_de_terminos(
    hoja: Worksheet, estilos: Estilos, glosario: Sequence[tuple[str, str]], fila: int
) -> int:
    """Escribe «Glosario» desde ``fila``; devuelve la fila del título siguiente."""
    hoja.cell(row=fila, column=1, value="Glosario").font = estilos.titulo(12)
    fila += 1
    for termino, definicion in glosario:
        _escribir_definicion(hoja, estilos, fila, termino, definicion)
        fila += 1
    return fila + 1


def _columnas_de_exportadores(
    hoja: Worksheet, estilos: Estilos, cfg: ConfigVinculacion, fila: int
) -> int:
    """Escribe «Columnas de EXPORTADORES» desde ``fila``; devuelve la fila del título siguiente."""
    hoja.cell(row=fila, column=1, value="Columnas de EXPORTADORES").font = estilos.titulo(12)
    for termino, definicion in _definiciones_de_exportadores(cfg):
        fila += 1
        _escribir_definicion(hoja, estilos, fila, termino, definicion, envolver_termino=True)
    return fila + 2


def _guia_para_llenar_la_revision(hoja: Worksheet, estilos: Estilos, fila: int) -> None:
    """Escribe «Cómo llenar la revisión» desde ``fila``."""
    hoja.cell(row=fila, column=1, value="Cómo llenar la revisión").font = estilos.titulo(12)
    for termino, definicion in _GUIA_DE_REVISION:
        fila += 1
        _escribir_definicion(hoja, estilos, fila, termino, definicion)


def hoja_leeme(
    libro: Workbook,
    estilos: Estilos,
    cfg: ConfigVinculacion,
    hojas: Sequence[tuple[str, str, str]],
    glosario: Sequence[tuple[str, str]],
) -> Worksheet:
    """Crea la hoja LEEME: título, tres pasos, hojas, glosario, columnas de EXPORTADORES y guía de revisión."""
    hoja = libro.create_sheet("LEEME")
    hoja.sheet_properties.tabColor = COLORES_PESTANA["negocio"]
    _titulo_y_subtitulo(hoja, estilos, cfg)
    fila = _pasos_de_lectura(hoja, estilos)
    fila = _tabla_de_hojas(hoja, estilos, hojas, fila)
    fila = _glosario_de_terminos(hoja, estilos, glosario, fila)
    fila = _columnas_de_exportadores(hoja, estilos, cfg, fila)
    _guia_para_llenar_la_revision(hoja, estilos, fila)
    hoja.column_dimensions["A"].width = 26
    hoja.column_dimensions["B"].width = 95
    hoja.column_dimensions["C"].width = 60
    return hoja

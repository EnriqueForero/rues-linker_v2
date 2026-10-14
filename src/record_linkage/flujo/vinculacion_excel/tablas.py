"""Volcado de un ``DataFrame`` en una hoja: celdas, formato de columnas, filtro y colores condicionales."""

from __future__ import annotations

import re
from collections.abc import Sequence

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from record_linkage.flujo.vinculacion_excel.estilos import (
    ALINEACION_ENCABEZADO,
    COLORES_PESTANA,
    Estilos,
    Pestana,
)

__all__ = [
    "colorear_columna",
    "escribir_celda",
    "escribir_tabla",
    "formatear_tabla",
    "hoja_desde_tabla",
]

#: Caracteres de control que Excel no admite (se conservan el tabulador, el salto de línea y el retorno).
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
#: Excel limita el nombre de una hoja a 31 caracteres.
_LARGO_MAXIMO_NOMBRE_HOJA = 31


def _valor_de_celda(valor: object) -> object:
    """Lo que se escribe en la celda: nativos de Python, ``SÍ``/``NO`` y texto sin caracteres de control."""
    if isinstance(valor, np.integer):
        return int(valor)
    if isinstance(valor, np.floating):
        return None if np.isnan(valor) else float(valor)
    if isinstance(valor, (np.bool_, bool)):
        return "SÍ" if bool(valor) else "NO"
    if isinstance(valor, str):
        return _CONTROL.sub(" ", valor)
    return valor


def escribir_celda(hoja: Worksheet, fila: int, columna: int, valor: object) -> Cell:
    """Escribe ``valor`` en una celda con las mismas garantías que ``escribir_tabla``.

    Nativos de Python, ausente = celda vacía, booleano = ``SÍ``/``NO``, caracteres de control cambiados por un
    espacio y un texto que empieza con ``=`` marcado como texto (no es una fórmula).
    """
    valor = _valor_de_celda(valor)
    celda = hoja.cell(row=fila, column=columna, value=valor)
    if isinstance(valor, str) and valor.startswith("="):
        celda.data_type = "s"  # texto que empieza con '=' no es una fórmula
    return celda


def escribir_tabla(hoja: Worksheet, tabla: pd.DataFrame, fila0: int = 1, col0: int = 1) -> None:
    """Escribe ``tabla`` desde la celda (``fila0``, ``col0``): el encabezado y luego una fila por registro.

    Encabezado y datos pasan por ``escribir_celda``: los tipos de numpy pasan a nativos, un ausente queda como
    celda vacía, un booleano se escribe ``SÍ``/``NO``, los caracteres de control se cambian por un espacio y un
    texto que empieza con ``=`` se marca como texto (no es una fórmula).
    """
    for j, c in enumerate(tabla.columns):
        escribir_celda(hoja, fila0, col0 + j, str(c))
    valores = tabla.astype(object).where(tabla.notna(), None).to_numpy()
    for i in range(len(tabla)):
        for j in range(len(tabla.columns)):
            escribir_celda(hoja, fila0 + 1 + i, col0 + j, valores[i, j])


def _largos_de_columnas(tabla: pd.DataFrame) -> list[int]:
    """Largo de cada columna: el mayor entre su nombre y el percentil 90 de sus primeros 300 valores."""
    largos = [len(str(c)) for c in tabla.columns]
    muestra = tabla.head(300)
    for j, c in enumerate(tabla.columns):
        # Los ausentes no cuentan: con pandas 3 ``astype(str)`` los deja en NaN (y una columna sin ningún valor
        # lanzaba ValueError) y con pandas 2 los vuelve 'None'/'nan'; así el ancho no depende de la versión.
        presentes = muestra[c].dropna().astype(str)
        if len(presentes):
            largos[j] = max(largos[j], int(presentes.str.len().quantile(0.9)))
    return largos


def _formato_numerico(nombre: str, serie: pd.Series) -> str | None:
    """Formato de número de una columna numérica según su nombre (``None`` si no es numérica)."""
    es_numerica = pd.api.types.is_numeric_dtype(serie) and not pd.api.types.is_bool_dtype(serie)
    if not es_numerica:
        return None
    if (
        nombre.startswith("PCT")
        or nombre in ("PRECISION", "RECALL", "F1")
        or nombre.startswith("RECALL_")
    ):
        return "0.0%"
    if any(k in nombre for k in ("SCORE", "SIM_", "JARO")):
        return "0.00"
    if "ANIO" in nombre or nombre.startswith("FILA") or nombre.endswith("_ID"):
        return "0"
    if serie.dropna().apply(lambda v: float(v).is_integer()).all():
        return "#,##0"
    return "#,##0.00"


def formatear_tabla(
    hoja: Worksheet,
    tabla: pd.DataFrame,
    estilos: Estilos,
    *,
    fila0: int = 1,
    anchos_max: int = 55,
    congelar: str | None = "A2",
    filtro: bool = True,
    envolver: bool = True,
) -> None:
    """Da formato a una tabla ya escrita desde ``fila0``: encabezado, anchos, números, panel y filtro.

    ``anchos_max`` topa el ancho de las columnas; ``envolver`` parte en líneas el texto de las que
    llegan al tope; ``congelar`` es la celda donde se inmovilizan los paneles (``None`` = no se
    inmoviliza); ``filtro`` pone el autofiltro sobre el encabezado.
    """
    ncol = len(tabla.columns)
    for j in range(1, ncol + 1):
        celda = hoja.cell(row=fila0, column=j)
        celda.font, celda.fill = estilos.encabezado, estilos.relleno_encabezado
        celda.alignment = ALINEACION_ENCABEZADO
    hoja.row_dimensions[fila0].height = 30
    largos = _largos_de_columnas(tabla)
    for j, c in enumerate(tabla.columns, 1):
        ancho = min(max(8, largos[j - 1] + 2), anchos_max)
        hoja.column_dimensions[get_column_letter(j)].width = ancho
        formato = _formato_numerico(str(c).upper(), tabla[c])
        alineacion = Alignment(wrap_text=envolver and ancho >= anchos_max, vertical="top")
        for i in range(fila0 + 1, fila0 + 1 + len(tabla)):
            celda = hoja.cell(row=i, column=j)
            celda.font = estilos.base
            if formato:
                celda.number_format = formato
            celda.alignment = alineacion
    if congelar:
        hoja.freeze_panes = congelar
    if filtro and len(tabla.columns):
        hoja.auto_filter.ref = f"A{fila0}:{get_column_letter(ncol)}{fila0 + max(1, len(tabla))}"


def colorear_columna(
    hoja: Worksheet,
    tabla: pd.DataFrame,
    columna: str,
    reglas: Sequence[tuple[str, PatternFill]],
    *,
    fila0: int = 1,
) -> None:
    """Pinta las celdas de ``columna`` cuyo texto EMPIEZA con cada ``texto`` de ``reglas`` (formato condicional).

    No hace nada si la tabla no tiene esa columna o no tiene filas.
    """
    if columna not in tabla.columns or not len(tabla):
        return
    letra = get_column_letter(list(tabla.columns).index(columna) + 1)
    rango = f"{letra}{fila0 + 1}:{letra}{fila0 + len(tabla)}"
    for texto, relleno in reglas:
        regla = FormulaRule(
            formula=[f'LEFT(${letra}{fila0 + 1},{len(texto)})="{texto}"'], fill=relleno
        )
        hoja.conditional_formatting.add(rango, regla)


def _nombre_de_hoja_unico(libro: Workbook, nombre: str) -> str:
    """``nombre`` recortado a 31 caracteres y distinto (sin distinguir mayúsculas, como Excel) de los ya usados.

    Si el recortado ya existe, se le pone ``_2``, ``_3``… al final, dentro de los 31 caracteres. Antes openpyxl
    renombraba la segunda hoja a 32 caracteres (con un aviso) y Excel podía no abrir el libro.
    """
    base = nombre[:_LARGO_MAXIMO_NOMBRE_HOJA]
    usados = {n.lower() for n in libro.sheetnames}
    candidato, k = base, 1
    while candidato.lower() in usados:
        k += 1
        sufijo = f"_{k}"
        candidato = base[: _LARGO_MAXIMO_NOMBRE_HOJA - len(sufijo)] + sufijo
    return candidato


def hoja_desde_tabla(
    libro: Workbook,
    estilos: Estilos,
    nombre: str,
    tabla: pd.DataFrame | None,
    pestana: Pestana,
    *,
    congelar: str | None = "A2",
    anchos_max: int = 55,
    envolver: bool = True,
) -> Worksheet:
    """Crea la hoja ``nombre`` (recortada a 31 caracteres y sin repetir el de otra hoja) con ``tabla`` escrita y formateada.

    Una tabla ausente o sin columnas deja la hoja con la sola celda ``(sin datos)``.
    """
    hoja = libro.create_sheet(_nombre_de_hoja_unico(libro, nombre))
    hoja.sheet_properties.tabColor = COLORES_PESTANA[pestana]
    if tabla is None or not len(tabla.columns):
        hoja["A1"] = "(sin datos)"
        hoja["A1"].font = estilos.base
        return hoja
    escribir_tabla(hoja, tabla)
    formatear_tabla(
        hoja, tabla, estilos, anchos_max=anchos_max, congelar=congelar, envolver=envolver
    )
    return hoja

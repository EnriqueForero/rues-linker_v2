"""Estilos del Excel de entrega: fuente, colores, rellenos y alineaciones.

Todo lo que decide CÓMO SE VE una hoja vive aquí. El contenido de cada hoja está en ``hojas.py`` y
``leeme.py``; el volcado de una tabla (anchos, formatos de número, filtros), en ``tablas.py``.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Literal

from openpyxl.styles import Alignment, Font, PatternFill

__all__ = [
    "ALINEACION_ARRIBA",
    "ALINEACION_ARRIBA_ENVUELTA",
    "ALINEACION_DERECHA",
    "ALINEACION_ENCABEZADO",
    "ALINEACION_ENVUELTA",
    "AUTOR_COMENTARIOS",
    "COLORES_PESTANA",
    "Estilos",
    "Pestana",
    "altura_de_fila",
]

#: Fuente de todo el libro.
_FUENTE = "Arial"
#: Azul de los encabezados y de los títulos de sección.
_AZUL = "1F3864"
_BLANCO = "FFFFFF"

#: Autor de los comentarios que el motor deja en los encabezados.
AUTOR_COMENTARIOS = "motor"

#: Familias de hojas: cada una lleva su color de pestaña.
Pestana = Literal["negocio", "revision", "detalle", "auditoria"]
COLORES_PESTANA: dict[Pestana, str] = {
    "negocio": "548235",
    "revision": "C65911",
    "detalle": "2F5597",
    "auditoria": "7F7F7F",
}

ALINEACION_ENCABEZADO = Alignment(wrap_text=True, vertical="center")
ALINEACION_ARRIBA = Alignment(vertical="top")
ALINEACION_ARRIBA_ENVUELTA = Alignment(wrap_text=True, vertical="top")
ALINEACION_ENVUELTA = Alignment(wrap_text=True)
ALINEACION_DERECHA = Alignment(horizontal="right")


@dataclass(frozen=True)
class Estilos:
    """Fuentes y rellenos del libro. ``Estilos()`` trae los del Excel de entrega.

    Los semáforos del libro son tres rellenos: ``relleno_ok`` (verde), ``relleno_advertencia`` (amarillo)
    y ``relleno_falla`` (naranja). ``relleno_entrada`` marca las celdas que una persona debe llenar.
    """

    base: Font = field(default_factory=lambda: Font(name=_FUENTE, size=10))
    negrita: Font = field(default_factory=lambda: Font(name=_FUENTE, size=10, bold=True))
    encabezado: Font = field(
        default_factory=lambda: Font(name=_FUENTE, size=10, bold=True, color=_BLANCO)
    )
    enlace: Font = field(
        default_factory=lambda: Font(name=_FUENTE, size=10, color="0563C1", underline="single")
    )
    relleno_encabezado: PatternFill = field(
        default_factory=lambda: PatternFill("solid", fgColor=_AZUL)
    )
    relleno_encabezado_entrada: PatternFill = field(
        default_factory=lambda: PatternFill("solid", fgColor="BF8F00")
    )
    relleno_entrada: PatternFill = field(
        default_factory=lambda: PatternFill("solid", fgColor="FFF2CC")
    )
    relleno_ok: PatternFill = field(default_factory=lambda: PatternFill("solid", fgColor="E2EFDA"))
    relleno_advertencia: PatternFill = field(
        default_factory=lambda: PatternFill("solid", fgColor="FFF2CC")
    )
    relleno_falla: PatternFill = field(
        default_factory=lambda: PatternFill("solid", fgColor="F8CBAD")
    )

    def titulo(self, tamano: int) -> Font:
        """Fuente de un título de hoja o de sección: negrita y azul, del tamaño pedido."""
        return Font(name=_FUENTE, size=tamano, bold=True, color=_AZUL)

    def cursiva(self, tamano: int) -> Font:
        """Fuente de una nota al pie o de un subtítulo: cursiva, del tamaño pedido."""
        return Font(name=_FUENTE, size=tamano, italic=True)


def altura_de_fila(lineas_completas: int) -> int:
    """Alto en puntos de una fila de texto envuelto.

    ``lineas_completas`` es cuántas veces cabe el ancho de la columna en el texto (la división entera
    del largo del texto por los caracteres que caben en una línea): 14 puntos por línea, contando la
    primera, más 2 de holgura.
    """
    return 14 * (lineas_completas + 1) + 2

"""Lectores locales seguros y de memoria acotada para fuentes heterogéneas."""

from __future__ import annotations

import codecs
import csv
import gzip
import importlib.util
import io
import math
import re
import shutil
import stat
import tempfile
import unicodedata
import zipfile
from collections.abc import Iterator, Sequence
from contextlib import contextmanager
from dataclasses import dataclass, field
from itertools import islice
from pathlib import Path, PurePosixPath
from types import MappingProxyType
from typing import Literal, Protocol, cast

import numpy as np
import pandas as pd

from ..matching.normalizadores import normalizar_identificador, normalizar_numero
from .errors import ArchiveSafetyError, EncodingDetectionError, IngestionError, SchemaError
from .specs import (
    ColumnType,
    Compression,
    IdentifierFormat,
    InputFormat,
    InvalidValuePolicy,
    LoadedSource,
    NumericFormat,
    SourceLoadReport,
    SourceSpec,
)

_TEXT_EXTENSIONS = {".csv", ".txt", ".tsv"}
_PARQUET_EXTENSIONS = {".parquet", ".pq"}
_XLSX_EXTENSIONS = {".xlsx", ".xlsm"}
_SUPPORTED_EXTENSIONS = _TEXT_EXTENSIONS | _PARQUET_EXTENSIONS | _XLSX_EXTENSIONS
_DELIMITER_CANDIDATES = (",", "\t", ";", "|")
_MAX_ENCODING_SCAN_CHUNK = 4 * 1024**2
_MAX_DELIMITER_SAMPLE_CHARS = 256 * 1024
_FLOAT_EXACT_INTEGER_LIMIT = 2**53


class _ReadableBytes(Protocol):
    def read(self, size: int = -1) -> bytes: ...


class _WritableBytes(Protocol):
    def write(self, data: bytes) -> int: ...


@dataclass(frozen=True)
class _MaterializedInput:
    path: Path
    logical_name: str
    compression: Compression
    archive_member: str | None = None


@dataclass
class _ReportState:
    input_format: InputFormat = InputFormat.AUTO
    compression: Compression = Compression.NONE
    encoding: str | None = None
    delimiter: str | None = None
    engine: str | None = None
    archive_member: str | None = None
    rows: int = 0
    columns: tuple[str, ...] = ()
    resolved_mapping: dict[str, str] = field(default_factory=dict)
    invalid_values: dict[str, int] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    missing_optional: tuple[str, ...] = ()


@dataclass(frozen=True)
class _ColumnPlan:
    selected: tuple[str, ...] | None
    rename: MappingProxyType
    output_order: tuple[str, ...]
    resolved_mapping: MappingProxyType
    missing_optional: tuple[str, ...]


def _preferred_string_dtype() -> Literal["string[pyarrow]", "string"]:
    """Usa Arrow cuando está disponible; evita objetos Python en tablas grandes."""
    return "string[pyarrow]" if importlib.util.find_spec("pyarrow") is not None else "string"


def _is_ooxml_workbook(path: Path) -> bool:
    if not zipfile.is_zipfile(path):
        return False
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            return "[Content_Types].xml" in names and "xl/workbook.xml" in names
    except (OSError, zipfile.BadZipFile):
        return False


def _detect_compression(path: Path, requested: Compression) -> Compression:
    if requested is not Compression.AUTO:
        return requested
    suffix = path.suffix.lower()
    if suffix in _XLSX_EXTENSIONS:
        return Compression.NONE
    if suffix == ".zip":
        return Compression.ZIP
    if suffix in {".gz", ".gzip"}:
        return Compression.GZIP

    with path.open("rb") as stream:
        magic = stream.read(4)
    if magic.startswith(b"\x1f\x8b"):
        return Compression.GZIP
    if magic.startswith(b"PK\x03\x04"):
        return Compression.NONE if _is_ooxml_workbook(path) else Compression.ZIP
    return Compression.NONE


def _safe_archive_name(name: str) -> PurePosixPath:
    if not name or "\\" in name or "\x00" in name:
        raise ArchiveSafetyError(f"Nombre inseguro dentro del ZIP: {name!r}.")
    posix = PurePosixPath(name)
    if posix.is_absolute() or any(part in {"", ".", ".."} for part in posix.parts):
        raise ArchiveSafetyError(f"Ruta insegura dentro del ZIP: {name!r}.")
    return posix


def _validate_zip_infos(
    infos: Sequence[zipfile.ZipInfo], spec: SourceSpec
) -> list[zipfile.ZipInfo]:
    limits = spec.safety_limits
    if len(infos) > limits.max_members:
        raise ArchiveSafetyError(f"ZIP con {len(infos):,} entradas; límite={limits.max_members:,}.")
    files: list[zipfile.ZipInfo] = []
    seen: set[str] = set()
    total = 0
    for info in infos:
        safe_name = _safe_archive_name(info.filename)
        normalized = safe_name.as_posix()
        if normalized in seen:
            raise ArchiveSafetyError(f"ZIP contiene nombre duplicado: {normalized!r}.")
        seen.add(normalized)
        mode = (info.external_attr >> 16) & 0xFFFF
        if stat.S_ISLNK(mode):
            raise ArchiveSafetyError(f"ZIP contiene enlace simbólico: {normalized!r}.")
        if info.flag_bits & 0x1:
            raise ArchiveSafetyError(f"ZIP contiene entrada cifrada: {normalized!r}.")
        if info.is_dir():
            continue
        if info.file_size > limits.max_member_bytes:
            raise ArchiveSafetyError(
                f"Miembro {normalized!r} mide {info.file_size:,} bytes; "
                f"límite={limits.max_member_bytes:,}."
            )
        total += info.file_size
        if total > limits.max_total_uncompressed_bytes:
            raise ArchiveSafetyError(
                f"ZIP declara {total:,} bytes sin comprimir; "
                f"límite={limits.max_total_uncompressed_bytes:,}."
            )
        ratio = info.file_size / max(1, info.compress_size)
        if ratio > limits.max_compression_ratio:
            raise ArchiveSafetyError(
                f"Miembro {normalized!r} tiene ratio {ratio:.1f}x; "
                f"límite={limits.max_compression_ratio:.1f}x."
            )
        files.append(info)
    if not files:
        raise ArchiveSafetyError("El ZIP no contiene archivos regulares.")
    return files


def _select_zip_member(files: Sequence[zipfile.ZipInfo], spec: SourceSpec) -> zipfile.ZipInfo:
    if spec.archive_member is not None:
        matches = [info for info in files if info.filename == spec.archive_member]
        if not matches:
            available = ", ".join(repr(info.filename) for info in files[:10])
            raise SchemaError(
                f"archive_member={spec.archive_member!r} no existe. Disponibles: {available}."
            )
        return matches[0]

    visible = [
        info
        for info in files
        if "__MACOSX" not in PurePosixPath(info.filename).parts
        and not PurePosixPath(info.filename).name.startswith(".")
    ]
    supported = [
        info for info in visible if Path(info.filename).suffix.lower() in _SUPPORTED_EXTENSIONS
    ]
    if len(supported) == 1:
        return supported[0]
    if not supported and len(visible) == 1:
        return visible[0]
    candidates = supported or visible
    names = ", ".join(repr(info.filename) for info in candidates[:10])
    raise SchemaError(
        f"El ZIP contiene múltiples archivos de datos; declare archive_member. Candidatos: {names}."
    )


def _copy_bounded(
    source: _ReadableBytes,
    target: _WritableBytes,
    *,
    max_bytes: int,
    chunk_bytes: int,
) -> int:
    total = 0
    while True:
        chunk = source.read(chunk_bytes)
        if not chunk:
            break
        total += len(chunk)
        if total > max_bytes:
            raise ArchiveSafetyError(f"La descompresión superó el límite de {max_bytes:,} bytes.")
        target.write(chunk)
    return total


@contextmanager
def _materialize_input(spec: SourceSpec) -> Iterator[_MaterializedInput]:
    path = Path(spec.path)
    if not path.exists():
        raise FileNotFoundError(f"Fuente '{spec.name}' no existe: {path}.")
    if not path.is_file():
        raise IngestionError(f"Fuente '{spec.name}' no es un archivo regular: {path}.")
    compression = _detect_compression(path, Compression(spec.compression))
    if compression is Compression.NONE:
        yield _MaterializedInput(path=path, logical_name=path.name, compression=compression)
        return

    temp_parent = Path(spec.temp_dir) if spec.temp_dir is not None else None
    if temp_parent is not None:
        temp_parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="rues-linker-", dir=temp_parent) as temp_name:
        materialized = Path(temp_name) / "payload"
        if compression is Compression.ZIP:
            try:
                with zipfile.ZipFile(path) as archive:
                    files = _validate_zip_infos(archive.infolist(), spec)
                    selected = _select_zip_member(files, spec)
                    free_bytes = shutil.disk_usage(materialized.parent).free
                    if selected.file_size > free_bytes:
                        raise ArchiveSafetyError(
                            f"No hay disco temporal suficiente: se requieren "
                            f"{selected.file_size:,} bytes y hay {free_bytes:,}."
                        )
                    with archive.open(selected, "r") as source, materialized.open("wb") as target:
                        copied = _copy_bounded(
                            source,
                            target,
                            max_bytes=min(
                                spec.safety_limits.max_member_bytes,
                                spec.safety_limits.max_total_uncompressed_bytes,
                            ),
                            chunk_bytes=spec.safety_limits.copy_chunk_bytes,
                        )
                    if copied != selected.file_size:
                        raise ArchiveSafetyError(
                            f"Tamaño descomprimido inesperado para {selected.filename!r}: "
                            f"declarado={selected.file_size:,}, leído={copied:,}."
                        )
            except ArchiveSafetyError:
                raise
            except (OSError, EOFError, RuntimeError, zipfile.BadZipFile) as exc:
                raise ArchiveSafetyError(f"ZIP inválido o con CRC incorrecto: {path}.") from exc
            yield _MaterializedInput(
                path=materialized,
                logical_name=selected.filename,
                compression=compression,
                archive_member=selected.filename,
            )
            return

        compressed_bytes = max(1, path.stat().st_size)
        gzip_limit = min(
            spec.safety_limits.max_member_bytes,
            spec.safety_limits.max_total_uncompressed_bytes,
            int(compressed_bytes * spec.safety_limits.max_compression_ratio),
        )
        try:
            with gzip.open(path, "rb") as source, materialized.open("wb") as target:
                copied = _copy_bounded(
                    source,
                    target,
                    max_bytes=gzip_limit,
                    chunk_bytes=spec.safety_limits.copy_chunk_bytes,
                )
        except ArchiveSafetyError:
            raise
        except (OSError, EOFError) as exc:
            raise ArchiveSafetyError(f"GZIP inválido o truncado: {path}.") from exc
        ratio = copied / compressed_bytes
        if ratio > spec.safety_limits.max_compression_ratio:
            raise ArchiveSafetyError(
                f"GZIP tiene ratio {ratio:.1f}x; "
                f"límite={spec.safety_limits.max_compression_ratio:.1f}x."
            )
        logical_name = path.name
        for suffix in (".gzip", ".gz"):
            if logical_name.lower().endswith(suffix):
                logical_name = logical_name[: -len(suffix)]
                break
        yield _MaterializedInput(
            path=materialized,
            logical_name=logical_name,
            compression=compression,
        )


def _detect_input_format(path: Path, logical_name: str, requested: InputFormat) -> InputFormat:
    if requested is not InputFormat.AUTO:
        return requested
    suffix = Path(logical_name).suffix.lower()
    if suffix == ".csv":
        return InputFormat.CSV
    if suffix in {".txt", ".tsv"}:
        return InputFormat.TXT
    if suffix in _XLSX_EXTENSIONS:
        return InputFormat.XLSX
    if suffix in _PARQUET_EXTENSIONS:
        return InputFormat.PARQUET
    with path.open("rb") as stream:
        magic = stream.read(4)
    if magic == b"PAR1":
        return InputFormat.PARQUET
    if magic.startswith(b"PK\x03\x04") and _is_ooxml_workbook(path):
        return InputFormat.XLSX
    if b"\x00" in magic:
        raise IngestionError(
            f"No se pudo inferir el formato binario de {logical_name!r}; declare format=."
        )
    return InputFormat.TXT


def _encoding_is_valid(path: Path, encoding: str) -> bool:
    decoder = codecs.getincrementaldecoder(encoding)(errors="strict")
    try:
        with path.open("rb") as stream:
            while chunk := stream.read(_MAX_ENCODING_SCAN_CHUNK):
                decoder.decode(chunk, final=False)
        decoder.decode(b"", final=True)
    except (LookupError, UnicodeDecodeError):
        return False
    return True


def _detect_encoding(path: Path, requested: str) -> str:
    requested_clean = str(requested).strip()
    if not requested_clean:
        raise ValueError("encoding no puede ser vacío.")
    with path.open("rb") as stream:
        bom = stream.read(4)
    if requested_clean.casefold() != "auto":
        try:
            codecs.lookup(requested_clean)
        except LookupError as exc:
            raise EncodingDetectionError(f"Codificación desconocida: {requested_clean!r}.") from exc
        if not _encoding_is_valid(path, requested_clean):
            raise EncodingDetectionError(
                f"{path.name!r} contiene bytes inválidos para {requested_clean!r}."
            )
        return requested_clean

    if bom.startswith(codecs.BOM_UTF8):
        candidates: tuple[str, ...] = ("utf-8-sig",)
    elif bom.startswith(codecs.BOM_UTF32_LE) or bom.startswith(codecs.BOM_UTF32_BE):
        candidates = ("utf-32",)
    elif bom.startswith(codecs.BOM_UTF16_LE) or bom.startswith(codecs.BOM_UTF16_BE):
        candidates = ("utf-16",)
    else:
        # CP1252 se prueba explícitamente; latin-1 nunca es fallback silencioso.
        candidates = ("utf-8", "cp1252")
    for encoding in candidates:
        if _encoding_is_valid(path, encoding):
            return encoding
    raise EncodingDetectionError(
        f"No se pudo decodificar {path.name!r} sin pérdida como UTF-8 o CP1252. "
        "Declare encoding= explícitamente si la fuente usa otra codificación."
    )


def _delimiter_score(sample: str, delimiter: str, header: int) -> tuple[float, int, int]:
    try:
        rows = list(islice(csv.reader(io.StringIO(sample), delimiter=delimiter), header + 31))
    except csv.Error:
        return (0.0, 0, 0)
    rows = [row for row in rows[header:] if row and any(str(value).strip() for value in row)]
    if not rows:
        return (0.0, 0, 0)
    header_width = len(rows[0])
    if header_width < 2:
        return (0.0, header_width, 0)
    comparable = rows[1:] or rows
    consistent = sum(len(row) == header_width for row in comparable)
    consistency = consistent / len(comparable)
    return (consistency, header_width, consistent)


def _detect_delimiter(path: Path, encoding: str, requested: str | None, header: int) -> str:
    if requested is not None:
        return requested
    with path.open("r", encoding=encoding, errors="strict", newline="") as stream:
        sample = stream.read(_MAX_DELIMITER_SAMPLE_CHARS)
    scores = {
        delimiter: _delimiter_score(sample, delimiter, header)
        for delimiter in _DELIMITER_CANDIDATES
    }
    best = max(
        _DELIMITER_CANDIDATES,
        key=lambda delimiter: (scores[delimiter][0], scores[delimiter][1], scores[delimiter][2]),
    )
    consistency, width, _ = scores[best]
    if width < 2 or consistency < 0.8:
        readable = {repr(key): value for key, value in scores.items()}
        raise SchemaError(
            "No fue posible inferir un delimitador estable entre coma, tab, punto y coma "
            f"y pipe. Declare delimiter=. Diagnóstico: {readable}."
        )
    return best


def _read_text_headers(path: Path, encoding: str, delimiter: str, header: int) -> list[str]:
    with path.open("r", encoding=encoding, errors="strict", newline="") as stream:
        reader = csv.reader(stream, delimiter=delimiter)
        try:
            row = next(islice(reader, header, header + 1))
        except StopIteration as exc:
            raise SchemaError(f"El archivo no contiene la fila de encabezado {header}.") from exc
        except csv.Error as exc:
            raise SchemaError("Encabezado CSV/TXT mal formado.") from exc
    headers = [str(value) for value in row]
    _validate_headers(headers)
    return headers


def _validate_headers(headers: Sequence[str]) -> None:
    if not headers:
        raise SchemaError("La fuente no contiene columnas.")
    counts: dict[str, int] = {}
    for header in headers:
        counts[header] = counts.get(header, 0) + 1
    duplicates = sorted(header for header, count in counts.items() if count > 1)
    if duplicates:
        raise SchemaError(f"Encabezados duplicados no permitidos: {duplicates}.")


def _normalize_header(value: str) -> str:
    ascii_value = (
        unicodedata.normalize("NFKD", str(value))
        .encode("ascii", errors="ignore")
        .decode("ascii")
        .casefold()
    )
    return re.sub(r"[^a-z0-9]+", "", ascii_value)


def _resolve_header(requested: str, headers: Sequence[str]) -> str:
    if requested in headers:
        return requested
    normalized = _normalize_header(requested)
    matches = [header for header in headers if _normalize_header(header) == normalized]
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1:
        raise SchemaError(
            f"La columna solicitada {requested!r} es ambigua al normalizar: {matches}. "
            "Use el encabezado exacto."
        )
    import difflib

    suggestions = difflib.get_close_matches(requested, list(headers), n=3, cutoff=0.45)
    suffix = f" Sugerencias: {suggestions}." if suggestions else ""
    raise SchemaError(f"No existe la columna solicitada {requested!r}.{suffix}")


def _resolve_optional_header(requested: str, headers: Sequence[str]) -> str | None:
    if requested in headers:
        return requested
    normalized = _normalize_header(requested)
    matches = [header for header in headers if _normalize_header(header) == normalized]
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1:
        raise SchemaError(
            f"La columna opcional {requested!r} es ambigua al normalizar: {matches}. "
            "Use el encabezado exacto."
        )
    return None


def _build_column_plan(headers: Sequence[str], spec: SourceSpec) -> _ColumnPlan:
    _validate_headers(headers)
    resolved_mapping = {
        canonical: _resolve_header(source, headers)
        for canonical, source in spec.column_mapping.items()
    }
    missing_optional: list[str] = []
    for canonical, source in spec.optional_column_mapping.items():
        resolved = _resolve_optional_header(source, headers)
        if resolved is None:
            missing_optional.append(canonical)
        else:
            resolved_mapping[canonical] = resolved
    resolved_passthrough = [_resolve_header(source, headers) for source in spec.passthrough_columns]
    mapped_sources = set(resolved_mapping.values())
    passthrough_unique = [source for source in resolved_passthrough if source not in mapped_sources]
    if spec.keep_unmapped:
        selected: tuple[str, ...] | None = None
        remaining = [header for header in headers if header not in mapped_sources]
    else:
        wanted = mapped_sources | set(passthrough_unique)
        selected = tuple(header for header in headers if header in wanted)
        remaining = passthrough_unique
    rename = {source: canonical for canonical, source in resolved_mapping.items()}
    output_order = tuple(resolved_mapping) + tuple(
        column for column in remaining if column not in resolved_mapping
    )
    if len(set(output_order)) != len(output_order):
        raise SchemaError(
            "El mapeo produce columnas de salida duplicadas. Revise columnas canónicas y passthrough."
        )
    return _ColumnPlan(
        selected=selected,
        rename=MappingProxyType(rename),
        output_order=output_order,
        resolved_mapping=MappingProxyType(resolved_mapping),
        missing_optional=tuple(missing_optional),
    )


def _register_plan(plan: _ColumnPlan, spec: SourceSpec, state: _ReportState) -> None:
    state.resolved_mapping = dict(plan.resolved_mapping)
    state.columns = plan.output_order
    state.missing_optional = tuple(plan.missing_optional)
    for canonical in plan.missing_optional:
        source = spec.optional_column_mapping[canonical]
        warning = (
            f"Columna opcional {canonical!r} (fuente {source!r}) ausente en "
            f"'{spec.name}'; se omitió."
        )
        if warning not in state.warnings:
            state.warnings.append(warning)


def _sentinel_mask(series: pd.Series, sentinels: Sequence[str]) -> pd.Series:
    string = series.astype("string").str.strip().str.casefold()
    normalized = {str(value).strip().casefold() for value in sentinels}
    mask = series.isna() | string.isin(normalized)
    if pd.api.types.is_numeric_dtype(series.dtype):
        numeric = pd.to_numeric(series, errors="coerce")
        for sentinel in normalized:
            try:
                numeric_sentinel = float(sentinel)
            except ValueError:
                continue
            if math.isfinite(numeric_sentinel):
                mask |= numeric.eq(numeric_sentinel)
    return mask.fillna(True)


def _identifier_input(series: pd.Series) -> tuple[pd.Series, pd.Series, bool]:
    """Convierte numéricos sin fabricar el sufijo '.0' ni ocultar precisión."""
    if pd.api.types.is_integer_dtype(series.dtype):
        return series.astype("string"), pd.Series(False, index=series.index), True
    if pd.api.types.is_float_dtype(series.dtype):
        numeric_array = pd.to_numeric(series, errors="coerce").to_numpy(
            dtype=float, na_value=np.nan
        )
        finite = np.isfinite(numeric_array)
        integral = np.mod(numeric_array, 1) == 0
        exact = np.abs(numeric_array) <= _FLOAT_EXACT_INTEGER_LIMIT
        valid = pd.Series(finite & integral & exact, index=series.index)
        numeric = pd.Series(numeric_array, index=series.index)
        converted = numeric.where(valid).astype("Int64").astype("string")
        invalid_numeric = series.notna() & ~valid
        return converted, invalid_numeric, True
    return series.astype("string"), pd.Series(False, index=series.index), False


def _record_invalid(
    column: str,
    raw: pd.Series,
    invalid: pd.Series,
    spec: SourceSpec,
    state: _ReportState,
) -> None:
    count = int(invalid.sum())
    if not count:
        return
    state.invalid_values[column] = state.invalid_values.get(column, 0) + count
    if spec.invalid_values is InvalidValuePolicy.RAISE:
        samples = raw.loc[invalid].astype("string").dropna().head(5).tolist()
        raise SchemaError(
            f"Fuente '{spec.name}', columna {column!r}: {count:,} valores no cumplen "
            f"el tipo declarado. Muestra: {samples}."
        )


def _transform_chunk(
    frame: pd.DataFrame,
    plan: _ColumnPlan,
    spec: SourceSpec,
    state: _ReportState,
) -> pd.DataFrame:
    frame.rename(columns=dict(plan.rename), inplace=True)
    output_order = list(plan.output_order)
    if spec.keep_unmapped:
        output_order = [
            *plan.output_order,
            *(c for c in frame.columns if c not in plan.output_order),
        ]
    if len(set(frame.columns)) != len(frame.columns):
        raise SchemaError("La renombrada produjo columnas de salida duplicadas.")
    frame = frame.loc[:, output_order]
    string_dtype = _preferred_string_dtype()

    for column in frame.columns:
        raw = frame[column]
        sentinels = (*spec.global_null_values, *spec.null_values.get(column, ()))
        missing = _sentinel_mask(raw, sentinels)
        column_type = spec.column_types.get(column, ColumnType.STRING)
        if column_type is ColumnType.IDENTIFIER:
            identifier_format = cast(
                IdentifierFormat,
                spec.identifier_formats.get(column, IdentifierFormat()),
            )
            source, invalid_numeric, was_numeric = _identifier_input(raw)
            normalized = normalizar_identificador(
                source.mask(missing, pd.NA),
                modo=identifier_format.mode,
                min_longitud=identifier_format.min_length,
                max_longitud=identifier_format.max_length,
            ).replace("", pd.NA)
            invalid = (~missing) & (invalid_numeric | normalized.isna())
            _record_invalid(column, raw, invalid, spec, state)
            if was_numeric:
                warning = (
                    f"{column}: la fuente almacenó el identificador como número; "
                    "los ceros iniciales ya no son recuperables."
                )
                if warning not in state.warnings:
                    state.warnings.append(warning)
            frame[column] = normalized.astype(string_dtype)
        elif column_type is ColumnType.NUMBER:
            numeric_format = cast(
                NumericFormat,
                spec.numeric_formats.get(column, NumericFormat()),
            )
            normalized = normalizar_numero(
                raw.mask(missing, pd.NA),
                separador_decimal=numeric_format.decimal_separator,
                separador_miles=numeric_format.thousands_separator,
            )
            numeric = pd.to_numeric(normalized, errors="coerce").astype("Float64")
            invalid = (~missing) & numeric.isna()
            _record_invalid(column, raw, invalid, spec, state)
            frame[column] = numeric
        else:
            # Los parsers preservan los saltos físicos del host dentro de campos
            # entrecomillados. Canonizarlos evita que el mismo archivo lógico
            # produzca valores distintos en Windows (CRLF) y POSIX (LF).
            cleaned = (
                raw.astype("string")
                .str.replace("\r\n", "\n", regex=False)
                .str.replace("\r", "\n", regex=False)
                .str.strip()
                .mask(missing, pd.NA)
            )
            frame[column] = cleaned.astype(string_dtype)

    state.rows += len(frame)
    state.columns = tuple(str(column) for column in frame.columns)
    return frame


def _iter_text(
    materialized: _MaterializedInput,
    spec: SourceSpec,
    state: _ReportState,
) -> Iterator[pd.DataFrame]:
    encoding = _detect_encoding(materialized.path, spec.encoding)
    delimiter = _detect_delimiter(materialized.path, encoding, spec.delimiter, spec.header)
    headers = _read_text_headers(materialized.path, encoding, delimiter, spec.header)
    plan = _build_column_plan(headers, spec)
    state.encoding = encoding
    state.delimiter = delimiter
    _register_plan(plan, spec, state)
    use_pyarrow = spec.text_engine == "pyarrow" or (
        spec.text_engine == "auto" and importlib.util.find_spec("pyarrow") is not None
    )
    if use_pyarrow:
        try:
            import pyarrow as pa
            import pyarrow.csv as pacsv
        except ImportError as exc:  # pragma: no cover - protegido por find_spec
            raise ImportError("text_engine='pyarrow' requiere pyarrow>=14.") from exc
        selected = list(plan.selected) if plan.selected is not None else list(headers)
        read_options = pacsv.ReadOptions(
            use_threads=True,
            block_size=spec.text_block_size_bytes,
            skip_rows=spec.header,
            encoding=encoding,
        )
        parse_options = pacsv.ParseOptions(
            delimiter=delimiter,
            newlines_in_values=spec.newlines_in_values,
        )
        convert_options = pacsv.ConvertOptions(
            column_types={column: pa.string() for column in selected},
            include_columns=selected,
            null_values=[],
            strings_can_be_null=False,
        )
        state.engine = "pyarrow"
        try:
            reader = pacsv.open_csv(
                materialized.path,
                read_options=read_options,
                parse_options=parse_options,
                convert_options=convert_options,
            )
            for batch in reader:
                yield _transform_chunk(
                    batch.to_pandas(types_mapper=pd.ArrowDtype), plan, spec, state
                )
        except (pa.ArrowInvalid, pa.ArrowNotImplementedError) as exc:
            raise SchemaError(
                f"CSV/TXT inválido para delimitador={delimiter!r}, encoding={encoding!r}: {exc}"
            ) from exc
        return

    state.engine = "pandas"
    reader = pd.read_csv(
        materialized.path,
        sep=delimiter,
        encoding=encoding,
        encoding_errors="strict",
        header=spec.header,
        usecols=list(plan.selected) if plan.selected is not None else None,
        dtype=_preferred_string_dtype(),
        keep_default_na=False,
        na_filter=False,
        chunksize=spec.chunksize,
        engine="c",
        memory_map=True,
        on_bad_lines="error",
    )
    for chunk in reader:
        yield _transform_chunk(chunk, plan, spec, state)


def _iter_xlsx(
    materialized: _MaterializedInput,
    spec: SourceSpec,
    state: _ReportState,
) -> Iterator[pd.DataFrame]:
    state.engine = "openpyxl-read-only"
    # XLSX/XLSM también son contenedores ZIP. Validarlos antes de que
    # openpyxl descomprima XML evita que esta ruta eluda los mismos límites
    # de miembros, tamaño, ratio, cifrado y rutas aplicados a un ZIP externo.
    try:
        with zipfile.ZipFile(materialized.path) as archive:
            _validate_zip_infos(archive.infolist(), spec)
    except ArchiveSafetyError:
        raise
    except (OSError, EOFError, RuntimeError, zipfile.BadZipFile) as exc:
        raise ArchiveSafetyError(f"XLSX/XLSM inválido: {materialized.path}.") from exc
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependencia declarada
        raise ImportError("Leer XLSX requiere openpyxl>=3.1.") from exc
    workbook = load_workbook(materialized.path, read_only=True, data_only=True)
    try:
        if isinstance(spec.sheet_name, int):
            try:
                sheet = workbook.worksheets[spec.sheet_name]
            except IndexError as exc:
                raise SchemaError(
                    f"sheet_name={spec.sheet_name} fuera de rango; hojas={workbook.sheetnames}."
                ) from exc
        else:
            if spec.sheet_name not in workbook.sheetnames:
                raise SchemaError(
                    f"No existe la hoja {spec.sheet_name!r}; hojas={workbook.sheetnames}."
                )
            sheet = workbook[spec.sheet_name]
        rows = sheet.iter_rows(values_only=True)
        try:
            header_row = next(islice(rows, spec.header, spec.header + 1))
        except StopIteration as exc:
            raise SchemaError(f"XLSX no contiene la fila de encabezado {spec.header}.") from exc
        headers = ["" if value is None else str(value) for value in header_row]
        plan = _build_column_plan(headers, spec)
        _register_plan(plan, spec, state)
        selected_names = list(headers) if plan.selected is None else list(plan.selected)
        positions = [headers.index(name) for name in selected_names]
        columns: dict[str, list[object]] = {name: [] for name in selected_names}
        buffered = 0
        for row in rows:
            values = [row[position] if position < len(row) else None for position in positions]
            # Una fila puede no traer ninguna de las columnas proyectadas y
            # aun así ser un registro físico real (hay datos no seleccionados).
            # Decidir sobre la fila completa evita pérdida silenciosa y conserva
            # la correspondencia una-fila-de-entrada ↔ una-fila-de-salida.
            if all(value is None for value in row):
                continue
            for name, value in zip(selected_names, values, strict=True):
                columns[name].append(value)
            buffered += 1
            if buffered >= spec.chunksize:
                yield _transform_chunk(pd.DataFrame(columns), plan, spec, state)
                columns = {name: [] for name in selected_names}
                buffered = 0
        if buffered:
            yield _transform_chunk(pd.DataFrame(columns), plan, spec, state)
    finally:
        workbook.close()


def _iter_parquet(
    materialized: _MaterializedInput,
    spec: SourceSpec,
    state: _ReportState,
) -> Iterator[pd.DataFrame]:
    state.engine = "pyarrow-parquet"
    try:
        import pyarrow.parquet as pq
    except ImportError as exc:  # pragma: no cover - dependencia declarada
        raise ImportError("Leer Parquet requiere pyarrow>=14.") from exc
    parquet = pq.ParquetFile(materialized.path)
    headers = list(parquet.schema_arrow.names)
    plan = _build_column_plan(headers, spec)
    _register_plan(plan, spec, state)
    selected = list(plan.selected) if plan.selected is not None else None
    for batch in parquet.iter_batches(batch_size=spec.chunksize, columns=selected):
        frame = batch.to_pandas(types_mapper=pd.ArrowDtype)
        yield _transform_chunk(frame, plan, spec, state)


def _iter_source_chunks(spec: SourceSpec, state: _ReportState) -> Iterator[pd.DataFrame]:
    with _materialize_input(spec) as materialized:
        input_format = _detect_input_format(
            materialized.path, materialized.logical_name, InputFormat(spec.format)
        )
        state.input_format = input_format
        state.compression = materialized.compression
        state.archive_member = materialized.archive_member
        if input_format in {InputFormat.CSV, InputFormat.TXT}:
            yield from _iter_text(materialized, spec, state)
        elif input_format is InputFormat.XLSX:
            yield from _iter_xlsx(materialized, spec, state)
        elif input_format is InputFormat.PARQUET:
            yield from _iter_parquet(materialized, spec, state)
        else:  # pragma: no cover - enum exhaustivo
            raise IngestionError(f"Formato no soportado: {input_format.value}.")


def iter_source_chunks(spec: SourceSpec) -> Iterator[pd.DataFrame]:
    """Itera una fuente por bloques; interfaz recomendada para Colab gratuito.

    El archivo ZIP/GZIP, si existe, se descomprime a disco temporal y nunca a
    RAM. El temporal permanece vivo hasta agotar o cerrar el iterador.
    """
    state = _ReportState()
    yield from _iter_source_chunks(spec, state)


def _empty_frame(columns: Sequence[str]) -> pd.DataFrame:
    return pd.DataFrame({column: pd.Series(dtype=_preferred_string_dtype()) for column in columns})


def load_source(spec: SourceSpec) -> LoadedSource:
    """Carga una fuente completa y devuelve datos más reporte auditable.

    Para archivos que no caben holgadamente en RAM use ``iter_source_chunks``.
    ``load_source`` materializa el resultado final por definición.
    """
    state = _ReportState()
    chunks = list(_iter_source_chunks(spec, state))
    if not chunks:
        output_columns = state.columns or (
            tuple(spec.column_mapping)
            + tuple(spec.optional_column_mapping)
            + tuple(spec.passthrough_columns)
        )
        data = _empty_frame(output_columns)
        state.columns = output_columns
    elif len(chunks) == 1:
        data = chunks[0].reset_index(drop=True)
    else:
        data = pd.concat(chunks, ignore_index=True)
    report = SourceLoadReport(
        source_name=spec.name,
        source_path=Path(spec.path),
        input_format=state.input_format,
        compression=state.compression,
        encoding=state.encoding,
        delimiter=state.delimiter,
        engine=state.engine,
        archive_member=state.archive_member,
        rows=state.rows,
        columns=state.columns,
        resolved_mapping=MappingProxyType(dict(state.resolved_mapping)),
        invalid_values=MappingProxyType(dict(state.invalid_values)),
        warnings=tuple(state.warnings),
        missing_optional=state.missing_optional,
    )
    return LoadedSource(data=data, report=report)


def load_sources(specs: Sequence[SourceSpec]) -> dict[str, LoadedSource]:
    """Carga contratos en orden, rechazando nombres duplicados."""
    names = [spec.name for spec in specs]
    if len(set(names)) != len(names):
        raise ValueError("Los nombres de SourceSpec deben ser únicos.")
    return {spec.name: load_source(spec) for spec in specs}

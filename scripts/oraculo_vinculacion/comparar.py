"""Comparador de paridad: tablas y Excel del motor original contra el paquete nuevo.

Uso:  python scripts/oraculo_vinculacion/comparar.py <etiqueta> <modulo_nuevo> [escenario ...]
      (el módulo nuevo debe exponer ``ejecutar`` y ``exportar_excel``; la referencia sale de ``correr_original.py``)

``TEXTOS_DECLARADOS``: frases de la salida del motor original que traían marcas reales de empresas y que el
paquete nuevo escribe con marcas inventadas. El repositorio no puede contener esas frases, así que viven en
``ORACULO_DIR/textos_declarados.json`` ({"texto original": "texto nuevo"}); se aplican a los DOS lados antes de
comparar, así que no ocultan ninguna otra diferencia. Si el arnés reporta una diferencia de texto de ejemplo
(«p. ej. …», glosario del LEEME, guía de revisión), agregue el par al archivo.
"""

from __future__ import annotations

import importlib
import json
import pickle
import re
import sys
import tempfile
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from _rutas import ORACULO_DIR, usar_apoyo_de_pruebas, usar_biblioteca

FECHA = re.compile(r"\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(:\d{2})?([+-]\d{2}:\d{2})?")
MOTOR = re.compile(r"motor \S+(?: \S+)? sobre rues-linker [\d.]+")
#: Filas de PARAMETROS que cambian en cada corrida.
PARAMETROS_VOLATILES = ("fecha_ejecucion_utc", "segundos", "motor")


def _cargar_textos_declarados() -> dict[str, str]:
    ruta = ORACULO_DIR / "textos_declarados.json"
    return json.loads(ruta.read_text(encoding="utf-8")) if ruta.exists() else {}


TEXTOS_DECLARADOS = _cargar_textos_declarados()


def _norm(v: Any) -> Any:
    if isinstance(v, str):
        for viejo, nuevo in TEXTOS_DECLARADOS.items():
            v = v.replace(viejo, nuevo)
        return MOTOR.sub("motor <MOTOR>", FECHA.sub("<FECHA>", v))
    return v


def _igual(x: Any, y: Any) -> bool:
    if x is None and y is None:
        return True
    if isinstance(x, float) and isinstance(y, float):
        return bool((np.isnan(x) and np.isnan(y)) or x == y)
    if isinstance(x, (list, tuple, set, frozenset)) or isinstance(y, (list, tuple, set, frozenset)):
        return bool(x == y)
    try:
        return bool(x == y)
    except (TypeError, ValueError):
        return str(x) == str(y)


def _valores(serie: pd.Series) -> list[Any]:
    return serie.astype(object).where(serie.notna(), None).map(_norm).tolist()


def _dif_columnas(
    nombre: str, a: pd.DataFrame, b: pd.DataFrame
) -> tuple[list[str], pd.DataFrame | None]:
    """Diferencias de columnas; devuelve ``b`` reordenado si solo cambia el orden (``None`` si no se puede seguir)."""
    if list(a.columns) == list(b.columns):
        return [], b
    solo_a = [c for c in a.columns if c not in b.columns]
    solo_b = [c for c in b.columns if c not in a.columns]
    dif = [f"{nombre}: columnas distintas\n   solo original: {solo_a}\n   solo nuevo: {solo_b}"]
    if sorted(map(str, a.columns)) != sorted(map(str, b.columns)):
        return dif, None
    return [*dif, f"{nombre}: mismas columnas en distinto ORDEN"], b[list(a.columns)]


def comparar_df(nombre: str, a: pd.DataFrame, b: pd.DataFrame, enmascarar: Any = ()) -> list[str]:
    """Lista de diferencias entre dos DataFrames (vacía = idénticos)."""
    dif, b2 = _dif_columnas(nombre, a, b)
    if b2 is None:
        return dif
    if len(a) != len(b2):
        return [*dif, f"{nombre}: filas {len(a)} (original) ≠ {len(b2)} (nuevo)"]
    for c in a.columns:
        if str(a[c].dtype) != str(b2[c].dtype):
            dif.append(f"{nombre}.{c}: dtype {a[c].dtype} (original) ≠ {b2[c].dtype} (nuevo)")
        va, vb = _valores(a[c]), _valores(b2[c])
        malas = [i for i, (x, y) in enumerate(zip(va, vb, strict=True)) if not _igual(x, y)]
        if malas and nombre in enmascarar and c in enmascarar[nombre]:
            continue
        if malas:
            i = malas[0]
            dif.append(
                f"{nombre}.{c}: {len(malas)} celdas distintas; fila {i}: original={va[i]!r} nuevo={vb[i]!r}"
            )
    return dif


def _dif_parametros(ref: pd.DataFrame, tabla: pd.DataFrame) -> list[str]:
    pa = ref.set_index("PARAMETRO")["VALOR"]
    pb = tabla.set_index("PARAMETRO")["VALOR"]
    return [
        f"PARAMETROS[{k}]: original={pa.get(k)!r} nuevo={pb.get(k)!r}"
        for k in sorted(set(pa.index) | set(pb.index))
        if k not in PARAMETROS_VOLATILES and pa.get(k) != pb.get(k)
    ]


def comparar_resultado(
    ref: dict,
    tablas: dict,
    resumen: pd.DataFrame,
    por_pais: pd.DataFrame,
    avisos: list,
    enmascarar: Any = None,
) -> list[str]:
    """Compara las 15 tablas, el resumen, el desglose por país y los avisos contra la referencia."""
    enmascarar = enmascarar or {"PARAMETROS": {"VALOR"}}
    dif: list[str] = []
    claves_a, claves_b = list(ref["tablas"]), list(tablas)
    if claves_a != claves_b:
        dif.append(f"tablas: orden/lista distinta\n  original={claves_a}\n  nuevo={claves_b}")
    for k in claves_a:
        if k in tablas:
            dif += comparar_df(k, ref["tablas"][k], tablas[k], enmascarar)
    if "PARAMETROS" in tablas:
        dif += _dif_parametros(ref["tablas"]["PARAMETROS"], tablas["PARAMETROS"])
    dif += comparar_df("resumen", ref["resumen"], resumen)
    dif += comparar_df("por_pais", ref["por_pais"], por_pais)
    if list(ref["avisos"]) != list(avisos):
        dif.append(f"avisos distintos: original={ref['avisos'][:3]} nuevo={list(avisos)[:3]}")
    return dif


def _rgb(color: Any) -> Any:
    return color.rgb if color else None


def _dif_propiedades(nombre: str, a: Any, b: Any) -> list[str]:
    """Dimensiones, paneles, color de pestaña, filtro, anchos, combinadas, formato condicional y validaciones."""
    dif: list[str] = []
    if a.dimensions != b.dimensions:
        dif.append(f"{nombre}: dimensiones {a.dimensions} ≠ {b.dimensions}")
    if a.freeze_panes != b.freeze_panes:
        dif.append(f"{nombre}: freeze_panes {a.freeze_panes} ≠ {b.freeze_panes}")
    if _rgb(a.sheet_properties.tabColor) != _rgb(b.sheet_properties.tabColor):
        dif.append(f"{nombre}: color de pestaña distinto")
    if a.auto_filter.ref != b.auto_filter.ref:
        dif.append(f"{nombre}: filtro {a.auto_filter.ref} ≠ {b.auto_filter.ref}")
    anchos_a = {k: (round(v.width, 2), v.hidden) for k, v in a.column_dimensions.items()}
    anchos_b = {k: (round(v.width, 2), v.hidden) for k, v in b.column_dimensions.items()}
    if anchos_a != anchos_b:
        n = sum(anchos_a.get(k) != anchos_b.get(k) for k in set(anchos_a) | set(anchos_b))
        dif.append(f"{nombre}: anchos/ocultas distintos ({n} columnas)")
    if sorted(a.merged_cells.ranges, key=str) != sorted(b.merged_cells.ranges, key=str):
        dif.append(f"{nombre}: celdas combinadas distintas")
    if len(a.conditional_formatting) != len(b.conditional_formatting):
        dif.append(
            f"{nombre}: formato condicional {len(a.conditional_formatting)} ≠ {len(b.conditional_formatting)}"
        )
    if len(a.data_validations.dataValidation) != len(b.data_validations.dataValidation):
        dif.append(f"{nombre}: validaciones de datos distintas")
    return dif


def _estilo(c: Any) -> tuple:
    return (
        c.number_format,
        c.font.b,
        c.font.name,
        c.font.sz,
        c.fill.fgColor.rgb,
        c.alignment.wrap_text,
        c.hyperlink.target if c.hyperlink else None,
        c.comment.text if c.comment else None,
    )


def _dif_celdas(nombre: str, a: Any, b: Any) -> list[str]:
    malas, primera = 0, None
    for fila_a, fila_b in zip(a.iter_rows(), b.iter_rows(), strict=False):
        if nombre == "PARAMETROS" and fila_a[0].value in PARAMETROS_VOLATILES:
            continue
        for ca, cb in zip(fila_a, fila_b, strict=False):
            va, vb = _norm(ca.value), _norm(cb.value)
            if not (_igual(va, vb) and _estilo(ca) == _estilo(cb)):
                malas += 1
                primera = primera or (ca.coordinate, va, vb, ca.number_format, cb.number_format)
    return (
        [f"{nombre}: {malas} celdas con valor/estilo distinto; primera {primera}"] if malas else []
    )


def comparar_excel(ruta_a: Path | str, ruta_b: Path | str) -> list[str]:
    """Hojas, propiedades y, celda por celda, valores y estilo de dos libros."""
    from openpyxl import load_workbook

    wa, wb = load_workbook(ruta_a), load_workbook(ruta_b)
    if wa.sheetnames != wb.sheetnames:
        return [f"hojas distintas: {wa.sheetnames} ≠ {wb.sheetnames}"]
    dif: list[str] = []
    for nombre in wa.sheetnames:
        dif += _dif_propiedades(nombre, wa[nombre], wb[nombre])
        dif += _dif_celdas(nombre, wa[nombre], wb[nombre])
    return dif


def main(etiqueta: str, modulo: str, escenarios: list[str]) -> int:
    usar_biblioteca()
    usar_apoyo_de_pruebas()
    import apoyo_vinculacion as av

    nuevo = importlib.import_module(modulo)
    carpeta = ORACULO_DIR / f"insumos_inventados_{etiqueta}"
    ins = av.InsumosVinculacion(
        carpeta, carpeta / "ancla.xlsx", carpeta / "descargas.xlsx", carpeta / "decisiones_1.csv",
        carpeta / "decisiones_2.csv",
    )  # fmt: skip
    if not ins.archivo_ancla.exists():
        raise SystemExit("falta correr correr_original.py para esta etiqueta")
    hubo = False
    for escenario in escenarios or list(av.ESCENARIOS):
        base = ORACULO_DIR / f"ref_{etiqueta}" / escenario
        with (base / "tablas.pkl").open("rb") as fh:
            ref = pickle.load(fh)
        res = nuevo.ejecutar(av.configurar(nuevo, ins, escenario=escenario), verboso=False)
        dif = comparar_resultado(ref, res.tablas, res.resumen, res.por_pais, res.avisos)
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "salida_nueva.xlsx"
            nuevo.exportar_excel(res, ruta)
            dif += ["EXCEL: " + d for d in comparar_excel(base / "salida.xlsx", ruta)]
        print(f"PARIDAD [{escenario}]:", "IDÉNTICA" if not dif else f"{len(dif)} DIFERENCIAS")
        for d in dif[:12]:
            print("   -", d)
        hubo = hubo or bool(dif)
    return 1 if hubo else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1], sys.argv[2], sys.argv[3:]))

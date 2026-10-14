#!/usr/bin/env python
"""generar_plantilla_base_real.py — Plantilla sintética para estructurar la BASE REAL y su VERDAD REAL.

Por qué existe
--------------
Las tablas de verdad que había (Ground_Truth_Robusto_V3, Ground_Truth_Exhaustivo,
Ground_Truth_Multicampo_v1) sirven para regresión pero no para estructurar una base
real: no traen un identificador de registro, ni fuente real, ni régimen ni caso por
columna, mezclan la verdad con la perturbación en texto libre, corrompen los NIT dentro
de los grupos positivos y no tienen el archivo de PARES etiquetados que exige
``docs/PROTOCOLO_GROUND_TRUTH.md``. Esta plantilla es el modelo a copiar: cada hoja
tiene los encabezados exactos que leen ``record_linkage.linkage``,
``record_linkage.evaluation.verdad_real.leer_etiquetas``, ``scripts/medir_verdad_real.py``
y ``scripts/banco.py``.

Todo es inventado: las razones sociales se arman con el vocabulario de prueba del
repositorio (KRONOS, ZELVORA, TRIVENTA, PELUVIKA, DULCEVENT, VALDORIA, ORBITAL, MARMOLERIA
QUINTAL…) y no identifican a ninguna empresa real; los correos y sitios web usan el dominio
reservado ``.example``; los NIT, cédulas y teléfonos son números al azar dentro de los rangos
reales (cualquier coincidencia es casual), y el dígito de verificación del NIT es el correcto
(``digito_verificacion_dian`` de la librería, módulo 11) cuando la fila dice que lo trae.

Hojas del Excel (``data/ejemplos/plantilla_base_real.xlsx``), en este orden
----------------------------------------------------------------------------
LEEME
    Texto en una columna: qué es cada hoja, cómo copiar la estructura con datos reales,
    columnas obligatorias y opcionales, cómo se escribe el NIT, qué va en IS_MATCH, la
    regla de decisión «misma empresa» del protocolo y el comando para medir.
FUENTE_CRM, FUENTE_RUES, FUENTE_ADUANAS
    Una hoja por fuente, con las mismas 15 columnas: ID_REGISTRO_FUENTE (único dentro de
    la fuente, p. ej. CRM-000001), NIT, DV, TIPO_IDENTIFICACION (NIT / CC / CE / VACIO),
    RAZON_SOCIAL, NOMBRE_COMERCIAL, DEPARTAMENTO, MUNICIPIO, DIRECCION, TELEFONO, EMAIL,
    SITIO_WEB, CIIU, PAIS (ISO-3) y FECHA_ACTUALIZACION (AAAA-MM-DD). Una misma empresa
    aparece en una, dos o tres fuentes.
VERDAD_GRUPOS
    La partición verdadera por registro: ID_REGISTRO_FUENTE, FUENTE, ID_GRUPO (G0001…;
    mismo grupo = misma empresa), TIPO_ENTIDAD (empresa / persona), REGIMEN (CON_ID si el
    registro trae identificador, SIN_ID si no), CASO (positivo_* / negativo_* / singleton /
    ambiguo_revisar), NIVEL_RUIDO (CLEAN / LIGHT / MODERATE / SEVERE) y VARIACION (qué se le
    cambió respecto del registro canónico del grupo). Con las hojas FUENTE_* es todo lo que
    necesita ``construir_banco`` para escribir el CSV plano de ``scripts/banco.py``. En una
    base real la hoja es opcional.
ETIQUETAS_PARES
    La plantilla de VERDAD REAL: pares etiquetados por personas, con ID_1, ID_2,
    IS_MATCH (1 / 0 / UNCLEAR), ANOTADOR (A, B o «resuelto»), BANDA_ESPERADA (la dificultad
    que el generador diseñó para el par: unido_holgado / unido_borde / no_unido_borde /
    no_unido_lejos; NO es la decisión del motor) y COMENTARIO. Es lo que lee
    ``leer_etiquetas`` y mide ``scripts/medir_verdad_real.py``. En la base real, la banda la
    escribe el muestreador de ``scripts/verificar_plantilla_base_real.py`` como
    BANDA_PUNTAJE, a partir de lo que el motor hizo de verdad.
DICCIONARIO
    Una fila por columna: en qué hoja va, si es obligatoria, tipo, formato, ejemplo y la
    regla de llenado (y qué pasa si falta).
CASOS
    El catálogo de casos que la plantilla ejercita: descripción, qué debe hacer el motor
    (unir / no unir / revisar), ids de ejemplo y cuántos registros lo ejercitan.

Archivos que escribe (``--salida``, por defecto ``data/ejemplos/``)
--------------------------------------------------------------------
- ``plantilla_base_real.xlsx`` con las ocho hojas.
- ``plantilla_<hoja>.csv`` por cada hoja de datos (fuente_crm, fuente_rues, fuente_aduanas,
  verdad_grupos, etiquetas_pares, diccionario, casos).
- ``plantilla_etiquetas_anotador_A.csv`` y ``_B.csv``: los dos archivos «a ciegas» de los
  que sale ETIQUETAS_PARES (difieren en los pares marcados «resuelto»), para probar el
  kappa con ``--etiquetas-b``.
- ``plantilla_banco.csv``: la misma verdad en el formato plano de ``scripts/banco.py``
  (ID_REGISTRO, ID_GROUP, REGIMEN, CASO, NIT, RAZON_SOCIAL, FUENTE, …), construido por
  ``construir_banco`` a partir de las hojas FUENTE_* y VERDAD_GRUPOS; la misma función la
  expone ``scripts/verificar_plantilla_base_real.py --banco`` para una base real.

Determinismo: con la misma ``--semilla`` (42 por defecto) los CSV salen byte a byte
iguales en cualquier proceso; ningún orden sale de un ``set``.

Uso::

    python scripts/generar_plantilla_base_real.py --salida data/ejemplos --semilla 42
"""

from __future__ import annotations

import argparse
import datetime as dt
import random
import sys
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "src"))

from record_linkage.matching.identificadores import (
    base_canonica,
    digito_verificacion_dian,
)

# ---------------------------------------------------------------------------
# Esquema
# ---------------------------------------------------------------------------

FUENTES: tuple[str, ...] = ("CRM", "RUES", "ADUANAS")
HOJAS: tuple[str, ...] = (
    "LEEME",
    "FUENTE_CRM",
    "FUENTE_RUES",
    "FUENTE_ADUANAS",
    "VERDAD_GRUPOS",
    "ETIQUETAS_PARES",
    "DICCIONARIO",
    "CASOS",
)
COLUMNAS_FUENTE: tuple[str, ...] = (
    "ID_REGISTRO_FUENTE",
    "NIT",
    "DV",
    "TIPO_IDENTIFICACION",
    "RAZON_SOCIAL",
    "NOMBRE_COMERCIAL",
    "DEPARTAMENTO",
    "MUNICIPIO",
    "DIRECCION",
    "TELEFONO",
    "EMAIL",
    "SITIO_WEB",
    "CIIU",
    "PAIS",
    "FECHA_ACTUALIZACION",
)
COLUMNAS_VERDAD: tuple[str, ...] = (
    "ID_REGISTRO_FUENTE",
    "FUENTE",
    "ID_GRUPO",
    "TIPO_ENTIDAD",
    "REGIMEN",
    "CASO",
    "NIVEL_RUIDO",
    "VARIACION",
)
COLUMNAS_PARES: tuple[str, ...] = (
    "ID_1",
    "ID_2",
    "IS_MATCH",
    "ANOTADOR",
    "BANDA_ESPERADA",
    "COMENTARIO",
)
#: Lo que reciben los anotadores en la base real: los dos registros y nada del motor.
COLUMNAS_PARES_A_ETIQUETAR: tuple[str, ...] = (
    "ID_1",
    "ID_2",
    "RAZON_SOCIAL_1",
    "NIT_1",
    "MUNICIPIO_1",
    "RAZON_SOCIAL_2",
    "NIT_2",
    "MUNICIPIO_2",
    "IS_MATCH",
    "ANOTADOR",
    "COMENTARIO",
)
#: Lo que el muestreador guarda aparte y se pega al archivo final: de qué banda salió cada par.
COLUMNAS_BANDAS: tuple[str, ...] = (
    "ID_1",
    "ID_2",
    "REGIMEN",
    "BANDA_PUNTAJE",
    "PUNTAJE",
    "UNIDO_POR_EL_MOTOR",
)
COLUMNAS_DICCIONARIO: tuple[str, ...] = (
    "COLUMNA",
    "HOJA",
    "OBLIGATORIA",
    "TIPO",
    "FORMATO",
    "EJEMPLO",
    "REGLA",
)
COLUMNAS_CASOS: tuple[str, ...] = (
    "CASO",
    "DESCRIPCION",
    "QUE_DEBE_HACER_EL_MOTOR",
    "EJEMPLO_IDS",
    "N_REGISTROS",
)
COLUMNAS_BANCO: tuple[str, ...] = (
    "ID_REGISTRO",
    "ID_GROUP",
    "ESTRATO",
    "FUENTE",
    "TIPO_ENTIDAD",
    "REGIMEN",
    "CASO",
    "NIVEL_RUIDO",
    "VARIACION",
    "NIT",
    "RAZON_SOCIAL",
    "DEPARTAMENTO",
    "MUNICIPIO",
    "CIIU",
    "TELEFONO",
    "DIRECCION",
    "EMAIL",
)
NIVELES: tuple[str, ...] = ("CLEAN", "LIGHT", "MODERATE", "SEVERE")
BANDAS: tuple[str, ...] = ("unido_holgado", "unido_borde", "no_unido_borde", "no_unido_lejos")
#: Ancho de la banda «al borde» a cada lado del umbral del motor (score_threshold del perfil).
MARGEN_BANDA = 0.10
#: Las hojas de fuente se reconocen por este prefijo: FUENTE_CRM, FUENTE_RUES, FUENTE_DIAN…
PREFIJO_HOJA_FUENTE = "FUENTE_"
SIN_CAMBIOS = "canónico (sin cambios)"

#: Catálogo de casos: (CASO, DESCRIPCION, QUE_DEBE_HACER_EL_MOTOR).
CATALOGO_CASOS: tuple[tuple[str, str, str], ...] = (
    (
        "positivo_nit_igual",
        "La misma empresa en dos o tres fuentes con el mismo NIT y el mismo nombre "
        "(a lo sumo cambia la forma del municipio o los espacios).",
        "unir",
    ),
    (
        "positivo_nit_con_y_sin_dv",
        "El mismo NIT escrito sin dígito de verificación, con el DV pegado y con puntos y guion.",
        "unir",
    ),
    (
        "positivo_nit_con_ceros_o_puntuacion",
        "El mismo NIT con ceros a la izquierda, con puntos, con guiones o con espacios de borde.",
        "unir",
    ),
    (
        "positivo_sufijo_legal",
        "Misma empresa y mismo NIT; el sufijo legal cambia (SAS / S.A.S. / LTDA / LIMITADA) "
        "o la empresa cambió de forma legal.",
        "unir",
    ),
    (
        "positivo_errata_1_letra",
        "Misma empresa; el nombre trae una letra cambiada, omitida, duplicada o transpuesta.",
        "unir",
    ),
    (
        "positivo_errata_2_letras",
        "Misma empresa; el nombre trae dos erratas.",
        "unir",
    ),
    (
        "positivo_orden_de_palabras",
        "Misma empresa; las palabras del nombre vienen en otro orden.",
        "unir",
    ),
    (
        "positivo_sigla",
        "La sigla frente al nombre largo, con el mismo NIT en las dos fuentes.",
        "unir",
    ),
    (
        "positivo_sigla_sin_nit",
        "La sigla frente al nombre largo sin NIT en ninguna fuente: con el nombre solo no se "
        "puede saber; va a revisión manual.",
        "revisar",
    ),
    (
        "positivo_truncado",
        "Misma empresa; una fuente cortó el nombre (campo de ancho fijo).",
        "unir",
    ),
    (
        "positivo_nombre_comercial_vs_razon_social",
        "Una fuente trae la razón social y la otra el nombre comercial, con el mismo NIT.",
        "unir",
    ),
    (
        "positivo_sin_nit_en_una_fuente",
        "Misma empresa y mismo nombre; el NIT solo viene en una fuente (par de régimen mixto).",
        "unir",
    ),
    (
        "positivo_sin_nit_en_ninguna",
        "Misma empresa sin NIT en ninguna fuente: solo nombre y municipio.",
        "unir",
    ),
    (
        "positivo_sucursal",
        "Mismo NIT; el nombre lleva la ciudad de la sucursal y el municipio cambia.",
        "unir",
    ),
    (
        "positivo_persona_natural",
        "Persona natural con cédula; nombres y apellidos en orden distinto según la fuente.",
        "unir",
    ),
    (
        "positivo_mojibake",
        "Misma empresa; una fuente trae la eñe y las tildes rotas (Ñ → Ã‘).",  # noqa: RUF001
        "unir",
    ),
    (
        "negativo_homonimo_persona",
        "Dos personas distintas con el mismo nombre completo y cédulas distintas.",
        "no unir",
    ),
    (
        "negativo_nombre_identico_nit_distinto",
        "Dos empresas distintas con razón social idéntica y NIT válidos distintos.",
        "no unir",
    ),
    (
        "negativo_nombre_generico",
        "Dos empresas distintas con nombre genérico (INVERSIONES SAS vs INVERSIONES S.A.S.) "
        "y NIT distintos.",
        "no unir",
    ),
    (
        "negativo_filial_matriz",
        "Matriz y filial comparten la raíz del nombre y tienen NIT distintos: son personas "
        "jurídicas distintas.",
        "no unir",
    ),
    (
        "negativo_nit_corrupto_un_digito",
        "Dos empresas distintas cuyos NIT difieren en un solo dígito.",
        "no unir",
    ),
    (
        "negativo_raiz_confundible",
        "Dos empresas distintas con raíces casi iguales (ZELVORA vs ZELVORO).",
        "no unir",
    ),
    (
        "singleton",
        "Empresas o personas que aparecen una sola vez en todas las fuentes.",
        "no unir",
    ),
    (
        "ambiguo_revisar",
        "Pares que una persona no puede decidir con los datos (sin NIT en un lado, nombre parcial o "
        "muy común): en ETIQUETAS_PARES van como UNCLEAR y en VERDAD_GRUPOS son grupos distintos.",
        "revisar",
    ),
)
CASOS: tuple[str, ...] = tuple(c for c, _, _ in CATALOGO_CASOS)

# ---------------------------------------------------------------------------
# Vocabulario inventado
# ---------------------------------------------------------------------------

RAICES: tuple[str, ...] = (
    "ZELVORA",
    "KRONOS",
    "TRIVENTA",
    "PELUVIKA",
    "DULCEVENT",
    "VALDORIA",
    "ORBITAL",
    "QUINTAL",
    "KRAMENTO",
    "QUELMAR",
    "VANTIKO",
    "MIRVALDO",
    "TOLSANA",
    "BRUMELIA",
    "OXIDRA",
    "FELTRAN",
    "NOVARQUE",
    "SILVENTO",
    "ALTAMIRIA",
    "CORVELUX",
    "ESMERAL",
    "FIORDANA",
    "GALVANTE",
    "HELIXTRA",
    "KALMIRA",
    "LUMBRERA",
    "MONTEVIK",
    "NARVALIA",
    "OLMEDAR",
    "PRISMAVOX",
    "RUBRALTA",
    "TARIMBA",
    "ULTRAVEX",
    "VERDINAL",
    "XILOTEX",
    "YARUMBA",
    "ZARCOLIN",
    "ANDAMAR",
    "CALDUMA",
    "DORVEX",
    "FULGORIA",
    "GRANVELA",
    "HORNALIA",
    "INDUVEX",
    "JARPELA",
    "KOLMARA",
    "LAVIRENA",
    "MOLTRIVA",
    "PALVERDE",
    "ROSALMA",
    "TRUVEL",
    "UMBRALIA",
    "VOLTEXA",
    "WENDRAL",
    "XAMPARA",
    "YUBERA",
    "ZAMBRUK",
    "CUMBREX",
    "DELVANIA",
    "ESTRAVIA",
    "FORNELIA",
    "GRISALTA",
    "JOLMENA",
    "KEVRANA",
    "LORVANTE",
    "ORVANIA",
    "PILTREX",
    "QUORANTE",
    "SOLMIRA",
    "WOLVARIA",
    "XENTRIA",
    "ZIBRALTA",
    "PIXOLAR",
    "GANDUMEL",
    "ZOQUITARA",
    "FUMBREL",
    "CORDALIX",
    "TABERNIO",
    "WIXANTE",
    "OLVERANA",
    "NIMBROSA",
    "KETALUM",
    "BRAVELIX",
    "CUNDAMAR",
    "DOLTIKA",
    "ESPIRAVE",
    "GOLBRANA",
    "HISTRALIA",
    "LIMBRATA",
    "NERVALIS",
    "OSTRALUX",
    "PUMBREVA",
    "QUISTRANA",
    "ROLDANIX",
    "TORVALIA",
    "VINDRAMA",
    "XOLTREVA",
    "YAMBRELO",
    "ZUNDAVIA",
)
#: Palabras de sector que van DELANTE de la raíz («COMERCIALIZADORA KRONOS»).
SECTORES_PREFIJO: tuple[str, ...] = (
    "COMERCIALIZADORA",
    "DISTRIBUIDORA",
    "INVERSIONES",
    "IMPORTACIONES",
    "EXPORTACIONES",
    "SOLUCIONES",
    "SERVICIOS",
    "SUMINISTROS",
    "CONSTRUCCIONES",
    "TRANSPORTES",
    "INDUSTRIAS",
    "LABORATORIOS",
    "CONSULTORES",
    "MANUFACTURAS",
    "PRODUCTOS",
)
#: Palabras de sector que van DETRÁS de la raíz («ZELVORA ALIMENTOS»).
SECTORES_SUFIJO: tuple[str, ...] = (
    "ALIMENTOS",
    "TEXTILES",
    "LOGISTICA",
    "QUIMICOS",
    "AGROINDUSTRIAL",
    "TECNOLOGIA",
    "INGENIERIA",
    "FARMACEUTICA",
    "PLASTICOS",
    "METALMECANICA",
    "EDITORIAL",
    "ENERGIA",
    "MINERALES",
    "EMPAQUES",
    "CAFETERA",
    "FLORES",
    "CALZADO",
    "MUEBLES",
    "COSMETICOS",
    "AGRICOLA",
    "PESQUERA",
    "AUTOMOTRIZ",
    "ELECTRICOS",
    "HIDRAULICA",
    "INMOBILIARIA",
    "SEGURIDAD",
    "MARMOLERIA",
)
REGIONES: tuple[str, ...] = (
    "DEL CARIBE",
    "DEL PACIFICO",
    "DEL ORIENTE",
    "DE OCCIDENTE",
    "DEL NORTE",
    "DEL SUR",
    "DE LOS LLANOS",
    "DEL CAUCA",
    "DEL MAGDALENA",
    "DE LA SABANA",
    "DE LOS ANDES",
    "DE COLOMBIA",
)
#: Familias de sufijo legal: la primera forma es la canónica.
FAMILIAS_SUFIJO: dict[str, tuple[str, ...]] = {
    "SAS": ("SAS", "S.A.S.", "S A S", "S.A.S"),
    "LTDA": ("LTDA", "LIMITADA", "LTDA.", "LIMITADA."),
    "SA": ("S.A.", "SA", "S A"),
    "SENC": ("S EN C", "S. EN C.", "S EN C."),
    "EU": ("E.U.", "EU", "E U"),
}
PESO_FAMILIAS: tuple[tuple[str, int], ...] = (
    ("SAS", 55),
    ("LTDA", 25),
    ("SA", 12),
    ("SENC", 5),
    ("EU", 3),
)
#: Cuántas entidades distintas pueden compartir una raíz (con sectores distintos).
USOS_MAXIMOS_RAIZ = 6
#: Cuántas raíces se apartan para entidades sin identificador (una por entidad, sin reuso).
RAICES_EXCLUSIVAS = 45
#: Raíces que usan los casos diseñados (filial, raíz confundible, ambiguos, nombres comerciales):
#: no entran al relleno para que ningún grupo de relleno se parezca a un caso por accidente.
RAICES_RESERVADAS: frozenset[str] = frozenset(
    {"KRONOS", "VALDORIA", "TRIVENTA", "ZELVORA", "KALMIRA", "TRUVEL", "ORBITAL", "QUELMAR"}
)
STOPWORDS_SIGLA: frozenset[str] = frozenset({"DE", "DEL", "LA", "LAS", "LOS", "Y", "&", "EN", "E"})

NOMBRES_PERSONA: tuple[str, ...] = (
    "MARIA FERNANDA",
    "JUAN CARLOS",
    "ANDRES FELIPE",
    "LUIS ALBERTO",
    "CAMILA ANDREA",
    "DIEGO ALEJANDRO",
    "PAULA ANDREA",
    "SANTIAGO",
    "LAURA VALENTINA",
    "JORGE ELIECER",
    "ANA MARIA",
    "CLAUDIA PATRICIA",
    "RICARDO",
    "GLORIA INES",
    "HERNANDO",
    "MARTHA LUCIA",
    "CARLOS ARTURO",
    "SANDRA MILENA",
    "OSCAR IVAN",
    "LILIANA",
)
APELLIDOS: tuple[str, ...] = (
    "OCAMPO",
    "RUIZ",
    "PEÑALOZA",
    "MUÑOZ",
    "CASTAÑEDA",
    "AVILA",
    "LONDOÑO",
    "QUIÑONES",
    "GOMEZ",
    "RODRIGUEZ",
    "MARTINEZ",
    "HERNANDEZ",
    "LOPEZ",
    "GARCIA",
    "RAMIREZ",
    "TORRES",
    "VARGAS",
    "MORENO",
    "JIMENEZ",
    "CASTRO",
    "ORTIZ",
    "ROJAS",
    "SUAREZ",
    "MEJIA",
    "ARANGO",
    "RESTREPO",
    "ZAPATA",
    "BOTERO",
    "CARDONA",
    "OSORIO",
    "VELEZ",
    "HOYOS",
    "GIRALDO",
    "MONTOYA",
    "BEDOYA",
)
NOMBRES_COMERCIALES: tuple[str, ...] = (
    "PANADERIA LA ESPIGA DORADA",
    "FERRETERIA EL TORNILLO DE ORO",
    "RESTAURANTE EL FOGON DE KALMIRA",
    "ALMACEN MODA TRIVENTA",
    "SUPERMERCADO LA CANASTA VERDE",
    "TALLER MOTORES DEL NORTE",
    "OPTICA VISION QUELMAR",
    "PAPELERIA EL LAPIZ AZUL",
    "DROGUERIA VIDA PLENA ZELVORA",
    "LAVANDERIA BURBUJAS DE ORBITAL",
    "CAFETERIA EL GRANO DE VALDORIA",
    "VETERINARIA PATITAS DE KRONOS",
)
#: (DEPARTAMENTO, MUNICIPIO canónico, variantes de escritura del municipio, indicativo).
LUGARES: tuple[tuple[str, str, tuple[str, ...], str], ...] = (
    ("BOGOTA D.C.", "BOGOTA D.C.", ("BOGOTÁ", "BOGOTA", "Bogotá D.C.", "BOGOTA, D.C."), "601"),
    ("ANTIOQUIA", "MEDELLIN", ("MEDELLÍN", "Medellín", "MEDELLIN ANT"), "604"),
    ("VALLE DEL CAUCA", "CALI", ("SANTIAGO DE CALI", "Cali", "CALI VALLE"), "602"),
    ("ATLANTICO", "BARRANQUILLA", ("Barranquilla", "B/QUILLA", "BARRANQUILLA ATL"), "605"),
    ("BOLIVAR", "CARTAGENA", ("CARTAGENA DE INDIAS", "Cartagena", "CARTAGENA BOL"), "605"),
    ("SANTANDER", "BUCARAMANGA", ("Bucaramanga", "B/MANGA", "BUCARAMANGA STDER"), "607"),
    ("RISARALDA", "PEREIRA", ("Pereira", "PEREIRA RDA"), "606"),
    ("CALDAS", "MANIZALES", ("Manizales", "MANIZALES CAL"), "606"),
    ("NORTE DE SANTANDER", "CUCUTA", ("CÚCUTA", "Cúcuta", "SAN JOSE DE CUCUTA"), "607"),
    ("TOLIMA", "IBAGUE", ("IBAGUÉ", "Ibagué"), "608"),
    ("MAGDALENA", "SANTA MARTA", ("Santa Marta", "STA MARTA", "STA. MARTA"), "605"),
    ("META", "VILLAVICENCIO", ("Villavicencio", "VILLAVICENCIO META"), "608"),
    ("NARIÑO", "PASTO", ("SAN JUAN DE PASTO", "Pasto"), "602"),
    ("HUILA", "NEIVA", ("Neiva", "NEIVA HUILA"), "608"),
    ("QUINDIO", "ARMENIA", ("Armenia", "ARMENIA QUINDIO"), "606"),
    ("CORDOBA", "MONTERIA", ("MONTERÍA", "Montería"), "604"),
    ("CUNDINAMARCA", "CHIA", ("CHÍA", "Chía"), "601"),
    ("ANTIOQUIA", "ITAGUI", ("ITAGÜÍ", "Itagüí", "ITAGUI ANT"), "604"),
    ("ANTIOQUIA", "ENVIGADO", ("Envigado", "ENVIGADO ANT"), "604"),
    ("VALLE DEL CAUCA", "YUMBO", ("Yumbo", "YUMBO VALLE"), "602"),
    ("CUNDINAMARCA", "SOACHA", ("Soacha", "SOACHA CUND"), "601"),
    ("ATLANTICO", "SOLEDAD", ("Soledad", "SOLEDAD ATL"), "605"),
)
PAISES_EXTRANJEROS: tuple[tuple[str, str], ...] = (
    ("USA", "INC"),
    ("PAN", "CORP"),
    ("MEX", "SA DE CV"),
    ("ESP", "SL"),
)
CIIU: tuple[str, ...] = (
    "1011", "1030", "1040", "1061", "1081", "1104", "1311", "1410", "1512", "1610",
    "1811", "2011", "2023", "2100", "2220", "2310", "2395", "2410", "2511", "2910",
    "3110", "4111", "4290", "4511", "4620", "4631", "4641", "4645", "4651", "4652",
    "4661", "4663", "4669", "4711", "4719", "4731", "4741", "4752", "4771", "4921",
    "4923", "5011", "5210", "5224", "5229", "5511", "5611", "5811", "6110", "6201",
    "6202", "6209", "6411", "6499", "6810", "6910", "6920", "7010", "7020", "7110",
    "7120", "7310", "7410", "7490", "7820", "8010", "8211", "8299", "8610", "8621",
    "9311", "9602",
)  # fmt: skip
TIPOS_VIA: tuple[str, ...] = ("CALLE", "CARRERA", "AVENIDA", "TRANSVERSAL", "DIAGONAL")
ABREV_VIA_CRM: dict[str, str] = {
    "CALLE": "Cl",
    "CARRERA": "Cra",
    "AVENIDA": "Av",
    "TRANSVERSAL": "Tv",
    "DIAGONAL": "Dg",
}
ABREV_VIA_ADUANAS: dict[str, str] = {
    "CALLE": "CL",
    "CARRERA": "KR",
    "AVENIDA": "AV",
    "TRANSVERSAL": "TV",
    "DIAGONAL": "DG",
}
COMPLEMENTOS: tuple[str, ...] = ("", "", "", "OF 201", "LOCAL 3", "BODEGA 12", "PISO 4", "INT 2")
PREFIJOS_EMAIL: tuple[str, ...] = (
    "contacto",
    "info",
    "ventas",
    "gerencia",
    "comercial",
    "administracion",
)
#: Proveedores «genéricos» de correo, con dominio reservado (RFC 2606): en la base real serán
#: gmail, hotmail…; aquí no puede existir ningún buzón. Cuatro, como los sufijos de dominio propio.
DOMINIOS_GENERICOS: tuple[str, ...] = (
    "correo.example",
    "buzon.example",
    "webmail.example",
    "mail.example",
)
#: Sufijos del dominio propio de una empresa (``zelvora.example``, ``zelvora-sas.example``).
SUFIJOS_DOMINIO: tuple[str, ...] = (".example", "-sas.example", ".example", ".example")
PREFIJOS_CELULAR: tuple[str, ...] = ("300", "301", "310", "311", "312", "313", "314", "315", "320", "321", "350")  # fmt: skip
#: UTF-8 leído como cp1252: la forma más común del «mojibake» en exportaciones.
MOJIBAKE: dict[str, str] = {"Ñ": "Ã‘", "É": "Ã‰", "Ó": "Ã“", "Ú": "Ãš", "Ü": "Ãœ"}  # noqa: RUF001
FECHA_INICIAL = dt.date(2024, 1, 1)
DIAS_RANGO_FECHAS = 1003  # hasta 2026-09-30

# ---------------------------------------------------------------------------
# Modelo
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class ConfiguracionPlantilla:
    """Perillas del generador (sin números mágicos en la lógica).

    Attributes:
        semilla: semilla única de todo el azar.
        salida: carpeta de los archivos.
        grupos_nit_igual: grupos de relleno del caso positivo_nit_igual.
        singletons: entidades que aparecen una sola vez.
        min_por_fuente, max_por_fuente: rango exigido de registros por fuente.
        min_pares, max_pares: rango exigido de pares en ETIQUETAS_PARES.
        pares_control_lejanos: negativos aleatorios «no_unido_lejos».
        fraccion_pares_nit_igual: qué fracción de los grupos nit_igual aporta un par.
    """

    semilla: int = 42
    salida: Path = RAIZ / "data" / "ejemplos"
    grupos_nit_igual: int = 118
    singletons: int = 60
    min_por_fuente: int = 150
    max_por_fuente: int = 300
    min_pares: int = 60
    max_pares: int = 120
    pares_control_lejanos: int = 12
    fraccion_pares_nit_igual: float = 0.10
    nombre_excel: str = "plantilla_base_real.xlsx"


@dataclass
class Entidad:
    """Una persona jurídica o natural inventada, con sus datos canónicos."""

    tipo: str  # empresa | persona
    nucleo: str  # razón social sin sufijo legal
    sufijo: str  # "" para personas
    identificador: str  # base sin DV, solo dígitos; "" si no tiene
    tipo_identificacion: str  # NIT | CC | CE | VACIO
    nombre_comercial: str
    lugar: tuple[str, str, tuple[str, ...], str]
    direccion: tuple[str, int, int, int, str]
    telefono: str
    email: str
    sitio_web: str
    ciiu: tuple[str, ...]
    pais: str = "COL"
    raiz: str = ""

    @property
    def razon_social(self) -> str:
        return f"{self.nucleo} {self.sufijo}".strip()

    @property
    def dv(self) -> str:
        if self.tipo_identificacion == "NIT" and self.identificador:
            return digito_verificacion_dian(self.identificador)
        return ""


@dataclass
class Registro:
    """Una fila de una hoja FUENTE_* más su verdad."""

    fuente: str
    campos: dict[str, str]
    caso: str
    nivel_ruido: str
    variacion: str
    tipo_entidad: str
    grupo: int = -1
    id_registro: str = ""

    @property
    def regimen(self) -> str:
        return "CON_ID" if self.campos["NIT"].strip() else "SIN_ID"


@dataclass
class ParPropuesto:
    """Un par para ETIQUETAS_PARES, antes de asignar ids."""

    a: Registro
    b: Registro
    is_match: str  # "1" | "0" | "UNCLEAR"
    banda: str
    comentario: str
    dificil: bool = False  # los anotadores no coincidieron: ANOTADOR = resuelto


@dataclass
class Bloque:
    """Lo que produce una receta de caso: uno o varios grupos y sus pares sugeridos."""

    caso: str
    grupos: list[list[Registro]]
    pares: list[ParPropuesto] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Utilidades de texto
# ---------------------------------------------------------------------------


def _sin_tildes(texto: str) -> str:
    plano = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in plano if not unicodedata.combining(c)).replace("Ñ", "N")


def _clave_nombre(texto: str) -> str:
    return " ".join(_sin_tildes(texto).upper().split())


def _distancia_hamming(a: str, b: str) -> int:
    if len(a) != len(b):
        return max(len(a), len(b))
    return sum(1 for x, y in zip(a, b, strict=True) if x != y)


def _mojibake(texto: str) -> str:
    for bueno, roto in MOJIBAKE.items():
        texto = texto.replace(bueno, roto)
    return texto


def _formato_nit(base: str, dv: str, forma: str) -> str:
    """Escrituras habituales de un NIT: la base, la base con DV, con puntos, con guiones…"""
    if forma == "base":
        return base
    if forma == "dv_pegado":
        return base + dv
    if forma == "puntos_guion_dv":
        return f"{base[:-6]}.{base[-6:-3]}.{base[-3:]}-{dv}"
    if forma == "puntos":
        return f"{base[:-6]}.{base[-6:-3]}.{base[-3:]}"
    if forma == "guiones":
        return f"{base[:-6]}-{base[-6:-3]}-{base[-3:]}"
    if forma == "guion_dv":
        return f"{base}-{dv}"
    if forma == "ceros":
        return "000" + base
    if forma == "espacios":
        return f" {base} "
    raise ValueError(f"forma de NIT desconocida: {forma}")


def _sigla(nucleo: str) -> str:
    return "".join(t[0] for t in nucleo.split() if t not in STOPWORDS_SIGLA)


# ---------------------------------------------------------------------------
# Fábrica de entidades
# ---------------------------------------------------------------------------


class Fabrica:
    """Inventa entidades sin colisiones: NIT a distancia ≥ 3, nombres distintos, DV válido."""

    def __init__(self, rng: random.Random) -> None:
        self.rng = rng
        self.identificadores: list[str] = []
        self.nombres_usados: dict[str, int] = {}
        self.usos_raiz: dict[str, int] = dict.fromkeys(RAICES, 0)
        barajadas = [r for r in RAICES if r not in RAICES_RESERVADAS]
        rng.shuffle(barajadas)
        self.pool_exclusivo = barajadas[:RAICES_EXCLUSIVAS]
        self.pool_compartido = barajadas[RAICES_EXCLUSIVAS:]
        self.cursor_raiz = 0

    # -- identificadores -------------------------------------------------------

    def _registrar(self, identificador: str) -> str:
        self.identificadores.append(identificador)
        return identificador

    def _sirve(self, candidato: str, distancia_minima: int = 3) -> bool:
        if base_canonica(candidato) != candidato.lstrip("0"):
            return False  # su último dígito validaría como DV por azar: ambiguo
        if base_canonica(candidato + digito_verificacion_dian(candidato)) != candidato:
            return False
        return all(
            _distancia_hamming(candidato, otro) >= distancia_minima for otro in self.identificadores
        )

    def nit(self) -> str:
        prefijos = ("800", "805", "811", "830", "860", "890", "900", "900", "901", "901")
        while True:
            candidato = self.rng.choice(prefijos) + f"{self.rng.randint(0, 999999):06d}"
            if self._sirve(candidato):
                return self._registrar(candidato)

    def nit_cercano(self, base: str) -> str:
        """Un NIT que difiere de ``base`` en exactamente un dígito (no el primero)."""
        while True:
            pos = self.rng.randint(1, len(base) - 1)
            digito = str((int(base[pos]) + self.rng.randint(1, 9)) % 10)
            candidato = base[:pos] + digito + base[pos + 1 :]
            resto = [x for x in self.identificadores if x != base]
            if (
                base_canonica(candidato) == candidato
                and base_canonica(candidato + digito_verificacion_dian(candidato)) == candidato
                and all(_distancia_hamming(candidato, otro) >= 3 for otro in resto)
            ):
                return self._registrar(candidato)

    def cedula(self) -> str:
        while True:
            if self.rng.random() < 0.5:
                candidato = str(self.rng.randint(2, 7)) + f"{self.rng.randint(0, 9999999):07d}"
            else:
                candidato = "10" + f"{self.rng.randint(0, 99999999):08d}"
            if self._sirve(candidato):
                return self._registrar(candidato)

    # -- nombres ---------------------------------------------------------------

    def raiz(self, *, exclusiva: bool = False) -> str:
        """Una raíz inventada.

        Las entidades CON identificador pueden compartir raíz (hasta USOS_MAXIMOS_RAIZ veces,
        siempre con otro sector): el veto por NIT distinto impide que el motor las una. Las
        entidades SIN identificador piden una raíz ``exclusiva`` que nadie más usa, para que
        ningún grupo de relleno se parezca a ellas por accidente (el nombre es toda su evidencia).
        """
        if exclusiva:
            if not self.pool_exclusivo:
                raise RuntimeError("se agotaron las raíces exclusivas: suba RAICES_EXCLUSIVAS")
            r = self.pool_exclusivo.pop(0)
            self.usos_raiz[r] = USOS_MAXIMOS_RAIZ
            return r
        for _ in range(len(self.pool_compartido) * USOS_MAXIMOS_RAIZ):
            r = self.pool_compartido[self.cursor_raiz % len(self.pool_compartido)]
            self.cursor_raiz += 1
            if self.usos_raiz[r] < USOS_MAXIMOS_RAIZ:
                self.usos_raiz[r] += 1
                return r
        raise RuntimeError("se agotaron las raíces compartidas: amplíe RAICES")

    def sufijo(self, familia: str | None = None) -> str:
        if familia is None:
            familia = self.rng.choices(
                [f for f, _ in PESO_FAMILIAS], weights=[p for _, p in PESO_FAMILIAS]
            )[0]
        return FAMILIAS_SUFIJO[familia][0]

    def nucleo_empresa(
        self, *, raiz: str | None = None, minimo_tokens: int = 2, exclusiva: bool = False
    ) -> str:
        """Un núcleo de razón social nuevo (raíz + sector [+ región])."""
        r = raiz or self.raiz(exclusiva=exclusiva)
        for _ in range(200):
            patron = self.rng.choice(("RS", "SR", "RSG", "SRG", "IR"))
            if patron == "RS":
                nucleo = f"{r} {self.rng.choice(SECTORES_SUFIJO)}"
            elif patron == "SR":
                nucleo = f"{self.rng.choice(SECTORES_PREFIJO)} {r}"
            elif patron == "RSG":
                nucleo = f"{r} {self.rng.choice(SECTORES_SUFIJO)} {self.rng.choice(REGIONES)}"
            elif patron == "SRG":
                nucleo = f"{self.rng.choice(SECTORES_PREFIJO)} {r} {self.rng.choice(REGIONES)}"
            else:
                nucleo = f"INVERSIONES {r} Y {self.rng.choice(APELLIDOS)}"
            if len(nucleo.split()) < minimo_tokens:
                continue
            if self.reservar_nombre(nucleo):
                return nucleo
        raise RuntimeError("no se pudo inventar un núcleo nuevo")

    def reservar_nombre(self, nucleo: str) -> bool:
        clave = _clave_nombre(nucleo)
        if clave in self.nombres_usados:
            return False
        self.nombres_usados[clave] = 1
        return True

    # -- datos de contacto -----------------------------------------------------

    def lugar(self) -> tuple[str, str, tuple[str, ...], str]:
        pesos = [8, 5, 4, 3, 2, 2, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1]
        return self.rng.choices(LUGARES, weights=pesos)[0]

    def direccion(self) -> tuple[str, int, int, int, str]:
        return (
            self.rng.choice(TIPOS_VIA),
            self.rng.randint(1, 170),
            self.rng.randint(1, 120),
            self.rng.randint(1, 99),
            self.rng.choice(COMPLEMENTOS),
        )

    def telefono(self, indicativo: str) -> str:
        if self.rng.random() < 0.45:
            return self.rng.choice(PREFIJOS_CELULAR) + f"{self.rng.randint(0, 9999999):07d}"
        return indicativo + f"{self.rng.randint(2000000, 8999999):07d}"

    def email_empresa(self, raiz: str, sector: str) -> tuple[str, str]:
        """(email, dominio propio o "")."""
        if self.rng.random() < 0.22:
            usuario = (raiz + sector[:4]).lower()
            return f"{usuario}@{self.rng.choice(DOMINIOS_GENERICOS)}", ""
        dominio = _sin_tildes(raiz).lower() + self.rng.choice(SUFIJOS_DOMINIO)
        return f"{self.rng.choice(PREFIJOS_EMAIL)}@{dominio}", dominio

    def email_persona(self, nombres: str, apellidos: str) -> str:
        n = _sin_tildes(nombres.split()[0]).lower()
        a = _sin_tildes(apellidos.split()[0]).lower()
        return f"{n}.{a}{self.rng.randint(1, 99)}@{self.rng.choice(DOMINIOS_GENERICOS)}"

    def ciiu(self, n: int | None = None) -> tuple[str, ...]:
        k = n if n is not None else self.rng.choices((1, 2, 3), weights=(55, 30, 15))[0]
        return tuple(self.rng.sample(CIIU, k))

    # -- entidades -------------------------------------------------------------

    def empresa(
        self,
        *,
        con_nit: bool = True,
        nucleo: str | None = None,
        familia_sufijo: str | None = None,
        sufijo: str | None = None,
        lugar: tuple[str, str, tuple[str, ...], str] | None = None,
        nombre_comercial: str = "",
        nit: str | None = None,
        pais: str = "COL",
        minimo_tokens: int = 2,
        raiz_exclusiva: bool = False,
    ) -> Entidad:
        if nucleo is None:
            nucleo = self.nucleo_empresa(
                minimo_tokens=minimo_tokens, exclusiva=raiz_exclusiva or not con_nit
            )
        else:
            self.reservar_nombre(nucleo)  # un nombre repetido a propósito no bloquea
        raiz = next((t for t in nucleo.split() if t in RAICES), nucleo.split()[0])
        sector = next(
            (t for t in nucleo.split() if t in SECTORES_SUFIJO or t in SECTORES_PREFIJO),
            "empresa",
        )
        lugar = lugar or self.lugar()
        email, dominio = self.email_empresa(raiz, sector)
        sitio = ""
        if dominio and self.rng.random() < 0.5:
            sitio = self.rng.choice(("www.", "https://www.", "https://")) + dominio
        identificador = "" if not con_nit else (nit if nit is not None else self.nit())
        return Entidad(
            tipo="empresa",
            nucleo=nucleo,
            sufijo=sufijo if sufijo is not None else self.sufijo(familia_sufijo),
            identificador=identificador,
            tipo_identificacion="NIT" if identificador else "VACIO",
            nombre_comercial=nombre_comercial,
            lugar=lugar,
            direccion=self.direccion(),
            telefono=self.telefono(lugar[3]),
            email=email,
            sitio_web=sitio,
            ciiu=self.ciiu(),
            pais=pais,
            raiz=raiz,
        )

    def persona(
        self,
        *,
        con_cedula: bool = True,
        nombres: str | None = None,
        apellidos: str | None = None,
        lugar: tuple[str, str, tuple[str, ...], str] | None = None,
        nombre_comercial: str = "",
        extranjera: bool = False,
    ) -> Entidad:
        nombres = nombres or self.rng.choice(NOMBRES_PERSONA)
        apellidos = apellidos or " ".join(self.rng.sample(APELLIDOS, 2))
        nucleo = f"{apellidos} {nombres}"  # orden oficial: apellidos y luego nombres
        self.reservar_nombre(nucleo)
        lugar = lugar or self.lugar()
        identificador = self.cedula() if con_cedula else ""
        tipo_id = "VACIO" if not identificador else ("CE" if extranjera else "CC")
        return Entidad(
            tipo="persona",
            nucleo=nucleo,
            sufijo="",
            identificador=identificador,
            tipo_identificacion=tipo_id,
            nombre_comercial=nombre_comercial,
            lugar=lugar,
            direccion=self.direccion(),
            telefono=self.telefono(lugar[3]),
            email=self.email_persona(nombres, apellidos),
            sitio_web="",
            ciiu=self.ciiu(1),
            raiz=apellidos.split()[0],
        )


# ---------------------------------------------------------------------------
# Variantes: un registro por fuente, con cada perturbación declarada
# ---------------------------------------------------------------------------


class Variante:
    """Construye el registro de una entidad en una fuente acumulando lo que se le cambia.

    ``NIVEL_RUIDO`` y ``VARIACION`` salen de lo aplicado, no de una etiqueta puesta a mano:
    una fila no-CLEAN siempre difiere del canónico en NIT, RAZON_SOCIAL o MUNICIPIO.
    """

    def __init__(self, gen: Generador, entidad: Entidad, fuente: str) -> None:
        self.gen = gen
        self.ent = entidad
        self.fuente = fuente
        self.campos = gen.campos_base(entidad, fuente)
        self.aplicadas: list[tuple[str, str]] = []

    def aplicar(self, descripcion: str, nivel: str, **cambios: str) -> Variante:
        self.campos.update(cambios)
        self.aplicadas.append((descripcion, nivel))
        return self

    # -- perturbaciones de nombre ---------------------------------------------

    def sufijo(self, forma: str, descripcion: str | None = None) -> Variante:
        nuevo = f"{self.ent.nucleo} {forma}".strip()
        return self.aplicar(
            descripcion or f"sufijo legal: {self.ent.sufijo} → {forma}", "LIGHT", RAZON_SOCIAL=nuevo
        )

    def errata(self, n: int) -> Variante:
        nucleo, detalle = self.gen.errata(self.ent.nucleo, n)
        nivel = "MODERATE" if n == 1 else "SEVERE"
        return self.aplicar(
            f"errata de {n} letra(s): {detalle}",
            nivel,
            RAZON_SOCIAL=f"{nucleo} {self.ent.sufijo}".strip(),
        )

    def orden(self) -> Variante:
        tokens = self.ent.nucleo.split()
        if len(tokens) == 2:
            nuevo = f"{tokens[1]} {tokens[0]}"
        else:
            nuevo = " ".join(tokens[1:] + tokens[:1])
        return self.aplicar(
            "orden de palabras alterado",
            "MODERATE",
            RAZON_SOCIAL=f"{nuevo} {self.ent.sufijo}".strip(),
        )

    def sigla(self) -> Variante:
        sigla = _sigla(self.ent.nucleo)
        return self.aplicar(
            f"sigla en lugar del nombre ({sigla})",
            "SEVERE",
            RAZON_SOCIAL=f"{sigla} {self.ent.sufijo}".strip(),
        )

    def truncado(self, largo: int) -> Variante:
        completo = self.ent.razon_social
        corte = min(largo, len(completo) - 4)
        return self.aplicar(
            f"nombre truncado a {corte} caracteres",
            "MODERATE",
            RAZON_SOCIAL=completo[:corte].rstrip(),
        )

    def mojibake(self) -> Variante:
        return self.aplicar(
            "mojibake: eñe y tildes rotas (UTF-8 leído como cp1252)",
            "MODERATE",
            RAZON_SOCIAL=_mojibake(self.ent.razon_social),
        )

    def caja(self) -> Variante:
        return self.aplicar(
            "caja: nombre en tipo título", "LIGHT", RAZON_SOCIAL=self.ent.razon_social.title()
        )

    def espacios(self) -> Variante:
        tokens = self.ent.razon_social.split()
        nuevo = " " + "  ".join(tokens[:2]) + " " + " ".join(tokens[2:]) + " "
        return self.aplicar(
            "espacios dobles y de borde", "LIGHT", RAZON_SOCIAL=nuevo.rstrip() + " "
        )

    def nombre_comercial_como_razon(self) -> Variante:
        return self.aplicar(
            "nombre comercial capturado como razón social",
            "SEVERE",
            RAZON_SOCIAL=self.ent.nombre_comercial,
            NOMBRE_COMERCIAL="",
        )

    def orden_persona(self) -> Variante:
        apellidos, nombres = self.gen.partes_persona(self.ent)
        return self.aplicar(
            "orden nombres-apellidos (el RUES escribe apellidos primero)",
            "MODERATE",
            RAZON_SOCIAL=f"{nombres} {apellidos}",
        )

    def sucursal(self, lugar: tuple[str, str, tuple[str, ...], str], forma: str) -> Variante:
        ciudad = lugar[1]
        nombre = {
            "SUCURSAL": f"{self.ent.razon_social} SUCURSAL {ciudad}",
            "GUION": f"{self.ent.razon_social} - {ciudad}",
            "AGENCIA": f"{self.ent.razon_social} AGENCIA {ciudad}",
        }[forma]
        return self.aplicar(
            f"sucursal: nombre con la ciudad ({ciudad}) y municipio distinto",
            "MODERATE",
            RAZON_SOCIAL=nombre,
            DEPARTAMENTO=lugar[0],
            MUNICIPIO=lugar[1],
        )

    # -- perturbaciones de identificador y lugar ------------------------------

    def nit(self, forma: str) -> Variante:
        texto = _formato_nit(self.ent.identificador, self.ent.dv, forma)
        descripcion = {
            "dv_pegado": "NIT con el dígito de verificación pegado",
            "puntos_guion_dv": "NIT con puntos, guion y DV",
            "puntos": "NIT con puntos",
            "guiones": "NIT con guiones",
            "guion_dv": "NIT con guion y DV",
            "ceros": "NIT con ceros a la izquierda",
            "espacios": "NIT con espacios de borde",
        }[forma]
        return self.aplicar(descripcion, "LIGHT", NIT=texto, DV="")

    def sin_identificador(self) -> Variante:
        return self.aplicar(
            "sin identificador en esta fuente",
            "MODERATE",
            NIT="",
            DV="",
            TIPO_IDENTIFICACION="VACIO",
        )

    def municipio_variante(self) -> Variante:
        variantes = self.ent.lugar[2]
        forma = variantes[self.gen.rng.randrange(len(variantes))]
        return self.aplicar(f"municipio escrito como «{forma}»", "LIGHT", MUNICIPIO=forma)

    # -- cierre ----------------------------------------------------------------

    def registro(self, caso: str) -> Registro:
        nivel = max((NIVELES.index(n) for _, n in self.aplicadas), default=0)
        variacion = "; ".join(d for d, _ in self.aplicadas) or SIN_CAMBIOS
        if nivel > 0:
            canonico = (self.ent.identificador, self.ent.razon_social, self.ent.lugar[1])
            actual = (self.campos["NIT"], self.campos["RAZON_SOCIAL"], self.campos["MUNICIPIO"])
            if actual == canonico:
                raise AssertionError(
                    f"la variación «{variacion}» no cambió nada en {self.ent.razon_social}"
                )
        return Registro(
            fuente=self.fuente,
            campos=dict(self.campos),
            caso=caso,
            nivel_ruido=NIVELES[nivel],
            variacion=variacion,
            tipo_entidad=self.ent.tipo,
        )


# ---------------------------------------------------------------------------
# Generador: convenciones por fuente y recetas de caso
# ---------------------------------------------------------------------------


class Generador:
    """Produce los bloques de todos los casos a partir de una configuración."""

    def __init__(self, config: ConfiguracionPlantilla) -> None:
        self.config = config
        self.rng = random.Random(config.semilla)
        self.fabrica = Fabrica(self.rng)
        self.bloques: list[Bloque] = []

    # -- convenciones por fuente ----------------------------------------------

    def fuentes(self, n: int) -> list[str]:
        """``n`` fuentes distintas, en el orden canónico CRM, RUES, ADUANAS."""
        elegidas = self.rng.sample(FUENTES, n)
        return sorted(elegidas, key=FUENTES.index)

    def fecha(self) -> str:
        return (
            FECHA_INICIAL + dt.timedelta(days=self.rng.randint(0, DIAS_RANGO_FECHAS))
        ).isoformat()

    def direccion_formateada(self, d: tuple[str, int, int, int, str], fuente: str) -> str:
        via, n1, n2, n3, comp = d
        if fuente == "CRM":
            base = f"{ABREV_VIA_CRM[via]} {n1} # {n2}-{n3}"
        elif fuente == "RUES":
            base = f"{via} {n1} No. {n2} - {n3}"
        else:
            base = f"{ABREV_VIA_ADUANAS[via]} {n1} {n2} {n3}"
        return f"{base} {comp}".strip()

    def telefono_formateado(self, t: str, fuente: str) -> str:
        if fuente == "CRM":
            return (
                f"({t[:3]}) {t[3:6]} {t[6:]}" if t.startswith("6") else f"{t[:3]} {t[3:6]} {t[6:]}"
            )
        if fuente == "RUES":
            return t
        return "57" + t

    def partes_persona(self, ent: Entidad) -> tuple[str, str]:
        tokens = ent.nucleo.split()
        return " ".join(tokens[:2]), " ".join(tokens[2:])

    def campos_base(self, ent: Entidad, fuente: str) -> dict[str, str]:
        """El registro canónico de ``ent`` tal como lo escribe cada fuente.

        La cobertura de las columnas opcionales cambia por fuente (el RUES trae CIIU y
        dirección; aduanas casi nunca trae correo); eso es cobertura, no ruido, y no
        cuenta en NIVEL_RUIDO. El formato del teléfono y la dirección también es el de
        cada fuente.
        """
        rng = self.rng
        cobertura = {
            "CRM": {"NC": 0.7, "DEP": 0.6, "MUN": 0.9, "DIR": 0.5, "TEL": 0.9, "EMAIL": 0.85, "WEB": 0.6, "CIIU": 0.15, "DV": 0.6},
            "RUES": {"NC": 1.0, "DEP": 1.0, "MUN": 1.0, "DIR": 0.95, "TEL": 0.7, "EMAIL": 0.6, "WEB": 0.2, "CIIU": 1.0, "DV": 1.0},
            "ADUANAS": {"NC": 0.0, "DEP": 0.4, "MUN": 0.75, "DIR": 0.6, "TEL": 0.3, "EMAIL": 0.1, "WEB": 0.0, "CIIU": 0.0, "DV": 0.0},
        }[fuente]  # fmt: skip

        def si(clave: str) -> bool:
            return rng.random() < cobertura[clave]

        ciiu = ";".join(ent.ciiu) if fuente == "RUES" else ent.ciiu[0]
        return {
            "NIT": ent.identificador,
            "DV": ent.dv if (ent.dv and si("DV")) else "",
            "TIPO_IDENTIFICACION": ent.tipo_identificacion,
            "RAZON_SOCIAL": ent.razon_social,
            "NOMBRE_COMERCIAL": ent.nombre_comercial if (ent.nombre_comercial and si("NC")) else "",
            "DEPARTAMENTO": ent.lugar[0] if si("DEP") else "",
            "MUNICIPIO": ent.lugar[1] if si("MUN") else "",
            "DIRECCION": self.direccion_formateada(ent.direccion, fuente) if si("DIR") else "",
            "TELEFONO": self.telefono_formateado(ent.telefono, fuente) if si("TEL") else "",
            "EMAIL": ent.email if si("EMAIL") else "",
            "SITIO_WEB": ent.sitio_web if (ent.sitio_web and si("WEB")) else "",
            "CIIU": ciiu if si("CIIU") else "",
            "PAIS": ent.pais,
            "FECHA_ACTUALIZACION": self.fecha(),
        }

    def errata(self, nucleo: str, n: int) -> tuple[str, str]:
        """``n`` erratas en tokens distintos del núcleo; devuelve (texto, detalle)."""
        tokens = nucleo.split()
        candidatos = [i for i, t in enumerate(tokens) if len(t) >= 5 and t not in STOPWORDS_SIGLA]
        posiciones = self.rng.sample(candidatos, min(n, len(candidatos)))
        detalles = []
        for i in posiciones:
            t = tokens[i]
            op = self.rng.choice(("sustituir", "omitir", "duplicar", "transponer"))
            p = self.rng.randint(1, len(t) - 2)
            if op == "sustituir":
                letra = self.rng.choice([c for c in "AEIOURSNLT" if c != t[p]])
                nuevo = t[:p] + letra + t[p + 1 :]
            elif op == "omitir":
                nuevo = t[:p] + t[p + 1 :]
            elif op == "duplicar":
                nuevo = t[:p] + t[p] + t[p:]
            else:
                if t[p] == t[p + 1]:
                    p -= 1
                nuevo = t[:p] + t[p + 1] + t[p] + t[p + 2 :]
            if nuevo == t:
                nuevo = t[:p] + ("X" if t[p] != "X" else "Z") + t[p + 1 :]
            tokens[i] = nuevo
            detalles.append(f"{t} → {nuevo}")
        return " ".join(tokens), ", ".join(detalles)

    # -- recetas ---------------------------------------------------------------

    def _par(
        self,
        a: Registro,
        b: Registro,
        is_match: str,
        banda: str,
        comentario: str,
        *,
        dificil: bool = False,
    ) -> ParPropuesto:
        return ParPropuesto(a, b, is_match, banda, comentario, dificil)

    def _grupo_variantes(
        self, ent: Entidad, fuentes: list[str], caso: str, planes: list[list[tuple[str, tuple]]]
    ) -> list[Registro]:
        """Un registro por fuente; ``planes[i]`` = lista de (método de Variante, args)."""
        registros = []
        for fuente, plan in zip(fuentes, planes, strict=True):
            v = Variante(self, ent, fuente)
            for metodo, args in plan:
                getattr(v, metodo)(*args)
            registros.append(v.registro(caso))
        return registros

    def receta_nit_igual(self, n: int) -> Bloque:
        bloque = Bloque("positivo_nit_igual", [])
        for i in range(n):
            ent = self.fabrica.empresa()
            fuentes = self.fuentes(self.rng.choices((2, 3), weights=(65, 35))[0])
            planes: list[list[tuple[str, tuple]]] = []
            for _ in fuentes:
                plan: list[tuple[str, tuple]] = []
                if self.rng.random() < 0.18:
                    plan.append(("municipio_variante", ()))
                elif self.rng.random() < 0.08:
                    plan.append(("espacios", ()))
                planes.append(plan)
            registros = self._grupo_variantes(ent, fuentes, bloque.caso, planes)
            if self.rng.random() < 0.08:  # duplicado interno del CRM, con espacios de más
                v = Variante(self, ent, "CRM").espacios()
                v.aplicadas.append(
                    ("duplicado dentro de la misma fuente (segunda fila del CRM)", "LIGHT")
                )
                registros.append(v.registro(bloque.caso))
            bloque.grupos.append(registros)
            if i % int(1 / self.config.fraccion_pares_nit_igual) == 0:
                bloque.pares.append(
                    self._par(
                        registros[0], registros[1], "1", "unido_holgado", "mismo NIT y mismo nombre"
                    )
                )
        return bloque

    def receta_nit_con_y_sin_dv(self) -> Bloque:
        bloque = Bloque("positivo_nit_con_y_sin_dv", [])
        for _ in range(3):
            ent = self.fabrica.empresa()
            regs = self._grupo_variantes(
                ent,
                ["CRM", "RUES", "ADUANAS"],
                bloque.caso,
                [[], [("nit", ("dv_pegado",))], [("nit", ("puntos_guion_dv",))]],
            )
            regs[0].campos["DV"] = ent.dv  # el CRM trae el DV en su columna
            bloque.grupos.append(regs)
            bloque.pares.append(
                self._par(
                    regs[0], regs[1], "1", "unido_holgado", "mismo NIT: sin DV y con DV pegado"
                )
            )
            bloque.pares.append(
                self._par(
                    regs[0], regs[2], "1", "unido_holgado", "mismo NIT: con puntos, guion y DV"
                )
            )
        return bloque

    def receta_ceros_o_puntuacion(self) -> Bloque:
        bloque = Bloque("positivo_nit_con_ceros_o_puntuacion", [])
        formas = (("ceros", "puntos"), ("espacios", "guiones"), ("ceros", "guion_dv"))
        for f_crm, f_adu in formas:
            ent = self.fabrica.empresa()
            regs = self._grupo_variantes(
                ent,
                ["CRM", "RUES", "ADUANAS"],
                bloque.caso,
                [[("nit", (f_crm,))], [], [("nit", (f_adu,))]],
            )
            bloque.grupos.append(regs)
            bloque.pares.append(
                self._par(regs[0], regs[1], "1", "unido_holgado", f"NIT {f_crm} vs limpio")
            )
            bloque.pares.append(
                self._par(regs[1], regs[2], "1", "unido_holgado", f"NIT limpio vs {f_adu}")
            )
        return bloque

    def receta_sufijo_legal(self) -> Bloque:
        bloque = Bloque("positivo_sufijo_legal", [])
        planes = (
            ("SAS", ("S.A.S.", "S A S")),
            ("LTDA", ("LIMITADA", "LTDA.")),
            ("SA", ("SA", "S A")),
            ("LTDA", ("SAS", None)),  # transformación de forma legal, mismo NIT
        )
        for familia, (forma_b, forma_c) in planes:
            ent = self.fabrica.empresa(familia_sufijo=familia)
            fuentes = ["CRM", "RUES", "ADUANAS"] if forma_c else self.fuentes(2)
            plan_b: list[tuple[str, tuple]] = [("sufijo", (forma_b,))]
            if familia == "LTDA" and forma_b == "SAS":
                plan_b = [("sufijo", (forma_b, "cambio de forma legal: LTDA → SAS (mismo NIT)"))]
            planes_v: list[list[tuple[str, tuple]]] = [[], plan_b]
            if forma_c:
                planes_v.append([("sufijo", (forma_c,))])
            regs = self._grupo_variantes(ent, fuentes, bloque.caso, planes_v)
            bloque.grupos.append(regs)
            bloque.pares.append(
                self._par(
                    regs[0], regs[1], "1", "unido_holgado", f"sufijo {ent.sufijo} vs {forma_b}"
                )
            )
        return bloque

    def receta_errata(self, n: int) -> Bloque:
        caso = "positivo_errata_1_letra" if n == 1 else "positivo_errata_2_letras"
        bloque = Bloque(caso, [])
        for i in range(4):
            con_nit = i < 3  # el cuarto grupo no tiene NIT: el nombre es toda la evidencia
            ent = self.fabrica.empresa(con_nit=con_nit, minimo_tokens=2)
            fuentes = self.fuentes(2)
            regs = self._grupo_variantes(ent, fuentes, caso, [[], [("errata", (n,))]])
            bloque.grupos.append(regs)
            banda = "unido_holgado" if con_nit else "unido_borde"
            bloque.pares.append(
                self._par(
                    regs[0],
                    regs[1],
                    "1",
                    banda,
                    regs[1].variacion,
                    dificil=(not con_nit and n == 2),
                )
            )
        return bloque

    def receta_orden(self) -> Bloque:
        bloque = Bloque("positivo_orden_de_palabras", [])
        for i in range(3):
            con_nit = i < 2
            ent = self.fabrica.empresa(con_nit=con_nit)
            regs = self._grupo_variantes(ent, self.fuentes(2), bloque.caso, [[], [("orden", ())]])
            bloque.grupos.append(regs)
            bloque.pares.append(
                self._par(
                    regs[0],
                    regs[1],
                    "1",
                    "unido_holgado" if con_nit else "unido_borde",
                    "mismas palabras en otro orden",
                )
            )
        return bloque

    def _nucleo_largo(self, *, exclusiva: bool) -> str:
        for _ in range(100):
            r = self.fabrica.raiz(exclusiva=exclusiva)
            s1, s2 = self.rng.sample(SECTORES_SUFIJO, 2)
            nucleo = f"{r} {s1} Y {s2}"
            if self.fabrica.reservar_nombre(nucleo):
                return nucleo
        raise RuntimeError("no se pudo inventar un nombre largo")

    def receta_sigla(self, con_nit: bool) -> Bloque:
        caso = "positivo_sigla" if con_nit else "positivo_sigla_sin_nit"
        bloque = Bloque(caso, [])
        for _ in range(3 if con_nit else 2):
            ent = self.fabrica.empresa(
                con_nit=con_nit,
                nucleo=self._nucleo_largo(exclusiva=not con_nit),
                familia_sufijo="SAS",
            )
            regs = self._grupo_variantes(ent, self.fuentes(2), caso, [[], [("sigla", ())]])
            bloque.grupos.append(regs)
            if con_nit:
                bloque.pares.append(
                    self._par(
                        regs[0], regs[1], "1", "unido_holgado", "sigla vs nombre largo, mismo NIT"
                    )
                )
            else:
                bloque.pares.append(
                    self._par(
                        regs[0],
                        regs[1],
                        "1",
                        "no_unido_borde",
                        "sigla vs nombre largo sin NIT: se confirmó por fuera (RUES) que es la misma",
                        dificil=True,
                    )
                )
        return bloque

    def receta_truncado(self) -> Bloque:
        bloque = Bloque("positivo_truncado", [])
        for i, largo in enumerate((20, 25, 30)):
            con_nit = i < 2
            ent = self.fabrica.empresa(con_nit=con_nit, minimo_tokens=3)
            regs = self._grupo_variantes(
                ent, self.fuentes(2), bloque.caso, [[], [("truncado", (largo,))]]
            )
            bloque.grupos.append(regs)
            bloque.pares.append(
                self._par(
                    regs[0],
                    regs[1],
                    "1",
                    "unido_holgado" if con_nit else "unido_borde",
                    regs[1].variacion,
                )
            )
        return bloque

    def receta_nombre_comercial(self) -> Bloque:
        bloque = Bloque("positivo_nombre_comercial_vs_razon_social", [])
        comerciales = list(NOMBRES_COMERCIALES[:3])
        for nombre_comercial in comerciales:
            ent = self.fabrica.empresa(nombre_comercial=nombre_comercial)
            regs = self._grupo_variantes(
                ent, ["CRM", "RUES"], bloque.caso, [[("nombre_comercial_como_razon", ())], []]
            )
            regs[1].campos["NOMBRE_COMERCIAL"] = nombre_comercial
            bloque.grupos.append(regs)
            bloque.pares.append(
                self._par(
                    regs[0],
                    regs[1],
                    "1",
                    "unido_holgado",
                    "nombre comercial en el CRM vs razón social en el RUES, mismo NIT",
                )
            )
        return bloque

    def receta_sin_nit_en_una_fuente(self) -> Bloque:
        bloque = Bloque("positivo_sin_nit_en_una_fuente", [])
        for i in range(8):
            ent = self.fabrica.empresa(raiz_exclusiva=True)
            fuentes = self.fuentes(2)
            plan_a: list[tuple[str, tuple]] = [("sin_identificador", ())]
            if i % 2 == 1:
                plan_a.append(("sufijo", (FAMILIAS_SUFIJO[self._familia(ent.sufijo)][1],)))
            regs = self._grupo_variantes(ent, fuentes, bloque.caso, [plan_a, []])
            bloque.grupos.append(regs)
            bloque.pares.append(
                self._par(
                    regs[0],
                    regs[1],
                    "1",
                    "unido_borde",
                    "mismo nombre; NIT solo en una fuente (par mixto)",
                )
            )
        return bloque

    def _familia(self, sufijo: str) -> str:
        return next(f for f, formas in FAMILIAS_SUFIJO.items() if sufijo in formas)

    def receta_sin_nit_en_ninguna(self) -> Bloque:
        bloque = Bloque("positivo_sin_nit_en_ninguna", [])
        for i in range(8):
            ent = self.fabrica.empresa(con_nit=False)
            fuentes = self.fuentes(2 if i % 4 != 3 else 3)
            planes: list[list[tuple[str, tuple]]] = [[] for _ in fuentes]
            if i % 4 == 1:
                planes[1] = [("sufijo", (FAMILIAS_SUFIJO[self._familia(ent.sufijo)][1],))]
            if i % 4 == 2:
                planes[1] = [("municipio_variante", ())]
            regs = self._grupo_variantes(ent, fuentes, bloque.caso, planes)
            bloque.grupos.append(regs)
            bloque.pares.append(
                self._par(
                    regs[0],
                    regs[1],
                    "1",
                    "unido_borde",
                    "sin NIT en ninguna fuente: mismo nombre y municipio",
                )
            )
        return bloque

    def receta_sucursal(self) -> Bloque:
        bloque = Bloque("positivo_sucursal", [])
        for forma in ("SUCURSAL", "GUION", "AGENCIA"):
            ent = self.fabrica.empresa()
            otros = [lg for lg in LUGARES if lg[1] != ent.lugar[1]]
            lugar_b = otros[self.rng.randrange(len(otros))]
            regs = self._grupo_variantes(
                ent, ["CRM", "RUES"], bloque.caso, [[("sucursal", (lugar_b, forma))], []]
            )
            bloque.grupos.append(regs)
            bloque.pares.append(
                self._par(
                    regs[0],
                    regs[1],
                    "1",
                    "unido_holgado",
                    "mismo NIT; sucursal con ciudad en el nombre",
                )
            )
        return bloque

    def receta_persona_natural(self) -> Bloque:
        bloque = Bloque("positivo_persona_natural", [])
        for i in range(4):
            ent = self.fabrica.persona(nombre_comercial=NOMBRES_COMERCIALES[3 + i] if i < 2 else "")
            plan_crm: list[tuple[str, tuple]] = [("orden_persona", ())]
            if i == 3:
                plan_crm.append(("sin_identificador", ()))
            regs = self._grupo_variantes(ent, ["CRM", "RUES"], bloque.caso, [plan_crm, []])
            bloque.grupos.append(regs)
            bloque.pares.append(
                self._par(
                    regs[0],
                    regs[1],
                    "1",
                    "unido_holgado" if i < 3 else "unido_borde",
                    "misma persona natural; orden de nombres y apellidos distinto",
                )
            )
        return bloque

    def receta_mojibake(self) -> Bloque:
        bloque = Bloque("positivo_mojibake", [])
        nucleos = ("MUEBLES PEÑALOZA", "CAFÉ DE LA MONTAÑA", "DISEÑOS LONDOÑO Y QUIÑONES")
        for i, nucleo in enumerate(nucleos):
            ent = self.fabrica.empresa(
                con_nit=i < 2, nucleo=nucleo, familia_sufijo=("SAS", "LTDA", "SENC")[i]
            )
            regs = self._grupo_variantes(
                ent, ["RUES", "ADUANAS"], bloque.caso, [[], [("mojibake", ())]]
            )
            bloque.grupos.append(regs)
            bloque.pares.append(
                self._par(
                    regs[0],
                    regs[1],
                    "1",
                    "unido_holgado" if i < 2 else "unido_borde",
                    "eñe y tildes rotas en aduanas",
                )
            )
        return bloque

    # -- negativos -------------------------------------------------------------

    def _registros_entidad(self, ent: Entidad, caso: str, n_fuentes: int) -> list[Registro]:
        return self._grupo_variantes(
            ent, self.fuentes(n_fuentes), caso, [[] for _ in range(n_fuentes)]
        )

    def receta_homonimo_persona(self) -> Bloque:
        bloque = Bloque("negativo_homonimo_persona", [])
        for i in range(3):
            nombres = self.rng.choice(NOMBRES_PERSONA)
            apellidos = " ".join(self.rng.sample(APELLIDOS, 2))
            lugares = self.rng.sample(LUGARES, 2)
            a = self.fabrica.persona(nombres=nombres, apellidos=apellidos, lugar=lugares[0])
            b = self.fabrica.persona(nombres=nombres, apellidos=apellidos, lugar=lugares[1])
            ra = self._registros_entidad(a, bloque.caso, 2 if i < 2 else 1)
            rb = self._registros_entidad(b, bloque.caso, 1 if i < 2 else 2)
            if i == 2:
                rb[0] = Variante(self, b, rb[0].fuente).orden_persona().registro(bloque.caso)
            bloque.grupos.extend([ra, rb])
            bloque.pares.append(
                self._par(
                    ra[0],
                    rb[0],
                    "0",
                    "no_unido_borde",
                    "mismo nombre completo, cédulas y ciudades distintas: dos personas",
                    dificil=(i == 1),
                )
            )
        return bloque

    def receta_nombre_identico(self) -> Bloque:
        bloque = Bloque("negativo_nombre_identico_nit_distinto", [])
        for i in range(3):
            nucleo = self.fabrica.nucleo_empresa()
            lugares = self.rng.sample(LUGARES, 2)
            a = self.fabrica.empresa(nucleo=nucleo, familia_sufijo="SAS", lugar=lugares[0])
            b = self.fabrica.empresa(nucleo=nucleo, familia_sufijo="SAS", lugar=lugares[1])
            ra = self._registros_entidad(a, bloque.caso, 2 if i != 1 else 1)
            rb = self._registros_entidad(b, bloque.caso, 1 if i != 1 else 2)
            bloque.grupos.extend([ra, rb])
            bloque.pares.append(
                self._par(
                    ra[0],
                    rb[0],
                    "0",
                    "no_unido_borde",
                    "razón social idéntica, NIT válidos distintos (ciudades distintas)",
                    dificil=(i == 0),
                )
            )
        return bloque

    def receta_nombre_generico(self) -> Bloque:
        bloque = Bloque("negativo_nombre_generico", [])
        pares = (
            ("INVERSIONES", "SAS", "S.A.S."),
            ("COMERCIALIZADORA", "LTDA", "LIMITADA"),
            ("DISTRIBUIDORA DEL NORTE", "SAS", "LTDA"),
        )
        for nucleo, suf_a, suf_b in pares:
            lugares = self.rng.sample(LUGARES, 2)
            a = self.fabrica.empresa(nucleo=nucleo, sufijo=suf_a, lugar=lugares[0])
            b = self.fabrica.empresa(nucleo=nucleo, sufijo=suf_b, lugar=lugares[1])
            ra = self._registros_entidad(a, bloque.caso, 1)
            rb = self._registros_entidad(b, bloque.caso, 2)
            bloque.grupos.extend([ra, rb])
            bloque.pares.append(
                self._par(ra[0], rb[0], "0", "no_unido_borde", "nombre genérico; NIT distintos")
            )
        return bloque

    def receta_filial_matriz(self) -> Bloque:
        bloque = Bloque("negativo_filial_matriz", [])
        familias = (
            ("KRONOS HOLDING", "KRONOS COLOMBIA", "SAS", "SAS"),
            ("VALDORIA GRUPO EMPRESARIAL", "VALDORIA SERVICIOS", "SA", "SAS"),
            ("TRIVENTA INTERNACIONAL", "TRIVENTA LOGISTICA", "SA", "SAS"),
        )
        for nuc_a, nuc_b, fam_a, fam_b in familias:
            lugar = self.fabrica.lugar()
            a = self.fabrica.empresa(nucleo=nuc_a, familia_sufijo=fam_a, lugar=lugar)
            b = self.fabrica.empresa(nucleo=nuc_b, familia_sufijo=fam_b, lugar=lugar)
            ra = self._registros_entidad(a, bloque.caso, 2)
            rb = self._registros_entidad(b, bloque.caso, 1)
            bloque.grupos.extend([ra, rb])
            bloque.pares.append(
                self._par(
                    ra[0],
                    rb[0],
                    "0",
                    "no_unido_borde",
                    "matriz y filial: misma raíz, NIT distintos, personas jurídicas distintas",
                    dificil=True,
                )
            )
        return bloque

    def receta_nit_corrupto(self) -> Bloque:
        bloque = Bloque("negativo_nit_corrupto_un_digito", [])
        for i in range(3):
            a = self.fabrica.empresa()
            if i == 0:
                nucleo_b = self.fabrica.nucleo_empresa()
            elif i == 1:
                sector = (
                    next((t for t in a.nucleo.split() if t in SECTORES_SUFIJO), None) or "LOGISTICA"
                )
                nucleo_b = self.fabrica.nucleo_empresa(raiz=self.fabrica.raiz()).replace(
                    a.nucleo.split()[-1], sector
                )
                if not self.fabrica.reservar_nombre(nucleo_b):
                    nucleo_b = f"{nucleo_b} {self.rng.choice(REGIONES)}"
            else:
                nucleo_b = f"{a.nucleo} {self.rng.choice(REGIONES)}"
                self.fabrica.reservar_nombre(nucleo_b)
            b = self.fabrica.empresa(nucleo=nucleo_b, nit=self.fabrica.nit_cercano(a.identificador))
            ra = self._registros_entidad(a, bloque.caso, 2 if i != 2 else 1)
            rb = self._registros_entidad(b, bloque.caso, 1 if i != 2 else 2)
            bloque.grupos.extend([ra, rb])
            bloque.pares.append(
                self._par(
                    ra[0],
                    rb[0],
                    "0",
                    "no_unido_borde",
                    f"NIT a un dígito ({a.identificador} vs {b.identificador}); empresas distintas",
                    dificil=(i == 2),
                )
            )
        return bloque

    def receta_raiz_confundible(self) -> Bloque:
        bloque = Bloque("negativo_raiz_confundible", [])
        pares = (
            ("ZELVORA", "ZELVORO", True, True),
            ("KALMIRA ALIMENTOS", "KALMIRO ALIMENTOS", False, False),
            ("TRUVEL ENERGIA", "TRUVELA ENERGIA", True, False),
        )
        for nuc_a, nuc_b, nit_a, nit_b in pares:
            lugar = self.fabrica.lugar()
            a = self.fabrica.empresa(nucleo=nuc_a, familia_sufijo="SAS", con_nit=nit_a, lugar=lugar)
            b = self.fabrica.empresa(nucleo=nuc_b, familia_sufijo="SAS", con_nit=nit_b, lugar=lugar)
            ra = self._registros_entidad(a, bloque.caso, 2)
            rb = self._registros_entidad(b, bloque.caso, 1)
            bloque.grupos.extend([ra, rb])
            bloque.pares.append(
                self._par(
                    ra[0],
                    rb[0],
                    "0",
                    "no_unido_borde",
                    f"raíces casi iguales ({nuc_a.split()[0]} vs {nuc_b.split()[0]}): empresas distintas",
                    dificil=True,
                )
            )
        return bloque

    def receta_singletons(self, n: int) -> Bloque:
        bloque = Bloque("singleton", [])
        for i in range(n):
            if i % 15 == 14:
                pais, sufijo = PAISES_EXTRANJEROS[(i // 15) % len(PAISES_EXTRANJEROS)]
                ent = self.fabrica.empresa(
                    con_nit=False, sufijo=sufijo, pais=pais, lugar=LUGARES[0]
                )
                ent.lugar = (pais, "", (), "601")
                regs = self._grupo_variantes(ent, ["ADUANAS"], bloque.caso, [[]])
                regs[0].campos.update(
                    {"DEPARTAMENTO": "", "MUNICIPIO": "", "DIRECCION": "", "TELEFONO": ""}
                )
            elif i % 15 == 13:
                ent = self.fabrica.persona(extranjera=(i % 30 == 28))
                regs = self._grupo_variantes(ent, self.fuentes(1), bloque.caso, [[]])
            elif i % 4 == 3:
                ent = self.fabrica.empresa(con_nit=False)
                regs = self._grupo_variantes(ent, self.fuentes(1), bloque.caso, [[]])
            else:
                ent = self.fabrica.empresa()
                regs = self._grupo_variantes(ent, self.fuentes(1), bloque.caso, [[]])
            bloque.grupos.append(regs)
        return bloque

    def receta_ambiguos(self) -> Bloque:
        """Tres pares que una persona no puede decidir con los datos: IS_MATCH = UNCLEAR.

        En VERDAD_GRUPOS cada registro es un grupo propio (CASO ambiguo_revisar); los pares
        UNCLEAR están exentos de la consistencia con los grupos porque no afirman nada.
        """
        bloque = Bloque("ambiguo_revisar", [])
        lugar = LUGARES[0]
        a1 = self.fabrica.empresa(
            con_nit=False, nucleo="COMERCIALIZADORA ORBITAL", sufijo="", lugar=lugar
        )
        b1 = self.fabrica.empresa(
            nucleo="ORBITAL COMERCIALIZADORA DEL CARIBE", familia_sufijo="SAS", lugar=LUGARES[3]
        )
        a2 = self.fabrica.persona(
            con_cedula=False, nombres="JUAN CARLOS", apellidos="PEREZ GOMEZ", lugar=lugar
        )
        b2 = self.fabrica.persona(nombres="JUAN CARLOS", apellidos="PEREZ GOMEZ", lugar=LUGARES[1])
        a3 = self.fabrica.empresa(
            con_nit=False, nucleo="INVERSIONES VALDORIA", sufijo="", lugar=LUGARES[2]
        )
        b3 = self.fabrica.empresa(
            nucleo="VALDORIA INVERSIONES Y CONSTRUCCIONES", familia_sufijo="SAS", lugar=LUGARES[2]
        )
        ra1 = self._grupo_variantes(a1, ["CRM"], bloque.caso, [[]])
        rb1 = self._grupo_variantes(b1, ["ADUANAS"], bloque.caso, [[]])
        ra2 = self._grupo_variantes(a2, ["CRM"], bloque.caso, [[("orden_persona", ())]])
        rb2 = self._grupo_variantes(b2, ["RUES"], bloque.caso, [[]])
        ra3 = self._grupo_variantes(a3, ["CRM"], bloque.caso, [[]])
        rb3 = self._grupo_variantes(b3, ["RUES"], bloque.caso, [[]])
        bloque.grupos.extend([ra1, rb1, ra2, rb2, ra3, rb3])
        bloque.pares.extend(
            [
                self._par(
                    ra1[0],
                    rb1[0],
                    "UNCLEAR",
                    "no_unido_borde",
                    "sin NIT en el CRM y ciudad distinta: no se puede decidir con los datos",
                ),
                self._par(
                    ra2[0],
                    rb2[0],
                    "UNCLEAR",
                    "no_unido_borde",
                    "nombre de persona muy común sin cédula en el CRM: no se puede decidir",
                ),
                self._par(
                    ra3[0],
                    rb3[0],
                    "UNCLEAR",
                    "no_unido_borde",
                    "nombre parcial sin NIT; podría ser la misma o un homónimo",
                ),
            ]
        )
        return bloque

    # -- orquestación ---------------------------------------------------------

    def generar_bloques(self) -> list[Bloque]:
        c = self.config
        self.bloques = [
            self.receta_nit_igual(c.grupos_nit_igual),
            self.receta_nit_con_y_sin_dv(),
            self.receta_ceros_o_puntuacion(),
            self.receta_sufijo_legal(),
            self.receta_errata(1),
            self.receta_errata(2),
            self.receta_orden(),
            self.receta_sigla(True),
            self.receta_sigla(False),
            self.receta_truncado(),
            self.receta_nombre_comercial(),
            self.receta_sin_nit_en_una_fuente(),
            self.receta_sin_nit_en_ninguna(),
            self.receta_sucursal(),
            self.receta_persona_natural(),
            self.receta_mojibake(),
            self.receta_homonimo_persona(),
            self.receta_nombre_identico(),
            self.receta_nombre_generico(),
            self.receta_filial_matriz(),
            self.receta_nit_corrupto(),
            self.receta_raiz_confundible(),
            self.receta_singletons(c.singletons),
            self.receta_ambiguos(),
        ]
        return self.bloques


# ---------------------------------------------------------------------------
# Ensamblaje: ids, grupos, pares y hojas
# ---------------------------------------------------------------------------


@dataclass
class Plantilla:
    """Las hojas ya armadas, listas para escribir."""

    fuentes: dict[str, pd.DataFrame]
    verdad: pd.DataFrame
    pares: pd.DataFrame
    pares_anotador_a: pd.DataFrame
    pares_anotador_b: pd.DataFrame
    diccionario: pd.DataFrame
    casos: pd.DataFrame
    leeme: pd.DataFrame
    banco: pd.DataFrame

    def hojas(self) -> dict[str, pd.DataFrame]:
        return {
            "LEEME": self.leeme,
            **{f"FUENTE_{f}": self.fuentes[f] for f in FUENTES},
            "VERDAD_GRUPOS": self.verdad,
            "ETIQUETAS_PARES": self.pares,
            "DICCIONARIO": self.diccionario,
            "CASOS": self.casos,
        }


def _clave_par(a: str, b: str) -> tuple[str, str]:
    return (a, b) if a <= b else (b, a)


def ensamblar(
    config: ConfiguracionPlantilla, bloques: list[Bloque], rng: random.Random
) -> Plantilla:
    """Asigna ID_GRUPO e ID_REGISTRO_FUENTE, arma los pares y construye las hojas."""
    grupos: list[list[Registro]] = [g for b in bloques for g in b.grupos]
    orden = list(range(len(grupos)))
    rng.shuffle(orden)
    for nuevo, viejo in enumerate(orden):
        for r in grupos[viejo]:
            r.grupo = nuevo + 1

    registros_por_fuente: dict[str, list[Registro]] = {f: [] for f in FUENTES}
    for g in grupos:
        for r in g:
            registros_por_fuente[r.fuente].append(r)
    for f in FUENTES:
        rng.shuffle(registros_por_fuente[f])
        for i, r in enumerate(registros_por_fuente[f], start=1):
            r.id_registro = f"{f}-{i:06d}"
        n = len(registros_por_fuente[f])
        if not config.min_por_fuente <= n <= config.max_por_fuente:
            raise AssertionError(
                f"la fuente {f} tiene {n} registros; se exigen entre {config.min_por_fuente} y "
                f"{config.max_por_fuente}. Ajuste grupos_nit_igual o singletons."
            )

    fuentes = {
        f: pd.DataFrame(
            [{"ID_REGISTRO_FUENTE": r.id_registro, **r.campos} for r in registros_por_fuente[f]],
            columns=list(COLUMNAS_FUENTE),
        )
        for f in FUENTES
    }

    todos = sorted(
        (r for g in grupos for r in g),
        key=lambda r: (r.grupo, FUENTES.index(r.fuente), r.id_registro),
    )
    verdad = pd.DataFrame(
        [
            {
                "ID_REGISTRO_FUENTE": r.id_registro,
                "FUENTE": r.fuente,
                "ID_GRUPO": f"G{r.grupo:04d}",
                "TIPO_ENTIDAD": r.tipo_entidad,
                "REGIMEN": r.regimen,
                "CASO": r.caso,
                "NIVEL_RUIDO": r.nivel_ruido,
                "VARIACION": r.variacion,
            }
            for r in todos
        ],
        columns=list(COLUMNAS_VERDAD),
    )

    pares = _armar_pares(config, bloques, grupos, rng)
    pares_a, pares_b = _archivos_anotadores(pares)
    casos = _hoja_casos(bloques, todos)
    banco = construir_banco(fuentes, verdad)
    return Plantilla(
        fuentes=fuentes,
        verdad=verdad,
        pares=pares,
        pares_anotador_a=pares_a,
        pares_anotador_b=pares_b,
        diccionario=_hoja_diccionario(),
        casos=casos,
        leeme=_hoja_leeme(config, len(todos), len(grupos), len(pares)),
        banco=banco,
    )


def _armar_pares(
    config: ConfiguracionPlantilla,
    bloques: list[Bloque],
    grupos: list[list[Registro]],
    rng: random.Random,
) -> pd.DataFrame:
    propuestos = [p for b in bloques for p in b.pares]
    # Negativos de control, lejos del borde: registros de grupos distintos elegidos al azar.
    todos = [r for g in grupos for r in g]
    vistos: dict[tuple[str, str], int] = {
        _clave_par(p.a.id_registro, p.b.id_registro): 1 for p in propuestos
    }
    while sum(1 for p in propuestos if p.banda == "no_unido_lejos") < config.pares_control_lejanos:
        a, b = rng.sample(todos, 2)
        clave = _clave_par(a.id_registro, b.id_registro)
        if a.grupo == b.grupo or clave in vistos:
            continue
        vistos[clave] = 1
        propuestos.append(
            ParPropuesto(a, b, "0", "no_unido_lejos", "control: nombre y NIT distintos")
        )
    rng.shuffle(propuestos)
    filas = []
    turno = 0
    for p in propuestos:
        if p.dificil:
            anotador = "resuelto"
        else:
            anotador = "AB"[turno % 2]
            turno += 1
        filas.append(
            {
                "ID_1": p.a.id_registro,
                "ID_2": p.b.id_registro,
                "IS_MATCH": p.is_match,
                "ANOTADOR": anotador,
                "BANDA_ESPERADA": p.banda,
                "COMENTARIO": p.comentario,
            }
        )
    pares = pd.DataFrame(filas, columns=list(COLUMNAS_PARES))
    if not config.min_pares <= len(pares) <= config.max_pares:
        raise AssertionError(
            f"ETIQUETAS_PARES tiene {len(pares)} pares; se exigen entre {config.min_pares} y {config.max_pares}"
        )
    return pares


def _archivos_anotadores(pares: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Los dos archivos «a ciegas»: coinciden salvo en los pares que hubo que resolver."""
    a = pares[["ID_1", "ID_2", "IS_MATCH"]].copy()
    a["ANOTADOR"] = "A"
    b = a.copy()
    b["ANOTADOR"] = "B"
    resueltos = (pares["ANOTADOR"] == "resuelto").to_numpy()
    b.loc[resueltos, "IS_MATCH"] = (
        b.loc[resueltos, "IS_MATCH"].map({"1": "0", "0": "1"}).fillna("UNCLEAR")
    )
    return a, b


def _hoja_casos(bloques: list[Bloque], todos: list[Registro]) -> pd.DataFrame:
    conteo: dict[str, int] = {}
    for r in todos:
        conteo[r.caso] = conteo.get(r.caso, 0) + 1
    ejemplo: dict[str, str] = {}
    for b in bloques:
        if b.caso in ejemplo:
            continue
        if b.pares:
            p = b.pares[0]
            ejemplo[b.caso] = f"{p.a.id_registro}; {p.b.id_registro}"
        else:
            ejemplo[b.caso] = "; ".join(r.id_registro for r in b.grupos[0])
    filas = [
        {
            "CASO": caso,
            "DESCRIPCION": descripcion,
            "QUE_DEBE_HACER_EL_MOTOR": accion,
            "EJEMPLO_IDS": ejemplo.get(caso, ""),
            "N_REGISTROS": conteo.get(caso, 0),
        }
        for caso, descripcion, accion in CATALOGO_CASOS
    ]
    return pd.DataFrame(filas, columns=list(COLUMNAS_CASOS))


def construir_banco(
    fuentes: dict[str, pd.DataFrame], verdad: pd.DataFrame, *, estrato: str = "PLANTILLA"
) -> pd.DataFrame:
    """VERDAD_GRUPOS más las hojas FUENTE_* en el formato plano de ``scripts/banco.py``.

    Una fila por fila de VERDAD_GRUPOS, en su orden, con los campos del registro tomados de
    su hoja de fuente. Sirve igual para la plantilla y para una base real: las columnas que
    la base no tenga quedan vacías, REGIMEN se deduce del NIT si VERDAD_GRUPOS no lo trae y
    TIPO_ENTIDAD de TIPO_IDENTIFICACION (CC / CE = persona) si tampoco viene.

    Raises:
        ValueError: faltan las columnas mínimas o un registro de la verdad no está en su fuente.
    """
    minimas = ("ID_REGISTRO_FUENTE", "FUENTE", "ID_GRUPO")
    faltan = [c for c in minimas if c not in verdad.columns]
    if faltan:
        raise ValueError(f"A VERDAD_GRUPOS le faltan {faltan}; se necesitan {list(minimas)}.")
    registros = pd.concat(
        [df.assign(FUENTE=f) for f, df in fuentes.items()], ignore_index=True, sort=False
    )
    marco = verdad.merge(
        registros,
        on=["ID_REGISTRO_FUENTE", "FUENTE"],
        how="left",
        validate="one_to_one",
        indicator=True,
    )
    huerfanos = marco.loc[marco["_merge"] != "both", "ID_REGISTRO_FUENTE"].tolist()
    if huerfanos:
        raise ValueError(
            f"{len(huerfanos)} registros de VERDAD_GRUPOS no están en su hoja FUENTE_* "
            f"(mismo ID_REGISTRO_FUENTE y FUENTE): {huerfanos[:5]}"
        )

    def columna(nombre: str) -> pd.Series:
        if nombre in marco.columns:
            return marco[nombre].fillna("").astype(str)
        return pd.Series("", index=marco.index, dtype=str)

    nit = columna("NIT")
    if "REGIMEN" in verdad.columns:
        regimen = columna("REGIMEN")
    else:
        regimen = nit.str.strip().ne("").map({True: "CON_ID", False: "SIN_ID"})
    if "TIPO_ENTIDAD" in verdad.columns:
        tipo_entidad = columna("TIPO_ENTIDAD")
    else:
        es_persona = columna("TIPO_IDENTIFICACION").str.strip().str.upper().isin(("CC", "CE"))
        tipo_entidad = es_persona.map({True: "persona", False: "empresa"})
    valores = {
        "ID_REGISTRO": columna("ID_REGISTRO_FUENTE"),
        "ID_GROUP": columna("ID_GRUPO"),
        "ESTRATO": pd.Series(estrato, index=marco.index, dtype=str),
        "FUENTE": columna("FUENTE"),
        "TIPO_ENTIDAD": tipo_entidad,
        "REGIMEN": regimen,
        "CASO": columna("CASO"),
        "NIVEL_RUIDO": columna("NIVEL_RUIDO"),
        "VARIACION": columna("VARIACION"),
        "NIT": nit,
        "RAZON_SOCIAL": columna("RAZON_SOCIAL"),
        "DEPARTAMENTO": columna("DEPARTAMENTO"),
        "MUNICIPIO": columna("MUNICIPIO"),
        "CIIU": columna("CIIU"),
        "TELEFONO": columna("TELEFONO"),
        "DIRECCION": columna("DIRECCION"),
        "EMAIL": columna("EMAIL"),
    }
    return pd.DataFrame(
        {c: valores[c].to_numpy() for c in COLUMNAS_BANCO}, columns=list(COLUMNAS_BANCO)
    )


def _hoja_diccionario() -> pd.DataFrame:
    hoja_f = "FUENTE_CRM / FUENTE_RUES / FUENTE_ADUANAS"
    bandas = " | ".join(BANDAS)
    no_se_usa = (
        "No se usa salvo que al correr el motor se pida (extra_features=['{col}'] en linkage() o "
        "--variables-extra {col}:{tipo}:0.10 en scripts/banco.py); entonces se compara {como}, "
        "nunca por igualdad exacta. Con los comandos de esta plantilla es solo informativa."
    )
    filas = [
        (
            "ID_REGISTRO_FUENTE",
            hoja_f,
            "Sí",
            "texto",
            "<PREFIJO_FUENTE>-<clave primaria>",
            "CRM-000017",
            "Único dentro de la fuente y estable entre corridas: es la llave que usan VERDAD_GRUPOS y ETIQUETAS_PARES. Use la clave primaria de su sistema con el prefijo de la fuente (en la plantilla, un consecutivo de 6 dígitos). Si falta, no hay forma de nombrar un par.",
        ),
        (
            "NIT",
            hoja_f,
            "No",
            "texto",
            "solo dígitos, SIN ceros a la izquierda; con o sin DV",
            "800123456",
            "Escríbalo solo con dígitos y sin ceros a la izquierda (quítelos antes de cargar: NIT.str.lstrip('0')). El motor tolera puntos, guiones, espacios y el DV pegado (800123456 y 8001234565 son el mismo), pero NO une un NIT con ceros a la izquierda (000800123456 queda sin pareja: hallazgo abierto, caso positivo_nit_con_ceros_o_puntuacion). Vacío = registro sin identificador (régimen SIN_ID). Guarde la columna como texto y léala con dtype=str. En la plantilla son números al azar con prefijos reales de la DIAN: cualquier coincidencia con un NIT existente es casual.",
        ),
        (
            "DV",
            hoja_f,
            "No",
            "texto (1 dígito)",
            "0-9",
            "5",
            "Dígito de verificación módulo 11 cuando la fuente lo trae aparte. Vacío si no aplica o si ya va pegado al NIT. Un DV que no valida es señal de NIT mal capturado.",
        ),
        (
            "TIPO_IDENTIFICACION",
            hoja_f,
            "Sí",
            "texto",
            "NIT | CC | CE | VACIO",
            "NIT",
            "NIT para empresas; CC (cédula) o CE (extranjería) para personas naturales, cuya cédula va en la columna NIT; VACIO cuando NIT está vacío. El motor no la lee: sirve para separar personas y empresas en el análisis y para deducir TIPO_ENTIDAD. Las cédulas de la plantilla son números al azar dentro de los rangos reales.",
        ),
        (
            "RAZON_SOCIAL",
            hoja_f,
            "Sí",
            "texto",
            "mayúsculas preferible; tal como viene",
            "ZELVORA ALIMENTOS SAS",
            "Obligatoria: sin ella el registro no se puede vincular y linkage() falla. Para personas naturales, el nombre completo (el RUES escribe apellidos y luego nombres).",
        ),
        (
            "NOMBRE_COMERCIAL",
            hoja_f,
            "No",
            "texto",
            "libre",
            "PANADERIA LA ESPIGA DORADA",
            "Nombre con que opera, si la fuente lo distingue de la razón social. Vacío si no se conoce. El motor no la lee.",
        ),
        (
            "DEPARTAMENTO",
            hoja_f,
            "No",
            "texto",
            "nombre del departamento",
            "ANTIOQUIA",
            "Informativa; el motor no la lee. Vacío si no se conoce.",
        ),
        (
            "MUNICIPIO",
            hoja_f,
            "No",
            "texto",
            "nombre del municipio (con o sin tilde)",
            "MEDELLIN",
            "Es la columna que linkage() recibe como col_ciudad; en la correlativa aparece con el nombre CIUDAD. Vacío si no se conoce; no invente un valor.",
        ),
        (
            "DIRECCION",
            hoja_f,
            "No",
            "texto",
            "libre",
            "CALLE 45 No. 12 - 34 OF 201",
            no_se_usa.format(col="DIRECCION", tipo="direccion_signed", como="con tolerancia")
            + " Los formatos CL/CRA/KR equivalentes son normales.",
        ),
        (
            "TELEFONO",
            hoja_f,
            "No",
            "texto",
            "dígitos; se admiten +57, espacios y paréntesis",
            "(604) 444 5566",
            no_se_usa.format(col="TELEFONO", tipo="telefono_signed", como="con tolerancia")
            + " En la plantilla: indicativos reales con números al azar; cualquier coincidencia con un teléfono existente es casual.",
        ),
        (
            "EMAIL",
            hoja_f,
            "No",
            "texto",
            "correo en minúsculas",
            "contacto@zelvora.example",
            no_se_usa.format(col="EMAIL", tipo="email_signed", como="con tolerancia")
            + " Un correo genérico compartido (gmail, hotmail…) NO prueba que sean la misma empresa. En la plantilla todos los dominios son .example (reservado): ningún buzón existe.",
        ),
        (
            "SITIO_WEB",
            hoja_f,
            "No",
            "texto",
            "dominio o URL",
            "www.zelvora.example",
            "Informativa; el motor no la lee. En la plantilla todos los dominios son .example (reservado).",
        ),
        (
            "CIIU",
            hoja_f,
            "No",
            "texto",
            "código de 4 dígitos; varios separados por ;",
            "4631;4711",
            no_se_usa.format(
                col="CIIU", tipo="conjunto_signed", como="por solapamiento de códigos"
            ),
        ),
        (
            "PAIS",
            hoja_f,
            "Sí",
            "texto",
            "ISO-3",
            "COL",
            "Obligatoria para separar entidades extranjeras en el análisis; el motor no la lee. COL para Colombia; entidades extranjeras sin NIT van con su país y TIPO_IDENTIFICACION VACIO.",
        ),
        (
            "FECHA_ACTUALIZACION",
            hoja_f,
            "No",
            "fecha",
            "AAAA-MM-DD",
            "2026-03-15",
            "Informativa; el motor no la usa hoy. Puede servir para elegir a mano el registro más reciente de un grupo.",
        ),
        (
            "ID_REGISTRO_FUENTE",
            "VERDAD_GRUPOS",
            "Sí",
            "texto",
            "igual que en FUENTE_*",
            "CRM-000017",
            "Una fila por registro verificado; cada id aparece una sola vez. En una base real la hoja entera es OPCIONAL (solo hace falta para scripts/banco.py y para la tabla caso por caso): si la llena, incluya solo los registros que alguien verificó.",
        ),
        (
            "FUENTE",
            "VERDAD_GRUPOS",
            "Sí",
            "texto",
            "CRM | RUES | ADUANAS",
            "RUES",
            "Nombre de la hoja de fuente sin el prefijo FUENTE_ (FUENTE_CRM → CRM).",
        ),
        (
            "ID_GRUPO",
            "VERDAD_GRUPOS",
            "Sí",
            "texto",
            "G + 4 dígitos",
            "G0042",
            "Todos los registros del mismo grupo son la misma empresa o persona (misma persona jurídica o natural). Un singleton es un grupo de un registro. En una base real sirve cualquier texto estable.",
        ),
        (
            "TIPO_ENTIDAD",
            "VERDAD_GRUPOS",
            "No",
            "texto",
            "empresa | persona",
            "empresa",
            "persona cuando TIPO_IDENTIFICACION es CC o CE, o cuando se sabe que es una persona natural sin cédula; empresa en otro caso. Pasa tal cual a plantilla_banco.csv; el motor no la lee. Si falta, construir_banco la deduce de TIPO_IDENTIFICACION.",
        ),
        (
            "REGIMEN",
            "VERDAD_GRUPOS",
            "Sí",
            "texto",
            "CON_ID | SIN_ID",
            "CON_ID",
            "CON_ID si el registro trae identificador (NIT o cédula no vacíos; válido = no vacío, el DV se revisa aparte), SIN_ID si no. scripts/banco.py corta el recall por este valor (prefijo SIN_). Si falta, construir_banco lo deduce del NIT.",
        ),
        (
            "CASO",
            "VERDAD_GRUPOS",
            "Sí",
            "texto",
            "positivo_* | negativo_* | singleton | ambiguo_revisar",
            "positivo_sufijo_legal",
            "Qué prueba el grupo; el catálogo está en la hoja CASOS. Un grupo tiene un solo CASO. En una base real puede quedar vacío o tomar el nombre del catálogo que mejor describa el grupo.",
        ),
        (
            "NIVEL_RUIDO",
            "VERDAD_GRUPOS",
            "No",
            "texto",
            "CLEAN | LIGHT | MODERATE | SEVERE",
            "LIGHT",
            "Qué tan lejos está la fila del registro canónico del grupo en NIT, RAZON_SOCIAL o MUNICIPIO (CLEAN = idéntico). El canónico es el registro del RUES o, si no hay, el más completo. En una base real es libre o vacío.",
        ),
        (
            "VARIACION",
            "VERDAD_GRUPOS",
            "No",
            "texto",
            "libre",
            "sufijo legal: SAS → S.A.S.",
            "Qué se le cambió a la fila respecto del registro canónico. En una base real: lo que observó quien etiquetó, o vacío.",
        ),
        (
            "ID_1",
            "ETIQUETAS_PARES",
            "Sí",
            "texto",
            "ID_REGISTRO_FUENTE",
            "CRM-000017",
            "Primer registro del par. Debe existir en alguna hoja FUENTE_*.",
        ),
        (
            "ID_2",
            "ETIQUETAS_PARES",
            "Sí",
            "texto",
            "ID_REGISTRO_FUENTE",
            "RUES-000103",
            "Segundo registro del par (de la misma fuente o de otra).",
        ),
        (
            "IS_MATCH",
            "ETIQUETAS_PARES",
            "Sí",
            "texto",
            "1 | 0 | UNCLEAR",
            "1",
            "1 = misma empresa o persona; 0 = distinta; UNCLEAR = no se puede decidir (se descarta y se cuenta). También se aceptan SAME/DIFFERENT, MISMA/DISTINTA, SI/NO.",
        ),
        (
            "ANOTADOR",
            "ETIQUETAS_PARES",
            "No",
            "texto",
            "A | B | resuelto",
            "A",
            "Quién dio la etiqueta final: A cuando los dos coincidieron (o el nombre de quien consolidó); «resuelto» cuando no coincidieron y se decidió en sesión. scripts/medir_verdad_real.py solo calcula el kappa con los dos archivos a ciegas (--etiquetas-b); con este archivo final no lo imprime.",
        ),
        (
            "BANDA_ESPERADA",
            "ETIQUETAS_PARES",
            "No",
            "texto",
            bandas,
            "unido_borde",
            "Solo en esta plantilla: la dificultad que el generador diseñó para el par, NO la decisión del motor (en 6 de 109 pares el motor hizo otra cosa; ver docs/PLANTILLA_BASE_REAL.md §5). No la use para decidir la etiqueta.",
        ),
        (
            "BANDA_PUNTAJE",
            "ETIQUETAS_PARES (base real)",
            "No",
            "texto",
            bandas,
            "no_unido_borde",
            "En la base real la escribe scripts/verificar_plantilla_base_real.py --pares-a-etiquetar en el archivo <nombre>_bandas.csv (se pega al archivo final; nunca va en el de los anotadores): unido_holgado = el motor unió el par con puntaje ≥ umbral + 0,10 (o por identificador); unido_borde = lo unió con puntaje entre el umbral y umbral + 0,10; no_unido_borde = lo consideró candidato y no lo unió (puntaje < umbral, o ≥ umbral pero vetado por identificador distinto o por la cobertura); no_unido_lejos = nunca fue candidato (par de control). El umbral es score_threshold del perfil (0,45 en produccion_estandar). Informativa.",
        ),
        (
            "COMENTARIO",
            "ETIQUETAS_PARES",
            "No",
            "texto",
            "libre",
            "mismo NIT con y sin DV",
            "Por qué se decidió así; obligatorio cuando ANOTADOR = resuelto.",
        ),
        (
            "CASO",
            "CASOS",
            "Sí",
            "texto",
            "nombre del caso",
            "positivo_sigla_sin_nit",
            "Coincide con VERDAD_GRUPOS.CASO. La hoja CASOS es opcional en una base real.",
        ),
        (
            "DESCRIPCION",
            "CASOS",
            "Sí",
            "texto",
            "libre",
            "La sigla frente al nombre largo sin NIT",
            "Qué situación real representa.",
        ),
        (
            "QUE_DEBE_HACER_EL_MOTOR",
            "CASOS",
            "Sí",
            "texto",
            "unir | no unir | revisar",
            "revisar",
            "Comportamiento esperado; «revisar» = el motor no debe unir solo y el par va a revisión manual.",
        ),
        (
            "EJEMPLO_IDS",
            "CASOS",
            "No",
            "texto",
            "ids separados por ;",
            "CRM-000012; RUES-000034",
            "Registros que ejercitan el caso, para buscarlos en las hojas FUENTE_*.",
        ),
        (
            "N_REGISTROS",
            "CASOS",
            "No",
            "entero",
            "≥ 0",
            "12",
            "Cuántos registros de VERDAD_GRUPOS tienen ese CASO.",
        ),
        (
            "ID_REGISTRO",
            "plantilla_banco.csv",
            "Sí",
            "texto",
            "igual que ID_REGISTRO_FUENTE",
            "CRM-000017",
            "La llave de scripts/banco.py: copia de ID_REGISTRO_FUENTE. El CSV lo escribe construir_banco (VERDAD_GRUPOS + FUENTE_*); para una base real: python scripts/verificar_plantilla_base_real.py --datos mi_base.xlsx --banco salida/banco.csv.",
        ),
        (
            "ID_GROUP",
            "plantilla_banco.csv",
            "Sí",
            "texto",
            "igual que ID_GRUPO",
            "G0042",
            "Copia de VERDAD_GRUPOS.ID_GRUPO.",
        ),
        (
            "ESTRATO",
            "plantilla_banco.csv",
            "No",
            "texto",
            "libre",
            "PLANTILLA",
            "Nombre del lote: scripts/banco.py reporta precisión, recall y F1 por estrato y el macro-F1. Se fija con --estrato en verificar_plantilla_base_real.py --banco.",
        ),
        (
            "TIPO_ENTIDAD",
            "plantilla_banco.csv",
            "No",
            "texto",
            "empresa | persona",
            "persona",
            "Copia de VERDAD_GRUPOS.TIPO_ENTIDAD (o deducida de TIPO_IDENTIFICACION). scripts/banco.py no la lee.",
        ),
        (
            "(las demás columnas)",
            "plantilla_banco.csv",
            "—",
            "—",
            "—",
            "—",
            "FUENTE, REGIMEN, CASO, NIVEL_RUIDO y VARIACION vienen de VERDAD_GRUPOS; NIT, RAZON_SOCIAL, DEPARTAMENTO, MUNICIPIO, CIIU, TELEFONO, DIRECCION y EMAIL, de la hoja FUENTE_* del registro (vacías si la base no las tiene). scripts/banco.py exige ID_REGISTRO, ID_GROUP, REGIMEN, NIT, RAZON_SOCIAL, FUENTE y CASO (CASO puede ir vacío).",
        ),
    ]
    return pd.DataFrame(filas, columns=list(COLUMNAS_DICCIONARIO))


def _hoja_leeme(
    config: ConfiguracionPlantilla, n_registros: int, n_grupos: int, n_pares: int
) -> pd.DataFrame:
    verificar = "python scripts/verificar_plantilla_base_real.py --datos mi_base.xlsx --dir-trabajo salida/trabajo"
    lineas = [
        "PLANTILLA DE BASE REAL Y VERDAD REAL — rues-linker (datos 100 % inventados; semilla "
        f"{config.semilla}; {n_registros} registros, {n_grupos} grupos, {n_pares} pares etiquetados).",
        "Para qué sirve: copie esta estructura con sus datos reales. Las hojas FUENTE_* son lo que lee "
        "record_linkage.linkage(); ETIQUETAS_PARES, lo que lee scripts/medir_verdad_real.py; VERDAD_GRUPOS + FUENTE_* "
        "se convierten al CSV plano de scripts/banco.py con scripts/verificar_plantilla_base_real.py --banco.",
        "Generada por scripts/generar_plantilla_base_real.py; el diccionario de columnas está en la hoja DICCIONARIO "
        "y la guía paso a paso en docs/PLANTILLA_BASE_REAL.md.",
        "",
        "QUÉ ES CADA HOJA",
        "FUENTE_CRM, FUENTE_RUES, FUENTE_ADUANAS: una hoja por fuente, con las mismas 15 columnas. Reemplace el "
        "contenido por sus registros reales (una fila por registro, sin consolidar). Puede renombrar, agregar o quitar "
        "hojas de fuente mientras el nombre empiece por FUENTE_ (FUENTE_DIAN, FUENTE_SUPERSOCIEDADES…): "
        "scripts/verificar_plantilla_base_real.py las descubre por ese prefijo.",
        "VERDAD_GRUPOS: la partición verdadera por registro (ID_GRUPO). Todos los registros del mismo grupo son la "
        "misma empresa o persona. En una base real esta hoja es OPCIONAL (solo hace falta para scripts/banco.py y para "
        "la tabla caso por caso): si la llena, incluya solo los registros que alguien verificó; CASO puede quedar vacío "
        "o tomar el nombre del catálogo que mejor describa el grupo; NIVEL_RUIDO y VARIACION son libres (el canónico "
        "es el registro del RUES o el más completo); TIPO_ENTIDAD = persona para CC/CE.",
        "ETIQUETAS_PARES: la VERDAD REAL del protocolo (docs/PROTOCOLO_GROUND_TRUTH.md): una muestra de pares "
        "etiquetados por dos personas. Es lo que mide scripts/medir_verdad_real.py. BANDA_ESPERADA es la dificultad "
        "diseñada por el generador, no lo que hizo el motor.",
        "DICCIONARIO: cada columna, si es obligatoria, formato, ejemplo y regla (incluye las de plantilla_banco.csv). "
        "CASOS: qué situación prueba cada caso y qué debe hacer el motor (unir / no unir / revisar); también es "
        "opcional en una base real.",
        "",
        "CÓMO COPIAR LA ESTRUCTURA CON DATOS REALES",
        "1. Exporte cada fuente a su hoja FUENTE_* con estos encabezados exactos (el orden no importa, el nombre sí). "
        "Las columnas que su fuente no tenga se dejan vacías, no se borran.",
        "2. Dé a cada registro un ID_REGISTRO_FUENTE único y estable (prefijo de la fuente + clave primaria): es la "
        "llave de la verdad. Nunca cambie un id después de etiquetar.",
        "3. Columnas obligatorias: ID_REGISTRO_FUENTE, RAZON_SOCIAL, TIPO_IDENTIFICACION y PAIS (las dos últimas no "
        "las lee el motor: separan personas y extranjeros en el análisis). Opcionales: NIT, DV, NOMBRE_COMERCIAL, "
        "DEPARTAMENTO, MUNICIPIO, DIRECCION, TELEFONO, EMAIL, SITIO_WEB, CIIU, FECHA_ACTUALIZACION. De las opcionales "
        "el motor solo usa MUNICIPIO (col_ciudad); TELEFONO, EMAIL, DIRECCION y CIIU solo cuentan si al correr el "
        "motor se piden (extra_features=['TELEFONO', 'EMAIL'] en linkage() o --variables-extra "
        "TELEFONO:telefono_signed:0.10 en scripts/banco.py); con los comandos de esta plantilla son informativas.",
        "4. NIT: solo dígitos y SIN ceros a la izquierda (quítelos antes de cargar: NIT.str.lstrip('0')). Puntos, "
        "guiones, espacios y el DV pegado los tolera el motor (800123456 y 8001234565 son el mismo), pero se "
        "recomienda quitarlos; los ceros a la izquierda NO los tolera hoy (hallazgo abierto: 000901156189 queda sin "
        "pareja, caso positivo_nit_con_ceros_o_puntuacion). Si la fuente trae el DV aparte, va en la columna DV. "
        "Guarde la columna como TEXTO y léala con dtype=str (paso 7). Vacío = sin identificador.",
        "5. Personas naturales: la cédula va en NIT y TIPO_IDENTIFICACION = CC (o CE); las cédulas no llevan DV.",
        "6. Guarde la base real FUERA del repositorio (Drive, carpeta verdad_real/): tiene nombres reales y datos "
        "licenciados. Al repositorio solo entran cifras.",
        "",
        "CÓMO PRODUCIR LA VERDAD REAL (ETIQUETAS_PARES)",
        f"7. Corra el motor y guarde la correlativa: {verificar} --correlativa salida/correlativa.parquet (descubre "
        "las hojas FUENTE_*; sin VERDAD_GRUPOS ni CASOS solo corre el motor). Si prefiere Python: lea cada hoja con "
        "pd.read_excel(ruta, sheet_name='FUENTE_CRM', dtype=str, keep_default_na=False) — sin dtype=str pandas "
        "vuelve numérico el NIT, pierde los ceros y los vacíos y el motor da otra partición — y llame "
        "linkage({'CRM': df_crm, …}, col_nit='NIT', col_name='RAZON_SOCIAL', col_ciudad='MUNICIPIO', "
        "work_dir='salida/trabajo'); luego resultado.correlativa.to_parquet('salida/correlativa.parquet'). No pase "
        "col_id (el motor duplicaría el prefijo: CRM-CRM-000001): la correlativa conserva ID_REGISTRO_FUENTE en cada "
        "fila, que es la columna con la que se mide; MUNICIPIO aparece en ella como CIUDAD.",
        f"8. Muestree los pares con el mismo script: {verificar} --pares-a-etiquetar salida/pares_a_etiquetar.csv "
        "--n-pares 500. Salen dos archivos: pares_a_etiquetar.csv (ID_1, ID_2 y los datos de los dos registros; sin "
        "ninguna columna del motor) para los anotadores, y pares_a_etiquetar_bandas.csv (BANDA_PUNTAJE, PUNTAJE, "
        "REGIMEN) que se pega al archivo final. La muestra se reparte por régimen (los dos con id / ninguno / mixto) y "
        "por banda (unido_holgado, unido_borde, no_unido_borde, no_unido_lejos; definidas en DICCIONARIO). Empiece "
        "por 500. scripts/active_labeling.py NO sirve aquí: no lee la correlativa y escribe ids internos que no son "
        "ID_REGISTRO_FUENTE.",
        "9. Dos personas etiquetan el MISMO archivo a ciegas (sin ver la decisión del motor ni la del otro), cada una "
        "en su copia (etiquetas_A.csv, etiquetas_B.csv). Cada una escribe en IS_MATCH: 1 = misma empresa, 0 = "
        "distinta, UNCLEAR = no se puede decidir; y en ANOTADOR su letra.",
        "10. REGLA DE DECISIÓN «misma empresa»: la misma persona jurídica o natural (misma matrícula / NIT raíz "
        "aunque cambie la sucursal, la forma legal o la grafía). Una SUCURSAL es la misma empresa. Una FILIAL, una "
        "MATRIZ o un grupo empresarial NO es la misma empresa. Si hace falta buscar en el RUES o en la web, se busca.",
        "11. Dos personas con el mismo nombre y cédulas distintas son distintas. Dos empresas con razón social "
        "idéntica y NIT válidos distintos son distintas. El mismo NIT con nombres muy distintos se marca UNCLEAR; si "
        "la consulta al RUES lo resuelve antes de cerrar el archivo, se escribe 1 o 0 con ANOTADOR = resuelto y el "
        "porqué en COMENTARIO.",
        "12. Mida el acuerdo con los dos archivos a ciegas: python scripts/medir_verdad_real.py --etiquetas "
        "etiquetas_A.csv --etiquetas-b etiquetas_B.csv --correlativa salida/correlativa.parquet --col-id "
        "ID_REGISTRO_FUENTE --col-nit NIT. Kappa > 0,80: verdad confiable; 0,60-0,80: resolver desacuerdos en "
        "sesión; < 0,60: reescribir esta regla y volver a etiquetar. El kappa solo se calcula con los dos archivos; "
        "el archivo final (ANOTADOR = A/B/resuelto) no lo imprime.",
        "13. El archivo final tiene UNA etiqueta por par; ANOTADOR = A cuando coincidieron (o quien consolidó) y "
        "«resuelto» si se decidió en sesión (con el porqué en COMENTARIO). Nunca se reetiqueta lo ya resuelto; la "
        "verdad crece con pares nuevos.",
        "",
        "CÓMO MEDIR",
        "14. python scripts/medir_verdad_real.py --etiquetas verdad_real/etiquetas_resueltas.csv --correlativa "
        "salida/correlativa.parquet --col-id ID_REGISTRO_FUENTE --col-grupo ID_GRUPO --col-nit NIT --salida "
        "docs/evidencia/verdad_real_<fecha>.json. El archivo de etiquetas es un CSV o un Excel con los pares en la "
        "primera hoja; si están en otra hoja (como ETIQUETAS_PARES en este libro), añada --hoja ETIQUETAS_PARES.",
        "15. Reporta pares, TP/FP/FN, precisión y recall con intervalo de Wilson al 95 % y F1, por régimen "
        "(CON_ID, SIN_ID, MIXTO) y global. Los UNCLEAR se descartan y se cuentan.",
        "16. Para regresión con la partición completa (VERDAD_GRUPOS) use el CSV plano: python scripts/banco.py "
        "--datos data/ejemplos/plantilla_banco.csv --etiqueta plantilla --dir-trabajo /tmp/banco_plantilla "
        f"--evidencia /tmp/evid_plantilla. El de su base lo escribe {verificar} --banco salida/banco.csv --estrato "
        "MI_LOTE (construir_banco: VERDAD_GRUPOS + FUENTE_*; las columnas están en DICCIONARIO).",
        "17. Para ver, caso por caso, qué hizo el motor con esta plantilla: python scripts/verificar_plantilla_base_real.py "
        "(la tabla resultante está en docs/PLANTILLA_BASE_REAL.md). Los fallos no se ocultan: la plantilla también "
        "muestra lo que el motor no sabe hacer hoy.",
        "",
        "NOTAS",
        "18. Los NIT, cédulas y teléfonos de esta plantilla son números al azar dentro de los rangos reales (prefijos "
        "de la DIAN, rangos de la Registraduría, indicativos de ciudad y de celular): cualquier coincidencia con un "
        "NIT, una cédula o un teléfono existente es casual. El DV sí es el correcto (módulo 11) cuando la fila lo "
        "trae, y ninguna pareja de NIT de empresas distintas difiere en un solo dígito salvo en el caso "
        "negativo_nit_corrupto_un_digito.",
        "19. La cobertura de las columnas opcionales cambia por fuente a propósito (el RUES trae CIIU y dirección; "
        "aduanas casi nunca trae correo). NIVEL_RUIDO describe solo NIT, RAZON_SOCIAL y MUNICIPIO.",
        "20. Los tres pares UNCLEAR están entre registros del caso ambiguo_revisar, que en VERDAD_GRUPOS son grupos "
        "distintos: ilustran que un par dudoso no afirma nada y no se cuenta en las métricas.",
        "21. Ninguna empresa real: las razones sociales se arman con el vocabulario de prueba del repositorio "
        "(KRONOS, ZELVORA, TRIVENTA, PELUVIKA, DULCEVENT, VALDORIA, ORBITAL, MARMOLERIA QUINTAL…) y no identifican a "
        "ninguna empresa concreta; los nombres de uso común (INVERSIONES SAS, DISTRIBUIDORA DEL NORTE SAS, PANADERIA "
        "LA ESPIGA DORADA…) están en negativo_nombre_generico y en los nombres comerciales justamente porque son "
        "comunes, y tampoco apuntan a nadie. Los correos y sitios web usan el dominio reservado .example (no existe "
        "ningún buzón). tests/test_sin_nombres_reales.py vigila los CSV de esta carpeta.",
    ]
    return pd.DataFrame({"LEEME": lineas})


# ---------------------------------------------------------------------------
# Escritura
# ---------------------------------------------------------------------------


def _nombre_csv(hoja: str) -> str:
    return f"plantilla_{hoja.lower()}.csv"


def escribir(plantilla: Plantilla, config: ConfiguracionPlantilla) -> list[Path]:
    """Escribe el Excel y los CSV; devuelve las rutas escritas."""
    salida = config.salida
    salida.mkdir(parents=True, exist_ok=True)
    rutas: list[Path] = []
    hojas = plantilla.hojas()
    ruta_xlsx = salida / config.nombre_excel
    with pd.ExcelWriter(ruta_xlsx, engine="openpyxl") as escritor:
        for nombre, df in hojas.items():
            df.to_excel(escritor, sheet_name=nombre, index=False)
            _dar_formato(escritor.sheets[nombre], df, es_texto=(nombre == "LEEME"))
    rutas.append(ruta_xlsx)
    for nombre, df in hojas.items():
        if nombre == "LEEME":
            continue
        ruta = salida / _nombre_csv(nombre)
        df.to_csv(ruta, index=False, encoding="utf-8", lineterminator="\n")
        rutas.append(ruta)
    for etiqueta, df in (("A", plantilla.pares_anotador_a), ("B", plantilla.pares_anotador_b)):
        ruta = salida / f"plantilla_etiquetas_anotador_{etiqueta}.csv"
        df.to_csv(ruta, index=False, encoding="utf-8", lineterminator="\n")
        rutas.append(ruta)
    ruta = salida / "plantilla_banco.csv"
    plantilla.banco.to_csv(ruta, index=False, encoding="utf-8", lineterminator="\n")
    rutas.append(ruta)
    return rutas


def _dar_formato(hoja, df: pd.DataFrame, *, es_texto: bool) -> None:
    """Encabezado en negrita, panel fijo, filtro y anchos legibles (solo estética)."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    relleno = PatternFill(start_color="DDE7F0", end_color="DDE7F0", fill_type="solid")
    for celda in hoja[1]:
        celda.font = Font(bold=True)
        celda.fill = relleno
    hoja.freeze_panes = "A2"
    if es_texto:
        hoja.column_dimensions["A"].width = 140
        for fila in hoja.iter_rows(min_row=2):
            fila[0].alignment = Alignment(wrap_text=True, vertical="top")
        return
    hoja.auto_filter.ref = hoja.dimensions
    for i, columna in enumerate(df.columns, start=1):
        largo = max([len(str(columna))] + [len(str(v)) for v in df[columna].head(200)])
        hoja.column_dimensions[get_column_letter(i)].width = min(max(10, largo + 2), 60)


# ---------------------------------------------------------------------------
# Punto de entrada
# ---------------------------------------------------------------------------


def generar(config: ConfiguracionPlantilla) -> Plantilla:
    """Genera la plantilla completa (sin escribir nada)."""
    gen = Generador(config)
    bloques = gen.generar_bloques()
    return ensamblar(config, bloques, gen.rng)


def resumen(plantilla: Plantilla, rutas: list[Path]) -> str:
    partes = [f"{f}: {len(df)} registros" for f, df in plantilla.fuentes.items()]
    partes.append(f"grupos: {plantilla.verdad['ID_GRUPO'].nunique()}")
    partes.append(
        f"pares: {len(plantilla.pares)} ({(plantilla.pares['IS_MATCH'] == 'UNCLEAR').sum()} UNCLEAR)"
    )
    partes.append(
        "casos: "
        + ", ".join(f"{c}={n}" for c, n in plantilla.casos.set_index("CASO")["N_REGISTROS"].items())
    )
    partes.append("archivos: " + ", ".join(str(r) for r in rutas))
    return "\n".join(partes)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument(
        "--salida", type=Path, default=ConfiguracionPlantilla.salida, help="carpeta de salida"
    )
    parser.add_argument(
        "--semilla", type=int, default=ConfiguracionPlantilla.semilla, help="semilla del azar"
    )
    args = parser.parse_args(argv)
    config = ConfiguracionPlantilla(semilla=args.semilla, salida=args.salida)
    plantilla = generar(config)
    rutas = escribir(plantilla, config)
    print(resumen(plantilla, rutas))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

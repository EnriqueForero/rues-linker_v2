"""``exportar_vistas``: la función libre del estándar para las VISTAS derivadas (F2.11).

Qué congela
-----------
* una vista por archivo: ``<carpeta>/<nombre>__<vista>.xlsx`` (hoja = vista,
  fórmulas neutralizadas, sin mutar la entrada) si cabe en Excel, y
  ``<nombre>__<vista>.csv.gz`` si no cabe —nunca un recorte—;
* ``libro=True``: UN ``<carpeta>/<nombre>.xlsx`` con una hoja por vista
  (nombres saneados y únicos); la vista que no cabe va a ``.csv.gz`` al lado;
* ``formato="csv"``: un ``.csv`` plano (neutralizado) por vista;
* sin ``nombre`` el archivo se llama ``<vista>.<ext>``; con ``libro`` hace falta;
* fail-fast con mensaje accionable: sin vistas, un valor que no es DataFrame,
  una vista o un nombre con separadores de ruta, ``libro`` con ``formato="csv"``;
* con ``diccionario=`` (el mismo que ``diccionario.csv``) cada ``.xlsx`` lleva
  la hoja ``DICCIONARIO`` de su vista (F2.16): las filas de la tabla homónima
  del diccionario o, si no la hay, las de sus columnas; en un libro, una sola
  hoja al final con la columna ``hoja``; ``formato="csv"`` no lo admite;
* ``PipelineResult.to_excel``/``to_csv`` son ALIAS de esta función: avisan con
  ``DeprecationWarning`` y ya no escriben nada por su cuenta.

Las empresas son inventadas. Ningún dato licenciado entra aquí.
"""

from __future__ import annotations

import gzip
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from record_linkage.exporters import escritor
from record_linkage.exporters.escritor import exportar_vistas


def _vistas() -> dict[str, pd.DataFrame]:
    return {
        "DUPLICADOS": pd.DataFrame(
            {"ID_GRUPO": [1, 1, 2], "RAZON_SOCIAL": ["=ACME", "ACME SAS", "+BETA"]}
        ),
        "RESUMEN": pd.DataFrame({"metrica": ["n_registros"], "valor": [3]}),
    }


def _hojas(ruta: Path) -> dict[str, list[list[object]]]:
    libro = load_workbook(ruta, read_only=True, data_only=False)
    try:
        return {
            hoja.title: [list(fila) for fila in hoja.iter_rows(values_only=True)]
            for hoja in libro.worksheets
        }
    finally:
        libro.close()


# ─────────────────────────────────────────────────────────────────────────────
# Una vista por archivo (notebooks 01 a 04)
# ─────────────────────────────────────────────────────────────────────────────


def test_una_vista_por_archivo_xlsx_con_hoja_y_neutralizada(tmp_path: Path) -> None:
    vistas = _vistas()
    antes = {k: v.copy(deep=True) for k, v in vistas.items()}
    carpeta = tmp_path / "vistas"  # no existe: la función la crea

    rutas = exportar_vistas(vistas, carpeta, "deduplicacion")

    assert rutas == [
        carpeta / "deduplicacion__DUPLICADOS.xlsx",
        carpeta / "deduplicacion__RESUMEN.xlsx",
    ]
    assert all(r.is_file() for r in rutas)
    hojas = _hojas(rutas[0])
    assert list(hojas) == ["DUPLICADOS"]
    assert hojas["DUPLICADOS"][0] == ["ID_GRUPO", "RAZON_SOCIAL"]
    assert hojas["DUPLICADOS"][1][1] == "'=ACME"
    assert hojas["DUPLICADOS"][3][1] == "'+BETA"
    assert list(_hojas(rutas[1])) == ["RESUMEN"]
    for clave, original in antes.items():
        pd.testing.assert_frame_equal(vistas[clave], original)


def test_sin_nombre_el_archivo_se_llama_como_la_vista(tmp_path: Path) -> None:
    rutas = exportar_vistas(_vistas(), tmp_path)
    assert [r.name for r in rutas] == ["DUPLICADOS.xlsx", "RESUMEN.xlsx"]


def test_la_vista_que_no_cabe_va_a_csv_gz_sin_recorte(tmp_path: Path) -> None:
    vistas = _vistas()
    rutas = exportar_vistas(vistas, tmp_path, "dedup", limite=2)
    assert [r.name for r in rutas] == ["dedup__DUPLICADOS.csv.gz", "dedup__RESUMEN.xlsx"]
    with gzip.open(rutas[0], "rt", encoding="utf-8") as f:
        texto = f.read()
    assert texto.count("\n") == 4  # encabezado + 3 filas: completo
    assert "'=ACME" in texto and "'+BETA" in texto  # neutralizado: se abre en hoja de cálculo


def test_una_vista_vacia_se_escribe_con_su_encabezado(tmp_path: Path) -> None:
    """«0 duplicados» es información: la vista vacía se escribe, no se calla."""
    vacia = pd.DataFrame({"ID_GRUPO": pd.Series([], dtype="int64"), "RAZON_SOCIAL": []})
    (ruta,) = exportar_vistas({"DUPLICADOS": vacia}, tmp_path, "dedup")
    assert _hojas(ruta)["DUPLICADOS"] == [["ID_GRUPO", "RAZON_SOCIAL"]]


def test_el_limite_por_defecto_es_el_del_estandar(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """``limite=None`` resuelve ``LIMITE_FILAS_EXCEL`` AL LLAMAR (una sola regla)."""
    monkeypatch.setattr(escritor.excel, "LIMITE_FILAS_EXCEL", 1)
    rutas = exportar_vistas(_vistas(), tmp_path, "dedup")
    assert [r.suffixes[-1] for r in rutas] == [".gz", ".xlsx"]


# ─────────────────────────────────────────────────────────────────────────────
# Un libro con una hoja por vista (notebooks 07 y 09)
# ─────────────────────────────────────────────────────────────────────────────


def test_libro_con_una_hoja_por_vista(tmp_path: Path) -> None:
    rutas = exportar_vistas(_vistas(), tmp_path, "resultado", libro=True)
    assert rutas == [tmp_path / "resultado.xlsx"]
    hojas = _hojas(rutas[0])
    assert list(hojas) == ["DUPLICADOS", "RESUMEN"]
    assert hojas["DUPLICADOS"][1][1] == "'=ACME"
    assert hojas["RESUMEN"][1] == ["n_registros", 3]


def test_libro_sanea_y_desduplica_los_nombres_de_hoja(tmp_path: Path) -> None:
    prefijo = "x" * 31
    vistas = {
        f"{prefijo}primero": pd.DataFrame({"x": [1]}),
        f"{prefijo}segundo": pd.DataFrame({"x": [2]}),
        "mal/nombre:*?[]": pd.DataFrame({"x": [3]}),
        "control\x01nombre": pd.DataFrame({"x": [4]}),
    }
    (ruta,) = exportar_vistas(vistas, tmp_path, "libro", libro=True)
    nombres = list(_hojas(ruta))
    assert len(nombres) == 4 and len({n.casefold() for n in nombres}) == 4
    assert all(len(n) <= 31 for n in nombres)
    assert all(
        not any(c in n for c in "[]:*?/\\") and not any(ord(c) < 32 for c in n) for n in nombres
    )


def test_libro_la_vista_que_no_cabe_va_a_csv_gz_al_lado(tmp_path: Path) -> None:
    rutas = exportar_vistas(_vistas(), tmp_path, "resultado", libro=True, limite=2)
    assert [r.name for r in rutas] == ["resultado.xlsx", "resultado__DUPLICADOS.csv.gz"]
    assert list(_hojas(rutas[0])) == ["RESUMEN"]


def test_libro_sin_ninguna_vista_que_quepa_no_deja_libro(tmp_path: Path) -> None:
    rutas = exportar_vistas(_vistas(), tmp_path, "resultado", libro=True, limite=0)
    assert [r.name for r in rutas] == ["resultado__DUPLICADOS.csv.gz", "resultado__RESUMEN.csv.gz"]
    assert not (tmp_path / "resultado.xlsx").exists()


def test_libro_exige_nombre(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="nombre"):
        exportar_vistas(_vistas(), tmp_path, libro=True)


# ─────────────────────────────────────────────────────────────────────────────
# CSV plano por vista
# ─────────────────────────────────────────────────────────────────────────────


def test_formato_csv_escribe_un_csv_plano_neutralizado_por_vista(tmp_path: Path) -> None:
    rutas = exportar_vistas(_vistas(), tmp_path, formato="csv")
    assert [r.name for r in rutas] == ["DUPLICADOS.csv", "RESUMEN.csv"]
    texto = rutas[0].read_text(encoding="utf-8")
    assert "'=ACME" in texto and "'+BETA" in texto
    assert len(pd.read_csv(rutas[0])) == 3


def test_formato_csv_no_admite_libro(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="libro"):
        exportar_vistas(_vistas(), tmp_path, "x", libro=True, formato="csv")


def test_formato_desconocido_falla_rapido(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="formato"):
        exportar_vistas(_vistas(), tmp_path, formato="parquet")  # type: ignore[arg-type]


# ─────────────────────────────────────────────────────────────────────────────
# Fail-fast
# ─────────────────────────────────────────────────────────────────────────────


def test_sin_vistas_falla_con_mensaje_accionable(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="ninguna vista"):
        exportar_vistas({}, tmp_path, "x")
    assert not any(tmp_path.iterdir())


def test_un_valor_que_no_es_dataframe_falla_antes_de_escribir(tmp_path: Path) -> None:
    with pytest.raises(TypeError, match="RESUMEN"):
        exportar_vistas({"DUPLICADOS": _vistas()["DUPLICADOS"], "RESUMEN": 12.5}, tmp_path, "x")  # type: ignore[dict-item]
    assert not any(tmp_path.iterdir())


@pytest.mark.parametrize("vista", ["../fuera", "sub/vista", r"sub\vista", "mala\x00"])
def test_una_vista_con_separadores_no_nombra_un_archivo(tmp_path: Path, vista: str) -> None:
    with pytest.raises(ValueError, match=r"nombre simple|control"):
        exportar_vistas({vista: pd.DataFrame({"x": [1]})}, tmp_path, "x")
    assert not (tmp_path.parent / "fuera.xlsx").exists()
    assert not any(tmp_path.iterdir())


@pytest.mark.parametrize("nombre", ["../fuera", "a/b", ""])
def test_un_nombre_con_separadores_falla_rapido(tmp_path: Path, nombre: str) -> None:
    with pytest.raises(ValueError, match="nombre"):
        exportar_vistas(_vistas(), tmp_path, nombre)


def test_esta_en_la_api_publica() -> None:
    import record_linkage

    assert record_linkage.exportar_vistas is exportar_vistas
    assert "exportar_vistas" in escritor.__all__


# ─────────────────────────────────────────────────────────────────────────────
# Hoja DICCIONARIO (F2.16)
# ─────────────────────────────────────────────────────────────────────────────


def _diccionario() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "tabla": ["correlativa", "correlativa", "golden"],
            "columna": ["ID_GRUPO", "RAZON_SOCIAL", "ID_GRUPO"],
            "tipo": ["int64", "string", "int64"],
            "significado": ["Grupo del registro.", "Nombre de la fuente.", "Grupo (golden)."],
            "origen": ["motor", "fuente", "motor"],
            "alias_es": ["ID_GRUPO", "RAZON_SOCIAL", "ID_GRUPO"],
        }
    )


ENCABEZADO_DICCIONARIO = ["columna", "alias_es", "significado", "origen"]


def test_con_diccionario_cada_xlsx_lleva_la_hoja_de_su_vista(tmp_path: Path) -> None:
    rutas = exportar_vistas(_vistas(), tmp_path, "dedup", diccionario=_diccionario())
    assert rutas == [tmp_path / "dedup__DUPLICADOS.xlsx", tmp_path / "dedup__RESUMEN.xlsx"]
    duplicados = _hojas(rutas[0])
    assert list(duplicados) == ["DUPLICADOS", "DICCIONARIO"]
    assert duplicados["DUPLICADOS"][0] == ["ID_GRUPO", "RAZON_SOCIAL"]  # nombres sin alias
    assert duplicados["DICCIONARIO"] == [
        ENCABEZADO_DICCIONARIO,
        ["ID_GRUPO", "ID_GRUPO", "Grupo del registro.", "motor"],
        ["RAZON_SOCIAL", "RAZON_SOCIAL", "Nombre de la fuente.", "fuente"],
    ]
    # RESUMEN no tiene ninguna columna en el diccionario: sin hoja vacía.
    assert list(_hojas(rutas[1])) == ["RESUMEN"]


def test_con_diccionario_la_vista_homonima_de_una_tabla_lleva_las_filas_de_la_tabla(
    tmp_path: Path,
) -> None:
    vista = pd.DataFrame({"ID_GRUPO": [1], "RAZON_SOCIAL": ["ACME"]})
    (ruta,) = exportar_vistas({"CORRELATIVA": vista}, tmp_path, diccionario=_diccionario())
    hojas = _hojas(ruta)
    assert list(hojas) == ["CORRELATIVA", "DICCIONARIO"]
    assert [f[0] for f in hojas["DICCIONARIO"][1:]] == ["ID_GRUPO", "RAZON_SOCIAL"]


def test_libro_con_diccionario_lleva_una_sola_hoja_al_final_con_la_columna_hoja(
    tmp_path: Path,
) -> None:
    (ruta,) = exportar_vistas(
        _vistas(), tmp_path, "resultado", libro=True, diccionario=_diccionario()
    )
    hojas = _hojas(ruta)
    assert list(hojas) == ["DUPLICADOS", "RESUMEN", "DICCIONARIO"]
    assert hojas["DICCIONARIO"] == [
        ["hoja", *ENCABEZADO_DICCIONARIO],
        ["DUPLICADOS", "ID_GRUPO", "ID_GRUPO", "Grupo del registro.", "motor"],
        ["DUPLICADOS", "RAZON_SOCIAL", "RAZON_SOCIAL", "Nombre de la fuente.", "fuente"],
    ]


def test_libro_una_vista_llamada_diccionario_no_choca_con_la_hoja(tmp_path: Path) -> None:
    vistas = {"DICCIONARIO": pd.DataFrame({"ID_GRUPO": [1]}), "RESUMEN": _vistas()["RESUMEN"]}
    (ruta,) = exportar_vistas(vistas, tmp_path, "x", libro=True, diccionario=_diccionario())
    nombres = list(_hojas(ruta))
    assert nombres[:2] == ["DICCIONARIO", "RESUMEN"] and len(nombres) == 3
    assert nombres[2] != "DICCIONARIO" and nombres[2].startswith("DICCIONARIO")


def test_sin_diccionario_nada_cambia(tmp_path: Path) -> None:
    rutas = exportar_vistas(_vistas(), tmp_path, "dedup")
    assert [list(_hojas(r)) for r in rutas] == [["DUPLICADOS"], ["RESUMEN"]]


def test_formato_csv_no_admite_diccionario(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="diccionario"):
        exportar_vistas(_vistas(), tmp_path, formato="csv", diccionario=_diccionario())
    assert list(tmp_path.iterdir()) == []


def test_diccionario_sin_las_columnas_del_estandar_falla_antes_de_escribir(
    tmp_path: Path,
) -> None:
    with pytest.raises(ValueError, match="diccionario"):
        exportar_vistas(_vistas(), tmp_path, "x", diccionario=pd.DataFrame({"columna": ["a"]}))
    assert list(tmp_path.iterdir()) == []

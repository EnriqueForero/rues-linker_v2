"""Regresiones para endurecimientos encontrados en la revisión final."""

from __future__ import annotations

import hashlib
import json
import os
import runpy
import sqlite3
import zipfile
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from record_linkage.api import _ordered_trusted_sources, linkage
from record_linkage.config.credentials import _load_from_file
from record_linkage.deduplication.unified import deduplicate_unified
from record_linkage.engine.lsh.disk_based import DiskBasedLSHEngine
from record_linkage.exporters.smart import SmartExporter
from record_linkage.golden.generator import GoldenRecordGeneratorV7
from record_linkage.ingestion import ArchiveSafetyError, SourceSpec, ZipSafetyLimits, load_source
from record_linkage.pipeline.orchestrator import Orchestrator
from record_linkage.reporting.data_handler import DataHandler


def test_golden_fallback_aborts_instead_of_losing_a_failed_chunk(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    frame = pd.DataFrame(
        {
            "ID_GRUPO": range(600),
            "NIT": [f"900{index:06d}" for index in range(600)],
            "RAZON_SOCIAL": [f"EMPRESA {index}" for index in range(600)],
            "SRC": ["TEST"] * 600,
            "ORIGINAL_INDEX": range(600),
            "MARK": range(600),
        }
    )
    original_to_sql = pd.DataFrame.to_sql
    call_count = 0

    def fail_primary_and_second_fallback_chunk(self, *args, **kwargs):
        nonlocal call_count
        call_count += 1
        if call_count in {1, 3}:
            raise sqlite3.OperationalError("fallo inyectado")
        return original_to_sql(self, *args, **kwargs)

    monkeypatch.setattr(pd.DataFrame, "to_sql", fail_primary_and_second_fallback_chunk)
    generator = GoldenRecordGeneratorV7(["TEST"])

    with pytest.raises(RuntimeError, match="no se publicará un resultado parcial"):
        generator._save_dataframe_optimized(frame, str(tmp_path / "input.db"))

    with sqlite3.connect(tmp_path / "input.db") as connection:
        table_exists = connection.execute(
            "SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name='correlative_table'"
        ).fetchone()[0]
    assert table_exists == 0


def test_golden_output_invariants_reject_missing_rows() -> None:
    frame = pd.DataFrame(
        {
            "ID_GRUPO": [0, 1],
            "NIT": ["900111111", "900222222"],
            "RAZON_SOCIAL": ["ALFA", "BETA"],
            "SRC": ["TEST", "TEST"],
            "ORIGINAL_INDEX": [0, 1],
        }
    )
    generator = GoldenRecordGeneratorV7(["TEST"])
    expected_index = generator._validate_input_contract(frame)
    golden = pd.DataFrame({"ID_GRUPO": [0]})
    partial = frame.iloc[[0]].copy()

    with pytest.raises(RuntimeError, match="no conserva todas las filas"):
        generator._validate_output_invariants(frame, expected_index, golden, partial)


def test_real_archive_proxy_compares_nit_base_instead_of_incompatible_raw_formats() -> None:
    verifier = runpy.run_path(
        str(Path(__file__).parents[1] / "scripts" / "verify_real_archives.py")
    )
    nit_consistency_proxies = verifier["_nit_consistency_proxies"]
    mapped = pd.DataFrame(
        {
            "ORIGINAL_INDEX": [0, 1],
            "ID_GRUPO": [7, 7],
            "EXPECTED_SRC": ["RUES", "EXPORTACIONES"],
            "EXPECTED_NIT": ["900123456", "9001234568"],
            "NIT_BASE": ["900123456", "900123456"],
        }
    )

    proxies = nit_consistency_proxies(mapped)

    assert proxies["canonical_nit_field"] == "NIT_BASE"
    assert proxies["common_canonical_nit_bases"] == 1
    assert proxies["common_base_group_consistency_rate"] == 1.0
    assert proxies["export_common_base_link_rate"] == 1.0
    assert proxies["cross_source_groups_with_conflicting_bases"] == 0


def test_trusted_sources_follow_source_order_and_reject_unknown() -> None:
    order = ["EXPORTACIONES", "RUES", "CRM"]
    assert _ordered_trusted_sources(order, {"CRM", "RUES"}) == ["RUES", "CRM"]
    with pytest.raises(ValueError, match="no declaradas"):
        _ordered_trusted_sources(order, {"RUES", "FANTASMA"})


@pytest.mark.skipif(os.name != "posix", reason="Los permisos POSIX no existen en Windows")
def test_credential_file_must_be_private_and_preserves_secret_bytes(tmp_path: Path) -> None:
    config = tmp_path / "config.json"
    payload = {
        "snowflake": {
            "account": "acct",
            "user": "user",
            "password": "  secret with spaces  ",
            "warehouse": "wh",
            "database": "db",
            "schema": "schema",
        }
    }
    config.write_text(json.dumps(payload), encoding="utf-8")
    config.chmod(0o644)
    assert _load_from_file(config) is None

    config.chmod(0o600)
    loaded = _load_from_file(config)
    assert loaded is not None
    assert loaded["password"] == "  secret with spaces  "


def test_xlsx_internal_zip_is_subject_to_uncompressed_size_limit(tmp_path: Path) -> None:
    workbook = tmp_path / "source.xlsx"
    pd.DataFrame({"id": ["123456"], "name": ["ACME"]}).to_excel(
        workbook, index=False, engine="openpyxl"
    )
    with zipfile.ZipFile(workbook) as archive:
        declared_size = sum(info.file_size for info in archive.infolist() if not info.is_dir())
    limits = ZipSafetyLimits(
        max_members=128,
        max_member_bytes=max(1, declared_size),
        max_total_uncompressed_bytes=max(1, declared_size - 1),
        max_compression_ratio=1_000,
    )
    spec = SourceSpec(
        name="xlsx",
        path=workbook,
        column_mapping={"NIT": "id", "RAZON_SOCIAL": "name"},
        safety_limits=limits,
    )
    with pytest.raises(ArchiveSafetyError, match="sin comprimir"):
        load_source(spec)


def test_xlsx_projection_preserves_rows_with_values_only_in_unmapped_columns(
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "projection.xlsx"
    pd.DataFrame({"ID": [None, "123456"], "UNMAPPED": ["ROW_EXISTS", "X"]}).to_excel(
        workbook, index=False, engine="openpyxl"
    )

    loaded = load_source(
        SourceSpec(
            name="xlsx",
            path=workbook,
            column_mapping={"NIT": "ID"},
            column_types={"NIT": "identifier"},
        )
    )

    assert len(loaded.data) == 2
    assert loaded.report.rows == 2
    assert loaded.data["NIT"].isna().iloc[0]
    assert loaded.data["NIT"].iloc[1] == "123456"


def test_unified_dedup_never_reuses_unverified_checkpoints(tmp_path: Path) -> None:
    output = tmp_path / "same-output"
    first = pd.DataFrame(
        {
            "NIT": ["900111111", "900222222"],
            "RAZON_SOCIAL": ["ALFA SAS", "BETA SAS"],
            "MARK": ["A0", "A1"],
        }
    )
    second = pd.DataFrame(
        {
            "NIT": ["901333333", "901444444", "901555555"],
            "RAZON_SOCIAL": ["GAMMA SAS", "DELTA SAS", "EPSILON SAS"],
            "MARK": ["B0", "B1", "B2"],
        }
    )

    deduplicate_unified(first, output_dir=str(output), validate_against_legacy=False)
    result, _ = deduplicate_unified(second, output_dir=str(output), validate_against_legacy=False)

    assert len(result) == len(second)
    assert set(result["MARK"]) == {"B0", "B1", "B2"}


def test_legacy_data_handler_fails_closed_instead_of_dropping_a_source() -> None:
    good = pd.DataFrame({"NIT": ["900123456"], "RAZON_SOCIAL": ["ACME"]})
    bad = pd.DataFrame({"NIT": ["900654321"], "WRONG_NAME": ["BETA"]})

    with pytest.raises(RuntimeError, match=r"resultado parcial.*BAD"):
        DataHandler().load_sources({"OK": good, "BAD": bad})

    loaded, report = DataHandler({"strict_source_loading": False}).load_sources(
        {"OK": good, "BAD": bad}
    )
    assert list(loaded) == ["OK"]
    assert len(report["errors"]) == 1


@pytest.mark.parametrize(
    "name", ["../outside", "sub/outside", r"sub\outside", "/tmp/x", "bad\x00x"]
)
def test_experiment_name_cannot_escape_work_directory(tmp_path: Path, name: str) -> None:
    orchestrator = Orchestrator.__new__(Orchestrator)
    orchestrator.work_dir = tmp_path / "work"

    with pytest.raises(ValueError, match="experiment"):
        orchestrator._save_experiment(name, {})
    assert not (tmp_path / "outside").exists()


def test_smart_exporter_supports_wide_and_duplicate_columns(tmp_path: Path) -> None:
    columns = [f"COL_{index}" for index in range(42)]
    columns[27] = columns[0]
    wide = pd.DataFrame([[None] * len(columns)], columns=columns)
    exporter = SmartExporter({"output_directory": str(tmp_path)})

    exported_paths = exporter.export({"wide": wide}, "multi", format="auto")
    assert isinstance(exported_paths, list) and len(exported_paths) == 1
    exported = Path(exported_paths[0])

    assert exported.suffix == ".xlsx"
    workbook = load_workbook(exported, read_only=False, data_only=False)
    try:
        sheet = workbook["wide"]
        assert sheet.max_column == 42
        assert sheet.column_dimensions["AA"].width is not None
        assert sheet.column_dimensions["AP"].width is not None
    finally:
        workbook.close()


def test_disk_lsh_invalidates_index_and_candidates_when_corpus_changes(
    tmp_path: Path,
) -> None:
    profile = {
        "lsh_permutations": 64,
        "lsh_threshold": 0.5,
        "lsh_ngram": 3,
        "force_disk_results": False,
        "enable_nit_blocking": False,
        "max_bucket_size": 500,
    }
    unique_a = [hashlib.sha256(f"A-{index}".encode()).hexdigest() for index in range(180)]
    unique_b = [hashlib.sha256(f"B-{index}".encode()).hexdigest() for index in range(180)]
    first_names = ["ALFA EMPRESA"] * 20 + unique_a
    second_names = unique_b + ["BETA EMPRESA"] * 20
    first = pd.DataFrame({"NOMBRE_LIMPIO": first_names})
    second = pd.DataFrame({"NOMBRE_LIMPIO": second_names})
    engine = DiskBasedLSHEngine(profile=profile)

    first_pairs = engine.find_candidates(first, output_dir=str(tmp_path))
    second_pairs = engine.find_candidates(second, output_dir=str(tmp_path))

    assert isinstance(first_pairs, set)
    assert isinstance(second_pairs, set)
    assert (0, 1) in first_pairs
    assert (0, 1) not in second_pairs
    assert (180, 181) in second_pairs


def test_disk_lsh_cross_source_uses_positions_not_dataframe_labels(tmp_path: Path) -> None:
    frame = pd.DataFrame(
        {
            "NOMBRE_LIMPIO": ["EMPRESA IDENTICA"] * 3,
            "FUENTE": ["A", "A", "B"],
        },
        index=[10, 20, 30],
    )
    engine = DiskBasedLSHEngine(
        profile={
            "lsh_permutations": 16,
            "lsh_threshold": 0.4,
            "lsh_ngram": 3,
            "enable_nit_blocking": False,
        }
    )

    pairs = engine.find_candidates(frame, output_dir=str(tmp_path), cross_source_only=True)

    assert isinstance(pairs, set)
    assert pairs == {(0, 2), (1, 2)}


def test_sourcespec_reserved_looking_column_cannot_control_clusters(tmp_path: Path) -> None:
    source_path = tmp_path / "reserved.csv"
    pd.DataFrame(
        {
            "tax": ["900111111", "900222222", "900333333"],
            "company": ["ALFA UNICA", "BETA UNICA", "GAMMA UNICA"],
            "_original_idx": [0, 0, 0],
        }
    ).to_csv(source_path, index=False)
    spec = SourceSpec(
        name="SOURCE",
        path=source_path,
        column_mapping={"NIT": "tax", "RAZON_SOCIAL": "company"},
        passthrough_columns=("_original_idx",),
    )

    result = linkage(
        [spec],
        profile="prueba_rapida",
        work_dir=str(tmp_path / "run"),
        skip_reporting=True,
    )

    assert len(result["correlative"]) == 3
    assert result["correlative"]["ID_GRUPO"].nunique() == 3
    assert len(result["golden"]) == 3

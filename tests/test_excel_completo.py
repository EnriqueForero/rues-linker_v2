"""F1.11 · Excel completo hasta 1.048.576 filas o ``<base>_LEEME.xlsx``; nunca un recorte.

Qué congela
-----------
* ``exporters.excel.escribir_excel_o_leeme``:
  (a) 1.100.000 filas sintéticas de 3 columnas → NO se escribe ningún xlsx
      de datos; se escribe ``correlativa_LEEME.xlsx`` con el conteo de filas
      y el nombre del parquet. Se mide tiempo y RAM (``MuestreadorRecursos``)
      y se exige que la escritura no duplique la tabla en memoria;
  (b) 120.000 filas → ``correlativa.xlsx`` completo con 120.000 filas de
      datos (contadas con openpyxl ``read_only=True``), también cuando la
      fuente es un ``pq.ParquetFile`` (flujo desde disco) y con hoja LEEME
      delante (alias de v1);
  (c) un valor que empieza por ``=`` o trae ``\\x1a`` se neutraliza, un
      ausente queda como celda vacía y un número negativo sigue siendo número;
* ``DataExportStrategy`` (alias de v1 de L6) ya NO escribe ``_MUESTRA_<n>k``:
  por encima del límite deja ``tabla_correlativa_LEEME.xlsx`` /
  ``golden_records_LEEME.xlsx`` y lo registra en ``omitidos``; la perilla
  ``export_settings.excel_max_rows`` ya no recorta nada y se avisa;
* en ``src/`` no queda ningún ``_MUESTRA_`` ni ``excel_limit``: una sola
  regla (``LIMITE_FILAS_EXCEL`` del escritor).

Por qué ``pd.ExcelWriter(engine="xlsxwriter", constant_memory=True)`` NO se usa
tal cual: pandas escribe las celdas COLUMNA a columna y ``constant_memory``
descarta cualquier celda de una fila ya volcada, así que el libro queda con
una sola columna completa (medido: 3 columnas → solo la última con datos). El
módulo escribe las filas con ``xlsxwriter`` directamente, por lotes; la prueba
(b) es la que lo vigila.

Los datos son sintéticos (numpy); ninguna empresa real entra aquí.
"""

from __future__ import annotations

import logging
import re
import time
from pathlib import Path

import numpy as np
import pandas as pd
import pyarrow.parquet as pq
import pytest
from openpyxl import load_workbook

from record_linkage.evaluation.banco import MuestreadorRecursos
from record_linkage.exporters import escritor, excel
from record_linkage.exporters.excel import LIMITE_FILAS_EXCEL, escribir_excel_o_leeme
from record_linkage.reporting import strategies
from record_linkage.reporting.strategies import DataExportStrategy

RAIZ = Path(__file__).resolve().parent.parent
SRC = RAIZ / "src" / "record_linkage"
N_GRANDE = 1_100_000
N_CABE = 120_000


def _sintetico(n: int, semilla: int = 42) -> pd.DataFrame:
    """n filas × 3 columnas inventadas: entero, decimal y texto (sin nombres reales)."""
    rng = np.random.default_rng(semilla)
    enteros = rng.integers(0, 1_000_000, size=n)
    return pd.DataFrame(
        {
            "ID": np.arange(n, dtype=np.int64),
            "VALOR": rng.random(n),
            "CODIGO": pd.Series(enteros).astype("string").radd("C-"),
        }
    )


def _contar_filas(ruta: Path, hoja: str) -> int:
    libro = load_workbook(ruta, read_only=True)
    try:
        return sum(1 for _ in libro[hoja].iter_rows(values_only=True))
    finally:
        libro.close()


def _texto_leeme(ruta: Path) -> str:
    libro = load_workbook(ruta, read_only=True)
    try:
        assert libro.sheetnames == ["LEEME"]
        celdas = [str(c) for fila in libro["LEEME"].iter_rows(values_only=True) for c in fila if c]
    finally:
        libro.close()
    return "\n".join(celdas)


# ─────────────────────────────────────────────────────────────────────────────
# (a) No cabe → LEEME, sin recorte
# ─────────────────────────────────────────────────────────────────────────────


def test_limite_es_el_del_escritor() -> None:
    """Una sola regla: la constante vive en ``exporters.excel`` y el escritor la reexporta."""
    assert LIMITE_FILAS_EXCEL == 1_048_575
    assert escritor.LIMITE_FILAS_EXCEL is LIMITE_FILAS_EXCEL


def test_1_1_millones_de_filas_no_escribe_xlsx_de_datos_sino_leeme(tmp_path: Path) -> None:
    df = _sintetico(N_GRANDE)
    assert len(df) > LIMITE_FILAS_EXCEL
    mib_tabla = df.memory_usage(deep=True).sum() / 1024**2
    ruta = tmp_path / "correlativa.xlsx"

    inicio = time.perf_counter()
    with MuestreadorRecursos(0.05) as muestreador:
        escrito = escribir_excel_o_leeme(df, ruta)
    segundos = time.perf_counter() - inicio

    assert escrito == tmp_path / "correlativa_LEEME.xlsx"
    assert not ruta.exists(), "no se escribe ningún xlsx de datos: ni completo ni recortado"
    assert not list(tmp_path.glob("*MUESTRA*"))
    assert sorted(p.name for p in tmp_path.iterdir()) == ["correlativa_LEEME.xlsx"]
    texto = _texto_leeme(escrito)
    assert "1.100.000" in texto
    assert "correlativa.parquet" in texto
    assert "read_parquet" in texto and "duckdb" in texto.lower() and "Power Query" in texto
    assert "1.048.576" in texto
    # Se mide y se reporta: el LEEME no toca la tabla, así que no la duplica.
    crecimiento = muestreador.pico_mib - muestreador.inicial_mib
    assert crecimiento < mib_tabla, (
        f"la escritura duplicó la tabla en RAM: +{crecimiento:.0f} MiB (tabla {mib_tabla:.0f} MiB)"
    )
    assert segundos < 60, f"{segundos:.1f} s: márquela slow"


def test_justo_en_el_limite_cabe_y_una_mas_no(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(excel, "LIMITE_FILAS_EXCEL", 5)
    cabe = escribir_excel_o_leeme(_sintetico(5), tmp_path / "a.xlsx")
    assert cabe == tmp_path / "a.xlsx" and _contar_filas(cabe, "DATOS") == 6
    no_cabe = escribir_excel_o_leeme(_sintetico(6), tmp_path / "b.xlsx")
    assert no_cabe == tmp_path / "b_LEEME.xlsx" and not (tmp_path / "b.xlsx").exists()
    assert "b.parquet" in _texto_leeme(no_cabe)
    # ``limite`` explícito manda sobre la constante.
    explicito = escribir_excel_o_leeme(_sintetico(6), tmp_path / "c.xlsx", limite=6)
    assert explicito == tmp_path / "c.xlsx"


# ─────────────────────────────────────────────────────────────────────────────
# (b) Cabe → Excel completo, en flujo
# ─────────────────────────────────────────────────────────────────────────────


def test_120k_filas_escribe_el_excel_completo(tmp_path: Path) -> None:
    df = _sintetico(N_CABE)
    ruta = tmp_path / "correlativa.xlsx"

    inicio = time.perf_counter()
    with MuestreadorRecursos(0.05) as muestreador:
        escrito = escribir_excel_o_leeme(df, ruta)
    segundos = time.perf_counter() - inicio

    assert escrito == ruta and ruta.is_file()
    assert not (tmp_path / "correlativa_LEEME.xlsx").exists()
    libro = load_workbook(ruta, read_only=True)
    try:
        assert libro.sheetnames == ["DATOS"]
        filas = libro["DATOS"].iter_rows(values_only=True)
        assert next(filas) == ("ID", "VALOR", "CODIGO")
        primera = next(filas)
        assert primera[0] == 0 and isinstance(primera[1], float) and primera[2].startswith("C-")
        n = 2 + sum(1 for _ in filas)
    finally:
        libro.close()
    assert n == N_CABE + 1  # 120.000 de datos + encabezado
    # Todas las columnas llegan completas (constant_memory con pandas las perdía).
    assert _contar_filas(ruta, "DATOS") == N_CABE + 1
    ultima = pd.read_excel(ruta, sheet_name="DATOS", skiprows=N_CABE, nrows=1, header=None)
    assert ultima.iloc[0, 0] == N_CABE - 1 and str(ultima.iloc[0, 2]).startswith("C-")
    assert segundos < 60, f"{segundos:.1f} s"
    assert muestreador.muestras >= 1


def test_desde_parquet_en_flujo_y_con_leeme_delante(tmp_path: Path) -> None:
    """La fuente puede ser un ``pq.ParquetFile`` (L6 desde disco): lotes, no la tabla entera."""
    df = _sintetico(N_CABE)
    parquet = tmp_path / "tabla_correlativa.parquet"
    escritor.escribir_parquet(df, parquet)
    del df
    ruta = tmp_path / "tabla_correlativa.xlsx"
    escrito = escribir_excel_o_leeme(
        pq.ParquetFile(parquet),
        ruta,
        leeme=escritor.leeme_alias_v1("excel/correlativa.xlsx"),
        filas_por_lote=7_000,
    )
    assert escrito == ruta
    libro = load_workbook(ruta, read_only=True)
    try:
        assert libro.sheetnames == ["LEEME", "DATOS"]
        assert "excel/correlativa.xlsx" in "\n".join(
            str(c) for fila in libro["LEEME"].iter_rows(values_only=True) for c in fila if c
        )
    finally:
        libro.close()
    assert _contar_filas(ruta, "DATOS") == N_CABE + 1
    # Por encima del límite, el parquet también va a LEEME (sin leerlo entero).
    no_cabe = escribir_excel_o_leeme(pq.ParquetFile(parquet), ruta, limite=10)
    assert no_cabe == tmp_path / "tabla_correlativa_LEEME.xlsx"
    assert "tabla_correlativa.parquet" in _texto_leeme(no_cabe)
    assert "120.000" in _texto_leeme(no_cabe)


def test_tabla_vacia_deja_solo_el_encabezado(tmp_path: Path) -> None:
    ruta = escribir_excel_o_leeme(_sintetico(0), tmp_path / "vacia.xlsx", hoja="golden")
    assert _contar_filas(ruta, "golden") == 1


# ─────────────────────────────────────────────────────────────────────────────
# (c) Neutralización: una regla escrita una vez (prepare_spreadsheet_data)
# ─────────────────────────────────────────────────────────────────────────────


def test_formulas_controles_y_ausentes(tmp_path: Path) -> None:
    df = pd.DataFrame(
        {
            "=HEADER": ["=2+2", "COMPA\x1aIA", None, "normal"],
            "N": [-7, 1, 2, 3],
            "F": [1.5, np.nan, 2.5, np.inf],
            "T": pd.to_datetime(
                ["2026-10-06 14:30:59", None, "2026-01-01", "2026-01-02"], format="ISO8601"
            ),
            "B": [True, False, True, False],
        }
    )
    copia = df.copy(deep=True)
    ruta = escribir_excel_o_leeme(df, tmp_path / "x.xlsx")
    pd.testing.assert_frame_equal(df, copia)  # la fuente no se muta
    libro = load_workbook(ruta, read_only=True, data_only=False)
    try:
        hoja = libro["DATOS"]
        filas = list(hoja.iter_rows(values_only=True))
    finally:
        libro.close()
    assert filas[0][0] == "'=HEADER"
    assert filas[1][0] == "'=2+2"
    assert filas[2][0] == "COMPAIA"  # el \x1a desaparece
    assert filas[3][0] is None  # ausente → celda vacía, no "nan" ni "None"
    assert filas[1][1] == -7 and isinstance(filas[1][1], int)
    assert filas[2][2] is None  # NaN → vacía
    assert filas[4][2] == "inf"  # ±inf → texto, como to_csv; nunca la fórmula =1/0
    assert filas[2][3] is None  # NaT → vacía
    assert str(filas[1][3]).startswith("2026-10-06 14:30:59")
    assert filas[1][4] is True and filas[2][4] is False
    # openpyxl no marca ninguna celda como fórmula.
    libro = load_workbook(ruta, read_only=False, data_only=False)
    try:
        assert all(c.data_type != "f" for fila in libro["DATOS"].iter_rows() for c in fila)
    finally:
        libro.close()


# ─────────────────────────────────────────────────────────────────────────────
# L6 (alias de v1): nunca más _MUESTRA_<n>k
# ─────────────────────────────────────────────────────────────────────────────


def _archivos(files: list[Path]) -> list[str]:
    return sorted(p.name for p in files)


def test_l6_desde_memoria_no_recorta(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(excel, "LIMITE_FILAS_EXCEL", 10)
    estrategia = DataExportStrategy()
    log = logging.getLogger(__name__)
    files = estrategia._export_from_memory(_sintetico(11), "tabla_correlativa", tmp_path, log)
    assert _archivos(files) == [
        "tabla_correlativa.csv.gz",
        "tabla_correlativa.parquet",
        "tabla_correlativa_LEEME.xlsx",
    ]
    assert not list(tmp_path.glob("*MUESTRA*"))
    texto = _texto_leeme(tmp_path / "tabla_correlativa_LEEME.xlsx")
    assert "tabla_correlativa.parquet" in texto and "11" in texto
    assert "excel/correlativa.xlsx" in texto  # el alias remite al archivo nuevo
    assert [o.artefacto for o in estrategia.omitidos] == ["tabla_correlativa.xlsx"]
    assert "LEEME" in estrategia.omitidos[0].motivo
    # Cabe → completo, con LEEME delante (alias).
    files = estrategia._export_from_memory(_sintetico(10), "golden_records", tmp_path, log)
    assert "golden_records.xlsx" in _archivos(files)
    assert _contar_filas(tmp_path / "golden_records.xlsx", "DATOS") == 11


def test_l6_desde_disco_no_recorta_ni_lee_entero(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(excel, "LIMITE_FILAS_EXCEL", 10)
    fuente = tmp_path / "golden.parquet"
    escritor.escribir_parquet(_sintetico(11), fuente)
    estrategia = DataExportStrategy()
    log = logging.getLogger(__name__)
    files = estrategia._export_from_disk_streaming(fuente, "golden_records", tmp_path, log)
    assert _archivos(files) == [
        "golden_records.csv.gz",
        "golden_records.parquet",
        "golden_records_LEEME.xlsx",
    ]
    assert not list(tmp_path.glob("*MUESTRA*"))
    assert "golden_records.parquet" in _texto_leeme(tmp_path / "golden_records_LEEME.xlsx")
    assert [o.artefacto for o in estrategia.omitidos] == ["golden_records.xlsx"]


def test_excel_max_rows_ya_no_recorta_y_se_avisa(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    """La perilla de v1 pedía recortar a N filas; ahora no hace nada y se dice."""
    strategies._PERILLA_EXCEL_AVISADA = False
    ctx = strategies.ReportingContext(
        correlative_df=_sintetico(30),
        golden_df=_sintetico(3),
        metrics={},
        config={"export_settings": {"excel_max_rows": 20}, "reporting_use_checkpoints": False},
        output_dir=tmp_path,
        start_time=time.time(),
    )
    with caplog.at_level(logging.WARNING, logger="record_linkage.reporting.strategies"):
        files = DataExportStrategy().execute(ctx, logging.getLogger(__name__))
    assert "tabla_correlativa.xlsx" in _archivos(files)
    assert _contar_filas(tmp_path / "tabla_correlativa.xlsx", "DATOS") == 31
    assert not list(tmp_path.glob("*MUESTRA*"))
    assert any("excel_max_rows" in r.getMessage() for r in caplog.records)
    assert not hasattr(DataExportStrategy, "EXCEL_ROW_LIMIT")


def test_ningun_modulo_de_src_recorta_ni_tiene_limite_propio() -> None:
    """``grep -rn MUESTRA_100k src`` → 0; ``excel_limit`` tampoco. Una sola regla."""
    # ``_MUESTRA_`` seguido de dígito, ``*``, ``<`` o ``{``: el nombre del recorte
    # (``PERILLAS_MUESTRA_RETIRADAS`` de reports.py no lo es).
    patron = re.compile(r"_MUESTRA_[\d*<{]|MUESTRA_100k|excel_limit|EXCEL_ROW_LIMIT")
    hallazgos = {
        p.relative_to(SRC).as_posix(): m.group(0)
        for p in sorted(SRC.rglob("*.py"))
        if (m := patron.search(p.read_text(encoding="utf-8")))
    }
    assert not hallazgos, hallazgos
    perfiles = (SRC / "config" / "profiles.py").read_text(encoding="utf-8")
    interno = (SRC / "pipeline" / "_internal.py").read_text(encoding="utf-8")
    assert '"excel_max_rows"' not in perfiles  # ya no es clave de ningún perfil
    assert '"excel_max_rows"' not in interno

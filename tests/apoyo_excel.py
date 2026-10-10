"""Apoyo compartido por las pruebas que leen libros ``.xlsx``.

Vive aquí —y no copiado en ``test_escritor.py`` y ``test_exportar_vistas.py``—
para que la lectura de un libro hoja por hoja se escriba una sola vez
(pulido de F2, hallazgo 3 de F2.16).
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


def hojas_del_libro(ruta: Path) -> dict[str, list[list[object]]]:
    """``{nombre de hoja: filas}`` de un ``.xlsx``, con las fórmulas sin evaluar."""
    libro = load_workbook(ruta, read_only=True, data_only=False)
    try:
        return {
            hoja.title: [list(fila) for fila in hoja.iter_rows(values_only=True)]
            for hoja in libro.worksheets
        }
    finally:
        libro.close()

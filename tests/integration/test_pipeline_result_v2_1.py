"""tests/integration/test_pipeline_result_v2_1.py

Tests para las features nuevas de v2.1.0 sobre PipelineResult:
    - F6.4: métodos to_excel() y to_csv() — desde F2.11 son ALIAS de la función
      libre del estándar ``exporters.escritor.exportar_vistas`` (avisan con
      ``DeprecationWarning``); aquí se congela que el alias conserva su contrato.
    - F6.5: items(), values(), __iter__, __len__
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from record_linkage.pipeline.result import PipelineResult


@pytest.fixture
def result_with_data(tmp_path: Path) -> PipelineResult:
    """PipelineResult con dos parquets reales en disco."""
    golden = pd.DataFrame({"ID_GRUPO": [1, 2, 3], "NIT_FINAL": ["900", "800", "700"]})
    correlative = pd.DataFrame(
        {
            "ID_GRUPO": [1, 1, 2, 3],
            "NIT": ["900", "900", "800", "700"],
            "RAZON_SOCIAL": ["ACME", "ACME SAS", "BETA", "GAMMA"],
        }
    )
    golden.to_parquet(tmp_path / "golden.parquet")
    correlative.to_parquet(tmp_path / "correlative.parquet")

    return PipelineResult(
        work_dir=tmp_path,
        golden_path=tmp_path / "golden.parquet",
        correlative_path=tmp_path / "correlative.parquet",
        extra={"execution_seconds": 12.5},
    )


# ─────────────────────────────────────────────────────────────────────────────
# F6.5 — Protocolo Mapping completo
# ─────────────────────────────────────────────────────────────────────────────
def test_items_yields_pairs(result_with_data: PipelineResult):
    items = list(result_with_data.items())
    keys = [k for k, _ in items]
    assert "golden_records" in keys
    assert "correlative_table" in keys
    assert "execution_seconds" in keys


def test_values_yields_values(result_with_data: PipelineResult):
    values = list(result_with_data.values())
    assert any(isinstance(v, pd.DataFrame) for v in values)
    # Verificar que el valor del `extra` está presente sin usar `in`
    # (lista con DataFrames no soporta `in` por la ambigüedad de bool(DataFrame))
    scalar_values = [v for v in values if not isinstance(v, pd.DataFrame)]
    assert 12.5 in scalar_values  # del `extra`


def test_iter_yields_keys(result_with_data: PipelineResult):
    keys = list(result_with_data)
    assert "golden_records" in keys
    assert "execution_seconds" in keys


def test_len_counts_available_keys(result_with_data: PipelineResult):
    n = len(result_with_data)
    # Al menos: golden_records, correlative_table, execution_seconds
    assert n >= 3


def test_dict_conversion(result_with_data: PipelineResult):
    """`dict(result)` debe funcionar gracias al protocolo Mapping."""
    d = dict(result_with_data)
    assert "golden_records" in d
    assert isinstance(d["golden_records"], pd.DataFrame)


# ─────────────────────────────────────────────────────────────────────────────
# F6.4 — Exportadores (alias de ``exportar_vistas`` desde F2.11)
# ─────────────────────────────────────────────────────────────────────────────
def test_to_excel_y_to_csv_son_alias_de_exportar_vistas(
    result_with_data: PipelineResult, tmp_path: Path
):
    """Avisan con DeprecationWarning que nombra la función libre del estándar."""
    with pytest.warns(DeprecationWarning, match="exportar_vistas"):
        result_with_data.to_excel(tmp_path / "alias.xlsx")
    with pytest.warns(DeprecationWarning, match="exportar_vistas"):
        result_with_data.to_csv(tmp_path / "alias_csv")


def test_to_excel_exige_xlsx_y_openpyxl(result_with_data: PipelineResult, tmp_path: Path):
    """El alias no escribe otra cosa que lo que ``exportar_vistas`` escribe."""
    with pytest.raises(ValueError, match=r"\.xlsx"):
        result_with_data.to_excel(tmp_path / "salida.xls")
    with pytest.raises(ValueError, match="openpyxl"):
        result_with_data.to_excel(tmp_path / "salida.xlsx", engine="xlsxwriter")


def test_to_csv_exige_utf8(result_with_data: PipelineResult, tmp_path: Path):
    with pytest.raises(ValueError, match=r"(?i)utf-8"):
        result_with_data.to_csv(tmp_path / "csvs", encoding="latin-1")


def test_to_excel_creates_file(result_with_data: PipelineResult, tmp_path: Path):
    out_path = tmp_path / "salida.xlsx"
    returned = result_with_data.to_excel(out_path)
    assert out_path.exists()
    assert returned == out_path.resolve()

    # Verificar contenido: dos hojas
    sheets = pd.read_excel(out_path, sheet_name=None)
    assert "golden_records" in sheets
    assert "correlative_table" in sheets
    assert len(sheets["golden_records"]) == 3
    assert len(sheets["correlative_table"]) == 4


def test_to_excel_with_custom_include(result_with_data: PipelineResult, tmp_path: Path):
    """Solo se exporta lo pedido en `include`."""
    out_path = tmp_path / "solo_golden.xlsx"
    result_with_data.to_excel(out_path, include=("golden_records",))
    sheets = pd.read_excel(out_path, sheet_name=None)
    assert "golden_records" in sheets
    assert "correlative_table" not in sheets


def test_to_excel_raises_on_empty_include(result_with_data: PipelineResult, tmp_path: Path):
    """Si `include` no resuelve a ningún DataFrame, debe levantar ValueError."""
    with pytest.raises(ValueError, match=r"(?i)ninguna"):
        result_with_data.to_excel(tmp_path / "x.xlsx", include=("clave_inexistente",))


def test_to_csv_creates_files(result_with_data: PipelineResult, tmp_path: Path):
    out_dir = tmp_path / "csvs"
    written = result_with_data.to_csv(out_dir)
    assert "golden_records" in written
    assert "correlative_table" in written
    assert (out_dir / "golden_records.csv").exists()
    assert (out_dir / "correlative_table.csv").exists()

    # Verificar contenido
    golden_back = pd.read_csv(out_dir / "golden_records.csv")
    assert len(golden_back) == 3


def test_to_csv_raises_on_empty_include(result_with_data: PipelineResult, tmp_path: Path):
    with pytest.raises(ValueError, match=r"(?i)ninguna"):
        result_with_data.to_csv(tmp_path / "csvs", include=("inexistente",))


def test_vistas_vacias_fallan_con_el_formato_del_modulo_de_errores(
    result_with_data: PipelineResult, tmp_path: Path
):
    """El error de ``_vistas`` se arma con ``mensaje_accionable`` (qué pasó · por
    qué importa · qué hacer) y cita las claves disponibles."""
    with pytest.raises(ValueError, match=r"\nQué hacer: ") as exc:
        result_with_data.to_csv(tmp_path / "csvs", include=("inexistente",))
    texto = str(exc.value)
    assert texto.startswith("Qué pasó: to_csv: ninguna de las claves")
    assert "golden_records" in texto and "exportar_vistas" in texto


def test_to_csv_con_indice_que_choca_con_una_columna_falla_con_remedio(tmp_path: Path):
    """Con ``index=True`` y un índice llamado como una columna, ``reset_index``
    lanzaría el ``ValueError`` crudo de pandas («cannot insert a, already
    exists»); el alias lo detecta antes y dice qué hacer. No deja archivo."""
    df = pd.DataFrame({"a": [1, 2], "b": [3, 4]}, index=pd.Index([10, 20], name="a"))
    result = PipelineResult(work_dir=tmp_path, extra={"payload": df})
    with pytest.raises(ValueError, match="rename_axis") as exc:
        result.to_csv(tmp_path / "csvs", include=("payload",), index=True)
    texto = str(exc.value)
    assert texto.startswith("Qué pasó: ") and "index=False" in texto
    assert "'payload'" in texto and "['a']" in texto
    assert not (tmp_path / "csvs" / "payload.csv").exists()


def test_to_csv_con_indice_sin_nombre_lo_llama_index(tmp_path: Path):
    """Documentado: con ``index=True`` un índice sin nombre sale como la
    columna ``index`` (lo que hace ``reset_index``), no como encabezado vacío."""
    result = PipelineResult(work_dir=tmp_path, extra={"payload": pd.DataFrame({"a": [1, 2]})})
    ruta = result.to_csv(tmp_path / "csvs", include=("payload",), index=True)["payload"]
    assert ruta.read_text(encoding="utf-8").splitlines()[0] == "index,a"


def test_exports_neutralize_spreadsheet_formulas_without_mutating_input(tmp_path: Path):
    dangerous = pd.DataFrame(
        {
            "=HEADER": ["=2+2", "+cmd", "-cmd", "@SUM(A1:A2)", "safe"],
            "number": [-7, 1, 2, 3, 4],
        },
        index=pd.Index(["=INDEX", "b", "c", "d", "e"], name="@INDEX_NAME"),
    )
    before = dangerous.copy(deep=True)
    result = PipelineResult(work_dir=tmp_path, extra={"payload": dangerous})

    xlsx_path = result.to_excel(tmp_path / "safe.xlsx", include=("payload",))
    workbook = load_workbook(xlsx_path, read_only=True, data_only=False)
    try:
        sheet = workbook["payload"]
        assert sheet["A1"].value == "'=HEADER"
        assert sheet["A2"].value == "'=2+2"
        assert sheet["A2"].data_type != "f"
        assert sheet["B2"].value == -7
    finally:
        workbook.close()

    csv_path = result.to_csv(tmp_path / "csvs", include=("payload",), index=True)["payload"]
    csv_text = csv_path.read_text(encoding="utf-8")
    assert "'@INDEX_NAME" in csv_text
    assert "'=INDEX" in csv_text
    assert "'=HEADER" in csv_text
    assert "'=2+2" in csv_text
    assert "'+cmd" in csv_text

    pd.testing.assert_frame_equal(dangerous, before)


@pytest.mark.parametrize("key", ["../escape", "sub/escape", r"sub\escape", "bad\x00key"])
def test_to_csv_rejects_include_path_traversal_and_controls(tmp_path: Path, key: str):
    result = PipelineResult(work_dir=tmp_path, extra={key: pd.DataFrame({"x": [1]})})

    with pytest.raises(ValueError, match=r"nombre simple|control|directorio"):
        result.to_csv(tmp_path / "out", include=(key,))

    assert not (tmp_path / "escape.csv").exists()


def test_to_excel_sanitizes_and_deduplicates_sheet_names(tmp_path: Path):
    prefix = "x" * 31
    frames = {
        f"{prefix}primero": pd.DataFrame({"x": [1]}),
        f"{prefix}segundo": pd.DataFrame({"x": [2]}),
        "bad/name:*?[]": pd.DataFrame({"x": [3]}),
        "control\x01name": pd.DataFrame({"x": [4]}),
    }
    result = PipelineResult(work_dir=tmp_path, extra=frames)

    path = result.to_excel(tmp_path / "sheets.xlsx", include=tuple(frames))
    workbook = load_workbook(path, read_only=True)
    try:
        names = workbook.sheetnames
        assert len(names) == len(frames)
        assert len({name.casefold() for name in names}) == len(names)
        assert all(len(name) <= 31 for name in names)
        assert all(not any(char in name for char in "[]:*?/\\") for name in names)
        assert all(not any(ord(char) < 32 for char in name) for name in names)
    finally:
        workbook.close()

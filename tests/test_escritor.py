"""Escritor único de la carpeta del estándar (F1.10): forma exacta, atómico, alias v1.

Qué congela
-----------
Sobre el conjunto sintético de ``test_contrato_salida`` (64 filas inventadas,
dos fuentes) y una corrida de ``linkage()`` sin L6:

* la carpeta ``<AAAA-MM-DD_HHMM>_<nombre>/`` tiene la forma EXACTA del
  estándar (lista fija de rutas relativas);
* ``manifest.json`` lista cada artefacto con su SHA-256 y tamaño reales y
  trae los bloques del estándar (insumos, conteos, invariantes, parámetros,
  tiempos y RSS por fase, renombres, prioridad de fuentes, omitidos);
* dos escrituras del mismo resultado con la misma ``marca_tiempo`` producen
  las mismas huellas en los artefactos deterministas (parquet, csv). Los
  ``.xlsx`` quedan fuera de la comparación: openpyxl escribe la fecha de
  creación en ``docProps/core.xml`` y el zip cambia aunque el contenido sea
  el mismo; por eso el manifiesto los registra pero la prueba no los exige;
* una escritura que falla a mitad (``to_parquet`` del golden) no deja carpeta
  definitiva ni pendiente;
* ``leer_resultado`` devuelve un ``ResultadoLinkage`` que ``validar()``
  acepta, aplica los alias en español y detecta un artefacto alterado;
* ``revision.csv`` y ``diccionario.csv`` son FIELES a los datos (sin la
  neutralización de hoja de cálculo): lo que se escribe es lo que
  ``leer_resultado`` devuelve, también con nombres que empiezan por
  ``=``, ``+``, ``-`` o ``@``;
* un ``_trabajo/`` ajeno en la pendiente (``res.dir_trabajo`` fuera) hace
  fallar la escritura en vez de publicarse como si fuera de esta corrida
  (también cuando ``res.dir_trabajo`` es None, sin decir «en None»);
* en ``revision.csv`` una celda vacía y un ausente son lo mismo: el CSV no
  los distingue, y ``leer_resultado`` devuelve ``pd.NA`` (lo que el contrato
  llama ausente), nunca ``''``;
* dos figuras con el mismo nombre no se escriben (``figuras/`` es plano):
  la escritura falla antes de tocar el disco;
* tras el ``rename`` de la pendiente ninguna ruta del resultado
  (``metricas['report_files']``, ``manifiesto['columnas_tecnicas']``, el
  ``origen`` de SCORE_PAR) ni del ``manifest.json`` publicado apunta a
  ``.<nombre>.pendiente/``: en memoria quedan bajo la carpeta definitiva y
  en el JSON relativas a ``_trabajo/``;
* un ``_trabajo/manifest.json`` ilegible deja ``tiempos_por_fase`` vacío y
  lo declara en ``omitidos``;
* ``linkage(carpeta_salida=...)`` escribe la carpeta y deja ``_trabajo/``
  dentro; con L6 activo los alias de v1 se escriben con ``DeprecationWarning``
  y su ``.xlsx`` lleva una primera hoja ``LEEME``;
* los alias parquet de v1 NO llevan el metadato ``contrato`` (no tienen su
  forma); el aviso de retiro sale también por ``logging`` (Colab silencia
  los ``DeprecationWarning``) y ``reporte_*.xlsx`` conserva la hoja
  ``Sheet1`` de v1;
* ningún módulo de ``src/`` fuera de la lista declarada llama a
  ``to_parquet``/``to_excel``/``to_csv``/``write_table``/``ExcelWriter``/
  ``ParquetWriter``/``write_to_dataset``/``Workbook``.

Las empresas son inventadas. Ningún dato licenciado entra aquí.
"""

from __future__ import annotations

import ast
import dataclasses
import hashlib
import json
import logging
import warnings
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
import pytest
from openpyxl import load_workbook
from test_contrato_salida import _correr_linkage, _silencio, conjunto_sintetico

import record_linkage as rl
from record_linkage import contrato
from record_linkage.exporters import escritor
from record_linkage.exporters.escritor import (
    LIMITE_FILAS_EXCEL,
    VERSION_RETIRO_ALIAS_V1,
    Manifiesto,
    escribir_resultado,
    leer_resultado,
)
from record_linkage.pipeline.errores import ContratoSalidaError, ErrorRuesLinker
from record_linkage.resultado import ResultadoLinkage

RAIZ = Path(__file__).resolve().parent.parent
SRC = RAIZ / "src" / "record_linkage"
MARCA = datetime(2026, 10, 6, 14, 30, 59)
CARPETA_ESPERADA = "2026-10-06_1430_prueba"

#: La forma exacta del estándar para una corrida de linkage() con golden,
#: sin enlaces (F3), con Excel y una figura. ``_trabajo/`` no está aquí
#: porque el resultado de la fixture tiene su work_dir fuera de la carpeta.
RUTAS_ESPERADAS = (
    "correlativa.parquet",
    "diccionario.csv",
    "entidades_ids.parquet",
    "excel/correlativa.xlsx",
    "excel/golden.xlsx",
    "figuras/dashboard.png",
    "golden.parquet",
    "manifest.json",
    "revision.csv",
)


@pytest.fixture(scope="module")
def res(tmp_path_factory: pytest.TempPathFactory) -> ResultadoLinkage:
    return _correr_linkage(tmp_path_factory.mktemp("trabajo"))


@pytest.fixture
def figura(tmp_path: Path) -> Path:
    ruta = tmp_path / "dashboard.png"
    ruta.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 16)
    return ruta


def _rutas_relativas(carpeta: Path) -> list[str]:
    return sorted(
        p.relative_to(carpeta).as_posix()
        for p in carpeta.rglob("*")
        if p.is_file() and "_trabajo" not in p.relative_to(carpeta).parts
    )


def _sha256(ruta: Path) -> str:
    return hashlib.sha256(ruta.read_bytes()).hexdigest()


def _con(res: ResultadoLinkage, **cambios: Any) -> ResultadoLinkage:
    """Copia de ``res`` con campos cambiados y su propio ``manifiesto`` (el
    escritor lo muta al publicar; la fixture es de módulo)."""
    return dataclasses.replace(res, manifiesto=dict(res.manifiesto), **cambios)


#: Revisión no vacía cuyos textos empiezan por los prefijos que una hoja de
#: cálculo tomaría por fórmula. Empresas inventadas.
REVISION_CON_PREFIJOS = pd.DataFrame(
    {
        "TIPO": ["cruce", "cruce", "duplicado", "cruce", "cruce"],
        "FUENTE": ["ORBIS", "ORBIS", "RUES", "ORBIS", "ORBIS"],
        "CLAVE_A": ["RUES-1", "RUES-2", "RUES-3", "RUES-4", "RUES-6"],
        "NOMBRE_A": ["=EMPRESA A SAS", "+EMPRESA C", "-EMPRESA E", "@EMPRESA G", "EMPRESA I"],
        "CLAVE_B": ["ORBIS-1", "ORBIS-2", "RUES-5", "ORBIS-4", "ORBIS-6"],
        "NOMBRE_B": ["-EMPRESA B LTDA", "@EMPRESA D", "=EMPRESA F", "+EMPRESA H", "EMPRESA J"],
        "DECISION": ["distinta", "misma_empresa", "distinta", "mismo_grupo", "distinta"],
        "AUTOR": ["asistida", "asistida", "asistida", "asistida", "asistida"],
        # La última fila no tiene RAZON (pd.NA): el CSV la escribe vacía y
        # leer_resultado la devuelve como NA, no como ''.
        "RAZON": ["+no es la misma", "=mismo NIT", "-sede distinta", "@matriz y filial", pd.NA],
    },
    dtype="string",
)


# ─────────────────────────────────────────────────────────────────────────────
# Forma de la carpeta y manifiesto
# ─────────────────────────────────────────────────────────────────────────────


def test_carpeta_tiene_la_forma_exacta(res: ResultadoLinkage, tmp_path: Path, figura: Path) -> None:
    man = escribir_resultado(res, tmp_path, "prueba", marca_tiempo=MARCA, figuras=[figura])
    carpeta = tmp_path / CARPETA_ESPERADA
    assert man.carpeta == carpeta
    assert carpeta.is_dir()
    assert tuple(_rutas_relativas(carpeta)) == RUTAS_ESPERADAS
    assert not (tmp_path / ".prueba.pendiente").exists()


def test_manifest_lista_cada_artefacto_con_sha256_y_bytes_reales(
    res: ResultadoLinkage, tmp_path: Path, figura: Path
) -> None:
    man = escribir_resultado(res, tmp_path, "prueba", marca_tiempo=MARCA, figuras=[figura])
    carpeta = man.carpeta
    texto = json.loads((carpeta / "manifest.json").read_text(encoding="utf-8"))
    assert texto["contrato"] == contrato.VERSION_CONTRATO
    assert texto["version"] == rl.__version__
    assert texto["nombre"] == "prueba"
    assert texto["carpeta"] == CARPETA_ESPERADA  # el nombre definitivo, no la pendiente
    assert texto["marca_tiempo"] == MARCA.isoformat(timespec="seconds")
    listados = {a["ruta"]: a for a in texto["artefactos"]}
    # Todo lo que hay en la carpeta, salvo el propio manifest (no puede
    # contener su propia huella), está listado con su huella y tamaño reales.
    assert set(listados) == set(RUTAS_ESPERADAS) - {"manifest.json"}
    for ruta, art in listados.items():
        archivo = carpeta / ruta
        assert art["bytes"] == archivo.stat().st_size, ruta
        assert art["sha256"] == _sha256(archivo), ruta
    # Bloques del estándar.
    assert texto["insumos"] == res.manifiesto["entradas"]
    assert texto["parametros"] == res.manifiesto["parametros"]
    assert texto["conteos"]["filas"] == len(res.correlativa)
    assert texto["conteos"]["grupos"] == res.correlativa["ID_GRUPO"].nunique()
    assert (
        texto["conteos"]["entidades_con_nit"] + texto["conteos"]["entidades_sin_nit"]
        == texto["conteos"]["entidades"]
    )
    assert texto["invariantes"]["ok"] is True and texto["invariantes"]["fallos"] == []
    # Tiempos y RSS vienen del manifest de _trabajo/ (L1…L5 corrieron).
    assert set(texto["tiempos_por_fase"]) >= {"L1_prep", "L5_golden"}
    assert set(texto["rss_por_fase"]) >= {"L1_prep", "L5_golden"}
    assert texto["renombres"] == {"ID_REGISTRO": "ID_REGISTRO_FUENTE"}
    assert texto["prioridad_fuentes"] == res.manifiesto["completar"]["prioridad_fuentes"]
    assert texto["omitidos"] == []
    assert texto["corrida"]["funcion"] == "linkage"
    assert isinstance(Manifiesto.desde_dict(texto), Manifiesto)


def test_entidades_ids_es_el_crosswalk_de_la_corrida(res: ResultadoLinkage, tmp_path: Path) -> None:
    man = escribir_resultado(res, tmp_path, "prueba", marca_tiempo=MARCA, excel=False)
    ids = pd.read_parquet(man.carpeta / "entidades_ids.parquet")
    assert list(ids.columns) == ["ID_ENTIDAD", "ID_GRUPO", "N_REGISTROS", "RETIRADO_EN"]
    assert len(ids) == res.correlativa["ID_GRUPO"].nunique()
    assert ids["ID_ENTIDAD"].is_unique and ids["ID_GRUPO"].is_unique
    assert ids["RETIRADO_EN"].isna().all()
    assert int(ids["N_REGISTROS"].sum()) == len(res.correlativa)
    esperado = res.correlativa.groupby("ID_GRUPO").size()
    assert ids.set_index("ID_GRUPO")["N_REGISTROS"].sort_index().tolist() == esperado.tolist()


def test_sin_excel_lo_declara_en_omitidos(res: ResultadoLinkage, tmp_path: Path) -> None:
    man = escribir_resultado(res, tmp_path, "prueba", marca_tiempo=MARCA, excel=False)
    assert not (man.carpeta / "excel").exists()
    assert [o["artefacto"] for o in man.omitidos] == [
        "excel/correlativa.xlsx",
        "excel/golden.xlsx",
    ]
    assert all("excel=False" in o["motivo"] for o in man.omitidos)


def test_excel_que_no_cabe_deja_leeme_y_no_recorta(
    res: ResultadoLinkage, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    assert LIMITE_FILAS_EXCEL == 1_048_575  # 1.048.576 filas de hoja menos el encabezado
    monkeypatch.setattr(escritor, "LIMITE_FILAS_EXCEL", 10)
    man = escribir_resultado(res, tmp_path, "prueba", marca_tiempo=MARCA)
    excel = man.carpeta / "excel"
    assert not (excel / "correlativa.xlsx").exists()
    assert not (excel / "golden.xlsx").exists()
    assert sorted(p.name for p in excel.iterdir()) == [
        "correlativa_LEEME.xlsx",
        "golden_LEEME.xlsx",
    ]
    libro = load_workbook(excel / "correlativa_LEEME.xlsx", read_only=True)
    try:
        assert libro.sheetnames == ["LEEME"]
        celdas = [str(c.value) for fila in libro["LEEME"].iter_rows() for c in fila if c.value]
    finally:
        libro.close()
    texto = "\n".join(celdas)
    assert f"{len(res.correlativa):,}".replace(",", ".") in texto
    assert "read_parquet" in texto and "duckdb" in texto.lower() and "Power Query" in texto
    assert {o["artefacto"] for o in man.omitidos} == {
        "excel/correlativa.xlsx",
        "excel/golden.xlsx",
    }
    assert [a["ruta"] for a in man.artefactos if a["ruta"].startswith("excel/")] == [
        "excel/correlativa_LEEME.xlsx",
        "excel/golden_LEEME.xlsx",
    ]


def test_dos_escrituras_producen_las_mismas_huellas(res: ResultadoLinkage, tmp_path: Path) -> None:
    """Los parquet y csv son deterministas; el xlsx no (fecha en el zip)."""
    m1 = escribir_resultado(res, tmp_path / "a", "prueba", marca_tiempo=MARCA)
    m2 = escribir_resultado(res, tmp_path / "b", "prueba", marca_tiempo=MARCA)
    h1 = {a["ruta"]: a["sha256"] for a in m1.artefactos if not a["ruta"].endswith(".xlsx")}
    h2 = {a["ruta"]: a["sha256"] for a in m2.artefactos if not a["ruta"].endswith(".xlsx")}
    assert h1 == h2
    assert set(h1) == {
        "correlativa.parquet",
        "golden.parquet",
        "entidades_ids.parquet",
        "revision.csv",
        "diccionario.csv",
    }
    assert {a["ruta"] for a in m1.artefactos} == {a["ruta"] for a in m2.artefactos}


def test_parquet_lleva_el_esquema_del_contrato(res: ResultadoLinkage, tmp_path: Path) -> None:
    import pyarrow.parquet as pq

    man = escribir_resultado(res, tmp_path, "prueba", marca_tiempo=MARCA, excel=False)
    esquema = pq.read_schema(man.carpeta / "correlativa.parquet")
    assert esquema.names[: len(contrato.COLUMNAS_CORRELATIVA)] == list(
        contrato.COLUMNAS_CORRELATIVA
    )
    for col in contrato.CORRELATIVA:
        assert esquema.field(col.nombre).type == col.tipo, col.nombre
    assert esquema.metadata[b"contrato"] == contrato.VERSION_CONTRATO.encode()
    assert b"pandas" not in esquema.metadata  # sin metadatos variables
    golden = pq.read_schema(man.carpeta / "golden.parquet")
    assert golden.names == list(contrato.COLUMNAS_GOLDEN)
    # Un parquet escrito sin columnas del contrato (los alias de v1, con sus
    # columnas técnicas) NO se etiqueta como contrato 1.0.
    alias = tmp_path / "tabla_correlativa.parquet"
    escritor.escribir_parquet(res.correlativa.assign(NOMBRE_LIMPIO="x"), alias)
    metadatos = pq.read_schema(alias).metadata or {}
    assert b"contrato" not in metadatos
    assert b"pandas" not in metadatos


# ─────────────────────────────────────────────────────────────────────────────
# Atomicidad y fallos
# ─────────────────────────────────────────────────────────────────────────────


def test_fallo_a_mitad_no_deja_carpeta_definitiva_ni_pendiente(
    res: ResultadoLinkage, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """El escritor no usa ``DataFrame.to_parquet`` (escribe con pyarrow para
    fijar el esquema), así que el «to_parquet del golden» que la especificación
    pide reventar es ``_escribir_parquet`` cuando le toca ``golden.parquet``."""
    original = escritor._escribir_parquet

    def revienta(df: pd.DataFrame, ruta: Path, *args: Any, **kwargs: Any) -> None:
        if ruta.name == "golden.parquet":
            raise OSError("disco lleno (simulado)")
        original(df, ruta, *args, **kwargs)

    monkeypatch.setattr(escritor, "_escribir_parquet", revienta)
    with pytest.raises(OSError, match="disco lleno"):
        escribir_resultado(res, tmp_path, "prueba", marca_tiempo=MARCA)
    assert not (tmp_path / CARPETA_ESPERADA).exists()
    assert not (tmp_path / ".prueba.pendiente").exists()
    assert list(tmp_path.iterdir()) == []


def test_fallo_conserva_trabajo_dentro_de_la_pendiente(
    res: ResultadoLinkage, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Si ``_trabajo/`` vive dentro de la pendiente (linkage(carpeta_salida=...)),
    un fallo al escribir NO borra los checkpoints: la siguiente corrida los
    reutiliza. Solo se borra lo que el escritor escribió."""
    pendiente = tmp_path / ".prueba.pendiente"
    trabajo = pendiente / "_trabajo"
    trabajo.mkdir(parents=True)
    (trabajo / "manifest.json").write_text("{}", encoding="utf-8")
    propio = _con(res, dir_trabajo=trabajo)  # como lo deja linkage(carpeta_salida=...)
    monkeypatch.setattr(
        escritor, "_escribir_parquet", lambda *a, **k: (_ for _ in ()).throw(OSError("x"))
    )
    with pytest.raises(OSError):
        escribir_resultado(propio, tmp_path, "prueba", marca_tiempo=MARCA)
    assert not (tmp_path / CARPETA_ESPERADA).exists()
    assert sorted(p.name for p in pendiente.iterdir()) == ["_trabajo"]
    assert (trabajo / "manifest.json").is_file()


def test_trabajo_ajeno_en_la_pendiente_falla_rapido(res: ResultadoLinkage, tmp_path: Path) -> None:
    """Una corrida interrumpida dejó ``.prueba.pendiente/_trabajo``; publicar
    un resultado cuyo ``dir_trabajo`` está en OTRA parte no debe arrastrar
    esos checkpoints a la carpeta definitiva como si fueran suyos."""
    pendiente = tmp_path / ".prueba.pendiente"
    ajeno = pendiente / "_trabajo"
    ajeno.mkdir(parents=True)
    (ajeno / "ajeno.txt").write_text("de otra corrida", encoding="utf-8")
    assert Path(res.dir_trabajo).resolve() != ajeno.resolve()
    with pytest.raises(ErrorRuesLinker, match="_trabajo") as exc:
        escribir_resultado(res, tmp_path, "prueba", marca_tiempo=MARCA, excel=False)
    assert "Qué hacer" in str(exc.value) and "carpeta_salida" in str(exc.value)
    assert str(res.dir_trabajo) in str(exc.value)
    assert not (tmp_path / CARPETA_ESPERADA).exists()
    # Nada se borra: el _trabajo/ ajeno sigue ahí para reanudar o borrar a mano.
    assert (ajeno / "ajeno.txt").is_file()
    assert sorted(p.name for p in pendiente.iterdir()) == ["_trabajo"]
    # Sin dir_trabajo el mensaje lo dice así, no «tiene su trabajo en None».
    with pytest.raises(ErrorRuesLinker, match="_trabajo") as exc:
        escribir_resultado(
            _con(res, dir_trabajo=None), tmp_path, "prueba", marca_tiempo=MARCA, excel=False
        )
    assert "None" not in str(exc.value), str(exc.value)
    assert "no tiene dir_trabajo" in str(exc.value)
    assert (ajeno / "ajeno.txt").is_file()
    assert not (tmp_path / CARPETA_ESPERADA).exists()


@pytest.mark.parametrize("contenido", ["{", "[]"])
def test_trabajo_manifest_ilegible_se_declara_en_omitidos(
    res: ResultadoLinkage, tmp_path: Path, contenido: str
) -> None:
    trabajo = tmp_path / "trabajo_roto"
    trabajo.mkdir()
    (trabajo / "manifest.json").write_text(contenido, encoding="utf-8")
    man = escribir_resultado(
        _con(res, dir_trabajo=trabajo),
        tmp_path / "salida",
        "prueba",
        marca_tiempo=MARCA,
        excel=False,
    )
    assert man.tiempos_por_fase == {} and man.rss_por_fase == {}
    omitidos = {o["artefacto"]: o["motivo"] for o in man.omitidos}
    assert "tiempos_por_fase" in omitidos and "rss_por_fase" in omitidos
    assert "_trabajo/manifest.json" in omitidos["tiempos_por_fase"]
    assert (
        "ilegible" in omitidos["tiempos_por_fase"]
        or "no es un objeto" in omitidos["tiempos_por_fase"]
    )
    texto = json.loads((man.carpeta / "manifest.json").read_text(encoding="utf-8"))
    assert {o["artefacto"] for o in texto["omitidos"]} >= {"tiempos_por_fase", "rss_por_fase"}


def test_figuras_homonimas_fallan_antes_de_escribir(res: ResultadoLinkage, tmp_path: Path) -> None:
    """``figuras/`` es plano: dos PNG con el mismo nombre se pisarían y el
    manifiesto listaría dos entradas para un archivo. Se falla con mensaje
    accionable y no queda ni carpeta definitiva ni pendiente. La misma ruta
    pasada dos veces NO es un choque: se copia una vez."""
    a, b = tmp_path / "a", tmp_path / "b"
    a.mkdir(), b.mkdir()
    (a / "fig.png").write_bytes(b"\x89PNG" + b"\x00" * 10)
    (b / "fig.png").write_bytes(b"\x89PNG" + b"\x00" * 20)
    salida = tmp_path / "salida"
    with pytest.raises(ErrorRuesLinker, match=r"fig\.png") as exc:
        escribir_resultado(
            res,
            salida,
            "prueba",
            marca_tiempo=MARCA,
            excel=False,
            figuras=[a / "fig.png", b / "fig.png"],
        )
    assert "Qué hacer" in str(exc.value) and "figuras=" in str(exc.value)
    assert not (salida / CARPETA_ESPERADA).exists()
    assert not (salida / ".prueba.pendiente").exists()
    # La misma figura dos veces se copia una vez y el manifiesto la lista una vez.
    man = escribir_resultado(
        res,
        salida,
        "prueba",
        marca_tiempo=MARCA,
        excel=False,
        figuras=[a / "fig.png", a / "fig.png"],
    )
    assert [x["ruta"] for x in man.artefactos if x["ruta"].startswith("figuras/")] == [
        "figuras/fig.png"
    ]
    assert leer_resultado(man.carpeta).validar().ok


def test_carpeta_definitiva_existente_falla_rapido(res: ResultadoLinkage, tmp_path: Path) -> None:
    escribir_resultado(res, tmp_path, "prueba", marca_tiempo=MARCA, excel=False)
    with pytest.raises(ErrorRuesLinker, match="ya existe"):
        escribir_resultado(res, tmp_path, "prueba", marca_tiempo=MARCA, excel=False)
    assert not (tmp_path / ".prueba.pendiente").exists()


def test_nombre_con_separadores_falla_rapido(res: ResultadoLinkage, tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="nombre"):
        escribir_resultado(res, tmp_path, "../fuga", marca_tiempo=MARCA)


def test_resultado_que_incumple_el_contrato_no_se_escribe(
    res: ResultadoLinkage, tmp_path: Path
) -> None:
    roto = ResultadoLinkage(
        correlativa=res.correlativa.drop(columns=["ID_ENTIDAD"]),
        golden=res.golden,
        manifiesto=dict(res.manifiesto),
        metricas=dict(res.metricas),
        diccionario=res.diccionario,
        dir_trabajo=res.dir_trabajo,
    )
    with pytest.raises(ContratoSalidaError):
        escribir_resultado(roto, tmp_path, "prueba", marca_tiempo=MARCA)
    assert list(tmp_path.iterdir()) == []


def test_golden_none_se_omite_y_se_declara(res: ResultadoLinkage, tmp_path: Path) -> None:
    sin_golden = ResultadoLinkage(
        correlativa=res.correlativa,
        golden=None,
        manifiesto=dict(res.manifiesto),
        metricas=dict(res.metricas),
        diccionario=res.diccionario,
        dir_trabajo=res.dir_trabajo,
    )
    man = escribir_resultado(sin_golden, tmp_path, "prueba", marca_tiempo=MARCA)
    assert not (man.carpeta / "golden.parquet").exists()
    assert not (man.carpeta / "excel" / "golden.xlsx").exists()
    assert (man.carpeta / "excel" / "correlativa.xlsx").is_file()
    assert {o["artefacto"] for o in man.omitidos} == {"golden.parquet", "excel/golden.xlsx"}
    leido = leer_resultado(man.carpeta)
    assert leido.golden is None and leido.validar().ok


# ─────────────────────────────────────────────────────────────────────────────
# leer_resultado
# ─────────────────────────────────────────────────────────────────────────────


def test_leer_resultado_devuelve_un_resultado_valido(res: ResultadoLinkage, tmp_path: Path) -> None:
    man = escribir_resultado(res, tmp_path, "prueba", marca_tiempo=MARCA)
    leido = leer_resultado(man.carpeta)
    assert isinstance(leido, ResultadoLinkage)
    reporte = leido.validar()
    assert reporte.ok, reporte.resumen()
    assert list(leido.correlativa.columns) == list(res.correlativa.columns)
    assert list(leido.golden.columns) == list(res.golden.columns)
    pd.testing.assert_frame_equal(
        leido.correlativa[["ID_REGISTRO", "ID_GRUPO", "ID_ENTIDAD", "SCORE_PAR"]],
        res.correlativa[["ID_REGISTRO", "ID_GRUPO", "ID_ENTIDAD", "SCORE_PAR"]],
        check_dtype=False,
    )
    assert leido.enlaces is None
    assert list(leido.revision.columns) == list(contrato.COLUMNAS_REVISION) and leido.revision.empty
    assert list(leido.diccionario.columns) == list(contrato.COLUMNAS_DICCIONARIO)
    assert leido.manifiesto["entradas"] == res.manifiesto["entradas"]
    assert leido.metricas["n_registros"] == len(res.correlativa)
    assert leido.manifiesto["manifest"]["nombre"] == "prueba"
    assert leido.dir_trabajo is None  # _trabajo/ no está dentro de esta carpeta


def test_revision_csv_es_fiel_a_los_datos(res: ResultadoLinkage, tmp_path: Path) -> None:
    """``revision.csv`` es la forma del archivo de decisiones y ``leer_resultado``
    lo lee de vuelta: no se neutraliza (eso es para los .xlsx y los alias
    .csv.gz, que se abren en hoja de cálculo). Nada se repara en silencio."""
    rev = REVISION_CON_PREFIJOS.copy()
    assert list(rev.columns) == list(contrato.COLUMNAS_REVISION)
    man = escribir_resultado(_con(res, revision=rev), tmp_path, "prueba", marca_tiempo=MARCA)
    texto = (man.carpeta / "revision.csv").read_text(encoding="utf-8")
    assert "'" not in texto, texto
    assert "=EMPRESA A SAS" in texto and "+no es la misma" in texto
    leido = leer_resultado(man.carpeta)
    pd.testing.assert_frame_equal(leido.revision, rev)
    assert leido.revision["RAZON"].isna().tolist() == [False, False, False, False, True]
    assert leido.validar().ok
    # El diccionario también es fiel: ningún apóstrofo añadido.
    dic = (man.carpeta / "diccionario.csv").read_text(encoding="utf-8")
    assert "'=" not in dic and "'-" not in dic and "'+" not in dic and "'@" not in dic


def test_revision_csv_celda_vacia_es_ausente(res: ResultadoLinkage, tmp_path: Path) -> None:
    """Un CSV no distingue ``''`` de ausente: en el estándar la celda vacía
    ES el ausente, y ``leer_resultado`` devuelve ``pd.NA`` (lo que una persona
    deja en blanco en el archivo de decisiones no es una cadena vacía). Una
    ``''`` en memoria vuelve como NA: queda documentado aquí y en el escritor,
    no se repara en silencio. Los textos ``NA``, ``null`` o ``nan`` siguen
    siendo texto."""
    rev = REVISION_CON_PREFIJOS.copy()
    rev.loc[0, "RAZON"] = ""
    rev.loc[1, "RAZON"] = "NA"
    rev.loc[2, "RAZON"] = "null"
    man = escribir_resultado(_con(res, revision=rev), tmp_path, "prueba", marca_tiempo=MARCA)
    leido = leer_resultado(man.carpeta).revision
    assert leido["RAZON"].isna().tolist() == [True, False, False, False, True]
    assert leido["RAZON"].tolist()[1:4] == ["NA", "null", "@matriz y filial"]
    esperado = rev.copy()
    esperado.loc[0, "RAZON"] = pd.NA
    pd.testing.assert_frame_equal(leido, esperado)


def test_leer_resultado_con_alias_en_espanol(res: ResultadoLinkage, tmp_path: Path) -> None:
    man = escribir_resultado(res, tmp_path, "prueba", marca_tiempo=MARCA, excel=False)
    leido = leer_resultado(man.carpeta, alias="es")
    assert list(leido.correlativa.columns[:3]) == ["ID_REGISTRO", "FUENTE", "FILA_ORIGEN"]
    assert "SIMILITUD_NOMBRE" in leido.correlativa.columns
    assert "FUENTE_PRINCIPAL" in leido.golden.columns and "N_REGISTROS" in leido.golden.columns
    # Las columnas de la fuente conservan su nombre (alias_es = columna).
    assert "SECTOR" in leido.correlativa.columns
    with pytest.raises(ValueError, match="alias"):
        leer_resultado(man.carpeta, alias="en")


def test_leer_resultado_detecta_artefacto_alterado(res: ResultadoLinkage, tmp_path: Path) -> None:
    man = escribir_resultado(res, tmp_path, "prueba", marca_tiempo=MARCA, excel=False)
    ruta = man.carpeta / "revision.csv"
    with ruta.open("a", encoding="utf-8") as f:
        f.write("cruce,X,a,A,b,B,distinta,alguien,porque\n")
    with pytest.raises(ErrorRuesLinker, match=r"revision\.csv") as exc:
        leer_resultado(man.carpeta)
    assert "Qué hacer" in str(exc.value)
    ruta.unlink()
    with pytest.raises(ErrorRuesLinker, match=r"revision\.csv"):
        leer_resultado(man.carpeta)


def test_leer_resultado_sin_manifest_falla_con_mensaje(tmp_path: Path) -> None:
    with pytest.raises(ErrorRuesLinker, match=r"manifest\.json"):
        leer_resultado(tmp_path)


# ─────────────────────────────────────────────────────────────────────────────
# linkage(carpeta_salida=...) y alias v1 en L6
# ─────────────────────────────────────────────────────────────────────────────


def test_linkage_con_carpeta_salida_escribe_el_estandar(tmp_path: Path) -> None:
    fuentes = conjunto_sintetico()
    with _silencio():
        res = rl.linkage(
            fuentes,
            carpeta_salida=tmp_path / "salidas",
            nombre="cruce",
            skip_reporting=True,
            col_ciudad="CIUDAD",
            col_id="CODIGO",
        )
    carpetas = [p for p in (tmp_path / "salidas").iterdir()]
    assert len(carpetas) == 1 and carpetas[0].name.endswith("_cruce")
    carpeta = carpetas[0]
    assert not (tmp_path / "salidas" / ".cruce.pendiente").exists()
    assert (carpeta / "_trabajo" / "manifest.json").is_file()
    assert (carpeta / "_trabajo" / "L3_scoring" / "scored.db").is_file()
    assert res.dir_trabajo == carpeta / "_trabajo"
    assert res.manifiesto["dir_trabajo"] == str(carpeta / "_trabajo")
    assert res.manifiesto["carpeta_salida"] == str(carpeta)
    assert set(_rutas_relativas(carpeta)) == set(RUTAS_ESPERADAS) - {"figuras/dashboard.png"}
    # En memoria, ninguna ruta sigue apuntando a la pendiente que ya no existe.
    assert res.manifiesto["columnas_tecnicas"]["quedan_en"] == str(carpeta / "_trabajo")
    origen = res.manifiesto["completar"]["score_par"]["origen"]
    assert origen == str(carpeta / "_trabajo" / "L3_scoring" / "scored.db")
    assert Path(origen).is_file()
    texto = (carpeta / "manifest.json").read_text(encoding="utf-8")
    assert ".pendiente" not in texto
    man = json.loads(texto)
    # En el JSON publicado las rutas bajo la pendiente quedan relativas a la
    # carpeta (sobreviven a mover la carpeta o cambiar de máquina).
    assert man["corrida"]["dir_trabajo"] == "_trabajo"
    assert man["corrida"]["columnas_tecnicas"]["quedan_en"] == "_trabajo"
    assert man["corrida"]["completar"]["score_par"]["origen"] == "_trabajo/L3_scoring/scored.db"
    assert set(man["tiempos_por_fase"]) == {
        "L1_prep",
        "L2_lsh_candidates",
        "L3_scoring",
        "L4_clustering",
        "L5_golden",
    }
    leido = leer_resultado(carpeta)
    assert leido.validar().ok
    assert leido.dir_trabajo == carpeta / "_trabajo"
    # leer_resultado las resuelve contra la carpeta leída.
    assert leido.manifiesto["columnas_tecnicas"]["quedan_en"] == str(carpeta / "_trabajo")
    assert leido.manifiesto["completar"]["score_par"]["origen"] == origen


def _exigir_l6() -> None:
    import importlib.util

    faltan = [m for m in ("matplotlib", "seaborn") if importlib.util.find_spec(m) is None]
    if faltan:
        pytest.fail(f"Faltan {faltan}: L6 omitiría sus PNG en silencio.", pytrace=False)


@pytest.fixture(scope="module")
def corrida_con_l6(tmp_path_factory: pytest.TempPathFactory) -> dict[str, Any]:
    _exigir_l6()
    raiz = tmp_path_factory.mktemp("l6")
    fuentes = conjunto_sintetico()
    from record_linkage.reporting import strategies

    strategies._ALIAS_V1_AVISADO = False
    with _silencio(), warnings.catch_warnings(record=True) as avisos:
        warnings.simplefilter("always")
        res = rl.linkage(
            fuentes,
            carpeta_salida=raiz,
            nombre="conl6",
            skip_reporting=False,
            col_ciudad="CIUDAD",
            col_id="CODIGO",
        )
    carpeta = next(p for p in raiz.iterdir() if p.name.endswith("_conl6"))
    return {"res": res, "carpeta": carpeta, "avisos": list(avisos)}


def test_alias_v1_se_escriben_con_deprecation_una_vez(corrida_con_l6: dict[str, Any]) -> None:
    carpeta: Path = corrida_con_l6["carpeta"]
    l6 = carpeta / "_trabajo" / "L6_reporting"
    for nombre in ("tabla_correlativa", "golden_records"):
        for ext in (".parquet", ".csv.gz", ".xlsx"):
            assert (l6 / f"{nombre}{ext}").is_file(), f"{nombre}{ext}"
    de_alias = [
        a
        for a in corrida_con_l6["avisos"]
        if issubclass(a.category, DeprecationWarning) and "excel/correlativa.xlsx" in str(a.message)
    ]
    assert len(de_alias) == 1, [str(a.message) for a in de_alias]
    assert VERSION_RETIRO_ALIAS_V1 in str(de_alias[0].message)


def test_alias_v1_avisan_tambien_por_logging(caplog: pytest.LogCaptureFixture) -> None:
    """Python, IPython y Colab silencian los ``DeprecationWarning`` que no
    nacen en ``__main__``: el aviso sale además por ``logging`` (una vez)."""
    from record_linkage.reporting import strategies

    strategies._ALIAS_V1_AVISADO = False
    with caplog.at_level(logging.WARNING), warnings.catch_warnings():
        warnings.simplefilter("ignore", DeprecationWarning)
        strategies._avisar_alias_v1()
        strategies._avisar_alias_v1()
    registros = [r for r in caplog.records if "ALIAS de v1" in r.getMessage()]
    assert len(registros) == 1
    assert registros[0].levelno == logging.WARNING
    assert VERSION_RETIRO_ALIAS_V1 in registros[0].getMessage()
    assert "excel/correlativa.xlsx" in registros[0].getMessage()


def test_alias_parquet_v1_no_lleva_el_metadato_del_contrato(
    corrida_con_l6: dict[str, Any],
) -> None:
    import pyarrow.parquet as pq

    carpeta: Path = corrida_con_l6["carpeta"]
    l6 = carpeta / "_trabajo" / "L6_reporting"
    for nombre in ("tabla_correlativa", "golden_records"):
        metadatos = pq.read_schema(l6 / f"{nombre}.parquet").metadata or {}
        assert b"contrato" not in metadatos, nombre
    assert pq.read_schema(carpeta / "correlativa.parquet").metadata[b"contrato"] == b"1.0"


def test_reportes_l6_conservan_la_hoja_sheet1(corrida_con_l6: dict[str, Any]) -> None:
    """``reporte_*.xlsx`` (L6 de v1) no es un alias: conserva la forma de v1
    (``to_excel`` sin ``sheet_name`` → ``Sheet1``) hasta que F1.11 lo unifique."""
    carpeta: Path = corrida_con_l6["carpeta"]
    reportes = sorted((carpeta / "_trabajo" / "L6_reporting").glob("reporte_*.xlsx"))
    assert reportes, "L6 no dejó ningún reporte_*.xlsx"
    for ruta in reportes:
        libro = load_workbook(ruta, read_only=True)
        try:
            assert libro.sheetnames == ["Sheet1"], ruta.name
        finally:
            libro.close()


def test_alias_xlsx_lleva_hoja_leeme_primero(corrida_con_l6: dict[str, Any]) -> None:
    carpeta: Path = corrida_con_l6["carpeta"]
    ruta = carpeta / "_trabajo" / "L6_reporting" / "tabla_correlativa.xlsx"
    libro = load_workbook(ruta, read_only=True)
    try:
        assert libro.sheetnames[0] == "LEEME"
        assert len(libro.sheetnames) == 2
        texto = "\n".join(
            str(c.value) for fila in libro["LEEME"].iter_rows() for c in fila if c.value
        )
        datos = libro[libro.sheetnames[1]]
        encabezado = [c.value for c in next(datos.iter_rows(max_row=1))]
    finally:
        libro.close()
    assert "excel/correlativa.xlsx" in texto
    assert VERSION_RETIRO_ALIAS_V1 in texto
    assert "ID_GRUPO" in encabezado


def test_con_l6_las_figuras_van_a_figuras(corrida_con_l6: dict[str, Any]) -> None:
    carpeta: Path = corrida_con_l6["carpeta"]
    figuras = sorted(p.name for p in (carpeta / "figuras").iterdir())
    assert figuras and all(f.endswith(".png") for f in figuras)
    assert "dashboard_ejecutivo.png" in figuras
    man = json.loads((carpeta / "manifest.json").read_text(encoding="utf-8"))
    rutas = {a["ruta"] for a in man["artefactos"]}
    assert {f"figuras/{f}" for f in figuras} <= rutas
    res: ResultadoLinkage = corrida_con_l6["res"]
    assert res.metricas["report_files"]
    assert leer_resultado(carpeta).validar().ok


def test_con_l6_las_rutas_del_resultado_sobreviven_al_rename(
    corrida_con_l6: dict[str, Any],
) -> None:
    """El docstring de ``linkage()`` promete ``report_files`` con L6 activo:
    tras el ``rename`` de la pendiente esas rutas (y las demás del resultado)
    deben existir, no apuntar a ``.conl6.pendiente/``."""
    carpeta: Path = corrida_con_l6["carpeta"]
    res: ResultadoLinkage = corrida_con_l6["res"]
    report_files = [Path(f) for f in res.metricas["report_files"]]
    assert len(report_files) >= 20
    assert all(f.is_file() for f in report_files), [str(f) for f in report_files if not f.is_file()]
    assert all(carpeta / "_trabajo" in f.parents for f in report_files)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", DeprecationWarning)
        assert res.get("report_files") == res.metricas["report_files"]
    assert res.manifiesto["columnas_tecnicas"]["quedan_en"] == str(carpeta / "_trabajo")
    assert Path(res.manifiesto["completar"]["score_par"]["origen"]).is_file()
    assert ".pendiente" not in json.dumps(res.manifiesto, default=str)
    assert ".pendiente" not in json.dumps(res.metricas, default=str)
    assert ".pendiente" not in (carpeta / "manifest.json").read_text(encoding="utf-8")


# ─────────────────────────────────────────────────────────────────────────────
# Reubicar y relativizar rutas anidadas (lo que usa el rename)
# ─────────────────────────────────────────────────────────────────────────────


def test_reubicar_rutas_recorre_la_estructura_y_respeta_el_limite(tmp_path: Path) -> None:
    de = tmp_path / ".x.pendiente"
    a = tmp_path / "2026-10-06_1430_x"
    valor = {
        "lista": [str(de / "_trabajo" / "a.db"), de / "_trabajo", str(de)],
        "tupla": (str(de / "b"), 3, None),
        "otro": str(tmp_path / ".x.pendiente2" / "c"),  # NO está bajo de: límite en el separador
        "texto": "sin ruta",
        "anidado": {"quedan_en": str(de / "_trabajo")},
    }
    nuevo = escritor._reubicar_rutas(valor, de, a)
    assert nuevo == {
        "lista": [str(a / "_trabajo" / "a.db"), a / "_trabajo", str(a)],
        "tupla": (str(a / "b"), 3, None),
        "otro": str(tmp_path / ".x.pendiente2" / "c"),
        "texto": "sin ruta",
        "anidado": {"quedan_en": str(a / "_trabajo")},
    }
    assert isinstance(nuevo["lista"][1], Path)
    assert valor["anidado"]["quedan_en"] == str(de / "_trabajo")  # el original no se toca
    relativo = escritor._relativizar_rutas(valor, de)
    assert relativo["lista"] == ["_trabajo/a.db", Path("_trabajo"), "."]
    assert relativo["otro"] == valor["otro"]
    # El camino inverso (leer_resultado): de relativa a la carpeta leída.
    vuelta = escritor._reubicar_rutas(
        relativo, Path("_trabajo"), a / "_trabajo", tambien_resuelta=False
    )
    assert vuelta["lista"][:2] == [str(a / "_trabajo" / "a.db"), a / "_trabajo"]
    assert vuelta["anidado"] == {"quedan_en": str(a / "_trabajo")}


# ─────────────────────────────────────────────────────────────────────────────
# Un solo punto de escritura fuera de _trabajo/
# ─────────────────────────────────────────────────────────────────────────────

#: Llamadas que escriben parquet/excel/csv. ``to_sql`` queda fuera a propósito:
#: SQLite es ``_trabajo/`` (L3_scoring/scored.db), nunca el entregable.
_ESCRITORES = {
    "to_parquet",
    "to_excel",
    "to_csv",
    "write_table",
    "write_to_dataset",
    "ParquetWriter",
    "ExcelWriter",
    "Workbook",
}

#: Módulos de ``src/record_linkage`` que pueden llamar a un escritor, y por
#: qué. El objetivo (F1.10) es que la carpeta del estándar la escriba SOLO
#: ``exporters/escritor.py``; lo demás es ``_trabajo/`` (checkpoints), caminos
#: que F2 lleva al estándar, o herramientas que no producen el entregable.
#: Lo que se encontró de más respecto a la lista de la especificación
#: (``golden/generator.py`` y ``engine/`` NO escriben; ``reporting/suite.py``,
#: ``flujo/``, ``deduplication/colab.py``, ``exporters/smart.py``,
#: ``evaluation/``, ``optimization/`` y ``classifier/`` sí) se lista con su
#: motivo. Quitar una entrada exige quitar primero la llamada.
ESCRITORES_PERMITIDOS: dict[str, str] = {
    "exporters/escritor.py": "el escritor único del estándar (F1.10).",
    "pipeline/orchestrator.py": "checkpoints L1…L5 en _trabajo/.",
    "pipeline/linkage_pipeline.py": "checkpoints del pipeline heredado en _trabajo/.",
    "pipeline/result.py": "PipelineResult.to_excel/to_csv heredados (lo pide el usuario).",
    "reporting/suite.py": "reportes L6 de v1 (reporte_*.xlsx); F1.11 los unifica en "
    "informe_cruce.xlsx a través del escritor.",
    "exporters/smart.py": "SmartExporter: utilidad genérica de exportación, no el estándar.",
    "flujo/cruce.py": "camino cruce (ejecutar_cruce): F2 lo lleva al estándar.",
    "flujo/insumos.py": "camino cruce: caché de insumos en parquet.",
    "flujo/resultados_disco.py": "camino cruce: resultados en disco; F2 lo lleva al estándar.",
    "deduplication/colab.py": "camino dedupe por lotes en Colab; F2 lo lleva al estándar.",
    "evaluation/banco.py": "el banco (medición), no el entregable.",
    "evaluation/ground_truth.py": "ground truth (medición), no el entregable.",
    "optimization/engine.py": "reportes de Optuna, no el entregable.",
    "classifier/data_loader.py": "datos de entrenamiento del clasificador.",
}


def _llamadas_a_escritores(ruta: Path) -> list[str]:
    """Nombres de escritor llamados en el módulo (AST: ni docstrings ni comentarios)."""
    arbol = ast.parse(ruta.read_text(encoding="utf-8"))
    encontrados: list[str] = []
    for nodo in ast.walk(arbol):
        if not isinstance(nodo, ast.Call):
            continue
        f = nodo.func
        nombre = f.attr if isinstance(f, ast.Attribute) else f.id if isinstance(f, ast.Name) else ""
        if nombre in _ESCRITORES:
            encontrados.append(nombre)
    return encontrados


def test_solo_el_escritor_escribe_fuera_de_trabajo() -> None:
    con_llamadas = {
        p.relative_to(SRC).as_posix(): _llamadas_a_escritores(p) for p in sorted(SRC.rglob("*.py"))
    }
    con_llamadas = {k: v for k, v in con_llamadas.items() if v}
    de_mas = sorted(set(con_llamadas) - set(ESCRITORES_PERMITIDOS))
    assert not de_mas, (
        f"Módulos que escriben parquet/excel/csv sin estar en la lista: "
        f"{ {k: con_llamadas[k] for k in de_mas} }. La carpeta del estándar la escribe "
        f"SOLO exporters/escritor.py; si es _trabajo/ o una herramienta, añádalo a "
        f"ESCRITORES_PERMITIDOS con su motivo."
    )
    sin_llamadas = sorted(k for k in ESCRITORES_PERMITIDOS if k not in con_llamadas)
    assert not sin_llamadas, f"Ya no escriben; quítelos de la lista: {sin_llamadas}"
    # L6 escribe a través del escritor: strategies.py no llama a ningún escritor.
    assert "reporting/strategies.py" not in con_llamadas


def test_la_lista_de_permitidos_nombra_archivos_reales() -> None:
    """Una clave mal escrita no debe pasar como «ya no escribe»."""
    inexistentes = sorted(k for k in ESCRITORES_PERMITIDOS if not (SRC / k).is_file())
    assert not inexistentes, f"No existen bajo src/record_linkage: {inexistentes}"


def test_la_compuerta_ve_parquetwriter_y_workbook() -> None:
    """``exporters/smart.py`` escribe con ``pq.ParquetWriter`` y ``openpyxl.Workbook``
    (no con ``to_parquet``/``to_excel``); la compuerta debe verlo."""
    llamadas = set(_llamadas_a_escritores(SRC / "exporters" / "smart.py"))
    assert {"ParquetWriter", "Workbook"} <= llamadas, llamadas


def test_escribir_xlsx_neutraliza_formulas(tmp_path: Path) -> None:
    """La ruta de L6 (alias) pasa por aquí: conserva la neutralización de 0.22.4."""
    df = pd.DataFrame({"=HEADER": ["=2+2", "safe"], "number": [-7, 1]})
    ruta = tmp_path / "x.xlsx"
    escritor.escribir_xlsx(df, ruta, leeme=escritor.leeme_alias_v1("excel/correlativa.xlsx"))
    libro = load_workbook(ruta, read_only=True, data_only=False)
    try:
        assert libro.sheetnames == ["LEEME", "datos"]
        hoja = libro["datos"]
        assert hoja["A1"].value == "'=HEADER"
        assert hoja["A2"].value == "'=2+2"
        assert hoja["B2"].value == -7
    finally:
        libro.close()
    pd.testing.assert_frame_equal(
        df, pd.DataFrame({"=HEADER": ["=2+2", "safe"], "number": [-7, 1]})
    )


def test_escribir_csv_gz_por_lotes_neutraliza(tmp_path: Path) -> None:
    import gzip

    import pyarrow.parquet as pq

    origen = tmp_path / "o.parquet"
    pd.DataFrame({"=H": ["=2+2", "b"], "n": [1, 2]}).to_parquet(origen, index=False)
    destino = tmp_path / "o.csv.gz"
    filas = escritor.escribir_csv_gz_por_lotes(pq.ParquetFile(origen), destino, filas_por_lote=1)
    assert filas == 2
    with gzip.open(destino, "rt", encoding="utf-8") as f:
        texto = f.read()
    assert texto.count("'=H") == 1 and "'=2+2" in texto

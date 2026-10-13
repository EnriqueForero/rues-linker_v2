"""Arreglos declarados de F3.1b en el Excel de entrega (``vinculacion_excel``) y en la validación de sus entradas."""

from __future__ import annotations

import warnings

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook

from record_linkage.flujo.vinculacion.config import Fuente
from record_linkage.flujo.vinculacion.controles import comparar_fuentes_decision
from record_linkage.flujo.vinculacion.lectura import leer_fuente
from record_linkage.flujo.vinculacion_excel import estilos, hojas, tablas


def _celdas(hoja) -> dict[str, tuple[object, str]]:
    return {
        c.coordinate: (c.value, c.data_type)
        for f in hoja.iter_rows()
        for c in f
        if c.value is not None
    }


# ── E5 · los encabezados pasan por el mismo escritor seguro que los datos ───────────────────────────
def test_e5_un_encabezado_que_empieza_con_igual_es_texto_no_formula():
    # Antes: la columna «=MOTIVO» quedaba como fórmula.
    hoja = Workbook().active
    tablas.escribir_tabla(hoja, pd.DataFrame({"=MOTIVO": ["x"], "OK": ["y"]}))
    assert _celdas(hoja)["A1"] == ("=MOTIVO", "s")


def test_e5_un_encabezado_con_caracter_de_control_no_hace_caer_la_escritura():
    # Antes: IllegalCharacterError de openpyxl.
    hoja = Workbook().active
    tablas.escribir_tabla(hoja, pd.DataFrame({"MO\x00TI\x1fVO": ["x"]}))
    assert hoja["A1"].value == "MO TI VO"


# ── E4 · las advertencias del RESUMEN también pasan por el escritor seguro ──────────────────────────
def test_e4_las_advertencias_del_resumen_no_son_formulas_ni_hacen_caer_la_escritura():
    # Antes: FUENTE «=1+1» quedaba como fórmula y un control en HALLAZGO lanzaba IllegalCharacterError.
    calidad = pd.DataFrame(
        {
            "FUENTE": ["=1+1"],
            "TEMA": ['=HYPERLINK("x")'],
            "HALLAZGO": ["se vio \x01 algo."],
            "IMPACTO": ["Impacto."],
        }
    )
    hoja = Workbook().active
    hojas._advertencias_de_calidad(hoja, estilos.Estilos(), calidad, 1)
    escrito = _celdas(hoja)
    assert escrito["A2"] == ("=1+1", "s")
    assert escrito["B2"][1] == "s"
    assert escrito["C2"] == (
        "se vio   algo. Impacto.",
        "s",
    )  # el control \x01 pasó a ser un espacio


# ── E1 · una columna sin valores no tumba el ancho de columnas (pandas 3) ───────────────────────────
@pytest.mark.parametrize(
    "columna",
    [
        pytest.param([None, None], id="object_todo_None"),
        pytest.param([np.nan, np.nan], id="float_todo_NaN"),
        pytest.param(pd.array([pd.NA, pd.NA], dtype="string"), id="string_todo_NA"),
        pytest.param(pd.array([None, None], dtype="Int64"), id="Int64_todo_NA"),
    ],
)
def test_e1_el_ancho_de_una_columna_vacia_es_el_de_su_encabezado(columna):
    # Antes (pandas 3): ValueError «cannot convert float NaN to integer».
    tabla = pd.DataFrame({"COLUMNA_VACIA": columna, "LLENA": ["abcdefgh", "abcdefgh"]})
    assert tablas._largos_de_columnas(tabla) == [len("COLUMNA_VACIA"), 8]


def test_e1_los_ausentes_no_cuentan_en_el_ancho_ni_dependen_de_la_version_de_pandas():
    # Antes, con pandas 2 el ausente contaba como el texto 'None'/'nan' (4 o 3 letras).
    tabla = pd.DataFrame({"X": ["ab", None, None, None, None, None, None, None, None, None]})
    assert tablas._largos_de_columnas(tabla) == [2]


def test_e1_exportar_una_tabla_de_una_fila_con_columnas_sin_valor_no_cae():
    hoja = Workbook().active
    tabla = pd.DataFrame({"A": [np.nan], "B": [None], "C": ["texto"]})
    tablas.escribir_tabla(hoja, tabla)
    tablas.formatear_tabla(hoja, tabla, estilos.Estilos())
    assert hoja.column_dimensions["C"].width == 8


# ── E8 · dos hojas que coinciden en sus 31 primeros caracteres no se pisan ──────────────────────────
def test_e8_dos_nombres_iguales_en_sus_31_caracteres_reciben_nombres_distintos_y_sin_aviso():
    libro = Workbook()
    tabla = pd.DataFrame({"A": [1]})
    largo = "DETALLE_BASE_CON_NOMBRE_MUY_LARGO_"
    with warnings.catch_warnings():
        warnings.simplefilter("error")  # openpyxl avisaba al renombrar a 32 caracteres
        a = tablas.hoja_desde_tabla(libro, estilos.Estilos(), largo + "UNO", tabla, "negocio")
        b = tablas.hoja_desde_tabla(libro, estilos.Estilos(), largo + "DOS", tabla, "negocio")
        c = tablas.hoja_desde_tabla(libro, estilos.Estilos(), largo + "TRES", tabla, "negocio")
    nombres = [a.title, b.title, c.title]
    assert nombres[0] == (largo + "UNO")[:31]
    assert len({n.lower() for n in nombres}) == 3
    assert all(len(n) <= 31 for n in nombres)
    assert nombres[1].endswith("_2") and nombres[2].endswith("_3")


def test_e8_la_comparacion_de_nombres_ignora_mayusculas_como_excel():
    libro = Workbook()
    tabla = pd.DataFrame({"A": [1]})
    a = tablas.hoja_desde_tabla(libro, estilos.Estilos(), "Detalle", tabla, "negocio")
    b = tablas.hoja_desde_tabla(libro, estilos.Estilos(), "DETALLE", tabla, "negocio")
    assert a.title == "Detalle" and b.title.lower() != a.title.lower()


# ── D · claves_repeticion y metricas_repetidas deben ser columnas de la base ────────────────────────
def _csv(tmp_path, contenido):
    ruta = tmp_path / "ancla.csv"
    ruta.write_text(contenido, encoding="utf-8")
    return ruta


def test_d_una_clave_de_repeticion_inexistente_falla_al_leer_la_fuente(tmp_path):
    # Antes: la errata ('HS' por 'HS6') se ignoraba sin aviso y la cifra «sin repetidos» deduplicaba solo
    # por (empresa, valor): subestimaba.
    ruta = _csv(tmp_path, "Empresa;Pais;Kilos;HS6\nKronos;ESPAÑA;10;1905\n")
    f = Fuente(
        clave="EXPORTADORES",
        archivo=str(ruta),
        col_nombre="Empresa",
        col_pais="Pais",
        cols_metricas=("Kilos",),
        metricas_repetidas=("Kilos",),
        claves_repeticion=("HS",),
    )
    with pytest.raises(KeyError, match=r"(?s)Qué pasó.*claves_repeticion.*'HS'.*Qué hacer"):
        leer_fuente(f)


def test_d_con_la_clave_de_repeticion_presente_se_lee_igual_que_antes(tmp_path):
    ruta = _csv(tmp_path, "Empresa;Pais;Kilos;HS6\nKronos;ESPAÑA;10;1905\n")
    f = Fuente(
        clave="EXPORTADORES",
        archivo=str(ruta),
        col_nombre="Empresa",
        col_pais="Pais",
        cols_metricas=("Kilos",),
        metricas_repetidas=("Kilos",),
        claves_repeticion=("HS6",),
    )
    assert leer_fuente(f)["HS6"].tolist() == ["1905"]


def test_d_una_metrica_repetida_que_no_es_metrica_se_rechaza_al_declarar_la_fuente():
    with pytest.raises(ValueError, match=r"(?s)Qué pasó.*metricas_repetidas.*'Fletes'.*Qué hacer"):
        Fuente(
            clave="EXPORTADORES",
            archivo="a.xlsx",
            col_nombre="Empresa",
            cols_metricas=("Kilos",),
            metricas_repetidas=("Fletes",),
        )


# ── M-b · COMPARACION_REVISIONES tiene un solo esquema ──────────────────────────────────────────────
def _decisiones(archivos_y_decisiones):
    return pd.DataFrame(
        {
            "ARCHIVO": [a for a, _ in archivos_y_decisiones],
            "TIPO": "VINCULO",
            "FUENTE": "ORBIS",
            "CLAVE_A": [f"A{k}" for k, _ in enumerate(archivos_y_decisiones)],
            "CLAVE_B": "B1",
            "NOMBRE_B": "NOMBRE",
            "DECISION": [d for _, d in archivos_y_decisiones],
            "RAZON": "",
        }
    )


def test_m_b_la_comparacion_tiene_las_mismas_columnas_vacia_o_llena():
    sin_decisiones = comparar_fuentes_decision(_decisiones([]).iloc[0:0])
    un_archivo = comparar_fuentes_decision(_decisiones([("a.csv", "EMPRESA")]))
    dos_archivos_sin_pares_repetidos = comparar_fuentes_decision(
        _decisiones(
            [("a.csv", "EMPRESA"), ("b.csv", "EMPRESA")]
        )  # CLAVE_A distinta: ningún par se repite
    )
    llena = comparar_fuentes_decision(
        _decisiones([("a.csv", "EMPRESA"), ("b.csv", "GRUPO")]).assign(CLAVE_A="A0")
    )
    assert len(llena) == 1
    for vacia in (sin_decisiones, un_archivo, dos_archivos_sin_pares_repetidos):
        assert len(vacia) == 0
        assert list(vacia.columns) == list(llena.columns)
    assert list(llena.columns)[-1] == "RAZON_ULTIMA"

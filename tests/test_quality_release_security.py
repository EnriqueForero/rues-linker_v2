"""Focused regression gates for production hardening and hosted notebooks."""

from __future__ import annotations

import gzip
import io
import logging
import os
import pickle
import sqlite3
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from record_linkage.classifier.hybrid import ClasificadorHibridoOptimizado
from record_linkage.classifier.runner import (
    _load_classifier_model,
    _save_classifier_model,
    ejecutar_proceso_clasificacion_directo,
)
from record_linkage.config.credentials import SnowflakeCredentials
from record_linkage.exporters.smart import SmartExporter
from record_linkage.optimization.engine import OptimizationEngine
from record_linkage.pipeline._internal import _class_exists
from record_linkage.reporting.strategies import DataExportStrategy
from record_linkage.utils import memory as memory_module
from record_linkage.utils.logger import CustomLogger


def test_process_rss_survives_psutil_no_such_process(monkeypatch: pytest.MonkeyPatch) -> None:
    def vanished_process(*_args, **_kwargs):
        raise memory_module.psutil.NoSuchProcess(pid=os.getpid())

    monkeypatch.setattr(memory_module.psutil, "Process", vanished_process)
    assert memory_module.get_process_rss_bytes() > 0
    assert memory_module.MemoryManager().get_memory_status()["process_mb"] > 0


def test_custom_logger_preserves_host_root_handlers_and_is_idempotent() -> None:
    root = logging.getLogger()
    original_level = root.level
    stream = io.StringIO()
    host_handler = logging.StreamHandler(stream)
    root.addHandler(host_handler)
    name = f"rues_linker.quality_release.{id(host_handler)}"
    named_logger = logging.getLogger(name)
    try:
        first = CustomLogger(name)
        handler_count = len(first.logger.handlers)
        second = CustomLogger(name)

        assert host_handler in root.handlers
        assert root.level == original_level
        assert host_handler.stream is stream
        assert second.logger is first.logger
        assert len(second.logger.handlers) == handler_count == 1
    finally:
        root.removeHandler(host_handler)
        host_handler.close()
        for handler in named_logger.handlers[:]:
            named_logger.removeHandler(handler)
            handler.close()


def test_custom_logger_is_safe_on_cp1252_stream(monkeypatch: pytest.MonkeyPatch) -> None:
    class NarrowStream(io.StringIO):
        encoding = "cp1252"

        def write(self, value: str) -> int:
            value.encode(self.encoding)
            return super().write(value)

    stream = NarrowStream()
    monkeypatch.setattr(sys, "stdout", stream)
    name = f"rues_linker.cp1252.{id(stream)}"
    named_logger = logging.getLogger(name)
    try:
        logger = CustomLogger(name)
        logger.info("Resultado ✅: niño")
        rendered = stream.getvalue()
        assert "Resultado \\u2705: niño" in rendered
    finally:
        for handler in named_logger.handlers[:]:
            named_logger.removeHandler(handler)
            handler.close()


@pytest.mark.parametrize("filename", ["../escape", "sub/escape", r"sub\escape", "/tmp/escape"])
def test_exporter_rejects_path_traversal(tmp_path: Path, filename: str) -> None:
    exporter = SmartExporter({"output_directory": str(tmp_path / "out")})
    with pytest.raises(ValueError, match=r"nombre simple|directorio"):
        exporter.export(pd.DataFrame({"x": [1]}), filename, format="csv")


def test_spreadsheet_formulas_are_escaped_without_mutating_input(tmp_path: Path) -> None:
    exporter = SmartExporter({"output_directory": str(tmp_path)})
    dangerous = ["=2+2", "+cmd", "-cmd", "@SUM(A1:A2)", "\tcmd", "safe", None]
    original = pd.DataFrame({"=dangerous_header": dangerous, "number": [-7, 1, 2, 3, 4, 5, 6]})

    prepared = exporter._prepare_spreadsheet_data(original)

    assert original.columns[0] == "=dangerous_header"
    assert original.iloc[0, 0] == "=2+2"
    assert prepared.columns[0] == "'=dangerous_header"
    assert prepared.iloc[:5, 0].str.startswith("'").all()
    assert prepared.iloc[5, 0] == "safe"
    # Los faltantes siguen siendo faltantes y NO se convierten en texto
    # escapado. Se comprueba la semántica (`isna`) y no la identidad con
    # `None`: pandas >= 3 unifica None/NaN al escribir un bloque object, de
    # modo que exigir `is None` haría fallar la prueba por una decisión de
    # pandas y no por un defecto de esta librería (verificado en 2.2.x y 3.0.x).
    assert pd.isna(prepared.iloc[6, 0])
    assert not isinstance(prepared.iloc[6, 0], str)
    assert prepared["number"].tolist() == [-7, 1, 2, 3, 4, 5, 6]

    csv_path = Path(exporter.export(original, "safe", format="csv"))
    assert "'=2+2" in csv_path.read_text(encoding="utf-8")

    xlsx_path = Path(exporter.export(original, "safe_excel", format="xlsx"))
    workbook = load_workbook(xlsx_path, read_only=True, data_only=False)
    try:
        worksheet = workbook.active
        assert worksheet["A1"].value == "'=dangerous_header"
        assert worksheet["A2"].value == "'=2+2"
        assert worksheet["A2"].data_type != "f"
        assert worksheet["B2"].value == -7
    finally:
        workbook.close()


def test_l6_streaming_exports_escape_spreadsheet_formulas(tmp_path: Path) -> None:
    source = tmp_path / "checkpoint.parquet"
    original = pd.DataFrame({"=HEADER": ["=2+2", "safe"], "number": [-7, 1]})
    original.to_parquet(source, index=False)

    strategy = DataExportStrategy()
    files = strategy._export_from_disk_streaming(
        source, "resultado", tmp_path, logger=logging.getLogger(__name__)
    )

    csv_path = next(path for path in files if path.name.endswith(".csv.gz"))
    with gzip.open(csv_path, "rt", encoding="utf-8") as stream:
        csv_text = stream.read()
    assert "'=HEADER" in csv_text
    assert "'=2+2" in csv_text

    xlsx_path = next(path for path in files if path.suffix == ".xlsx")
    workbook = load_workbook(xlsx_path, read_only=True, data_only=False)
    try:
        # F1.10: el alias de v1 lleva una primera hoja LEEME; los datos van después.
        assert workbook.sheetnames[0] == "LEEME"
        sheet = workbook[workbook.sheetnames[1]]
        assert sheet["A1"].value == "'=HEADER"
        assert sheet["A2"].value == "'=2+2"
        assert sheet["A2"].data_type != "f"
        assert sheet["B2"].value == -7
    finally:
        workbook.close()

    pd.testing.assert_frame_equal(original, pd.read_parquet(source))


def test_sqlite_export_quotes_catalog_table_and_streams_without_rowid(tmp_path: Path) -> None:
    db_path = tmp_path / "source.db"
    conn = sqlite3.connect(db_path)
    try:
        conn.execute('CREATE TABLE "odd"" table" ("=header" TEXT, value INTEGER)')
        conn.execute('INSERT INTO "odd"" table" VALUES (?, ?)', ("=2+2", -3))
        conn.execute('CREATE TABLE "without rowid" (id TEXT PRIMARY KEY, value TEXT) WITHOUT ROWID')
        conn.executemany(
            'INSERT INTO "without rowid" VALUES (?, ?)', [("b", "safe"), ("a", "@cmd")]
        )
        conn.commit()
    finally:
        conn.close()

    exporter = SmartExporter({"output_directory": str(tmp_path / "out")})
    odd_path = exporter.export_from_db(str(db_path), "odd", format="csv", table_name='odd" table')
    odd = pd.read_csv(odd_path, keep_default_na=False)
    assert odd.columns.tolist() == ["'=header", "value"]
    assert odd.iloc[0].tolist() == ["'=2+2", -3]

    without_rowid_path = exporter.export_from_db(
        str(db_path), "without_rowid", format="csv", table_name="without rowid"
    )
    without_rowid = pd.read_csv(without_rowid_path, keep_default_na=False)
    assert without_rowid["id"].tolist() == ["a", "b"]
    assert without_rowid["value"].tolist() == ["'@cmd", "safe"]

    with pytest.raises(ValueError, match="no encontrada"):
        exporter.export_from_db(
            str(db_path), "attack", format="csv", table_name='odd" table; DROP TABLE x;--'
        )


def test_class_lookup_never_executes_input(tmp_path: Path) -> None:
    marker = tmp_path / "executed"
    payload = f"__import__('pathlib').Path({str(marker)!r}).write_text('bad')"
    assert _class_exists(payload) is False
    assert not marker.exists()


def test_snowflake_password_is_redacted_from_repr() -> None:
    credentials = SnowflakeCredentials("account", "user", "TOP_SECRET", "wh", "db", "schema")
    assert "TOP_SECRET" not in repr(credentials)
    assert credentials.to_connector_kwargs()["password"] == "TOP_SECRET"


def test_classifier_json_roundtrip_and_legacy_pickle_is_rejected(tmp_path: Path) -> None:
    model = ClasificadorHibridoOptimizado({"umbral_empresa": 1, "umbral_persona": -1})
    model.word_weights = {"SOCIEDAD": 1.25}
    model.is_trained = True
    model_path = tmp_path / "model.pkl"  # Legacy suffix remains accepted for callers.
    _save_classifier_model(model, str(model_path))

    assert model_path.read_bytes().startswith(b"{")
    loaded = _load_classifier_model(str(model_path))
    assert loaded.word_weights == model.word_weights
    assert loaded.is_trained is True

    marker = tmp_path / "pickle_executed"

    class MaliciousPayload:
        def __reduce__(self):
            return os.system, (f"touch {marker}",)

    legacy_path = tmp_path / "legacy.pkl"
    legacy_path.write_bytes(pickle.dumps(MaliciousPayload()))
    with pytest.raises(ValueError, match="pickle legado"):
        ejecutar_proceso_clasificacion_directo(
            pd.DataFrame({"name": ["ACME"]}),
            "name",
            False,
            {"umbral_empresa": 1, "umbral_persona": -1},
            ruta_guardado_modelo=str(legacy_path),
        )
    assert not marker.exists()


def test_optimization_pickle_requires_explicit_trust() -> None:
    engine = OptimizationEngine.__new__(OptimizationEngine)
    with pytest.raises(ValueError, match="trusted=True"):
        engine.load_checkpoint("untrusted.pkl")
